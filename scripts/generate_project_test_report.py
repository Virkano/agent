#!/usr/bin/env python3
"""Generate the project test report as a styled, font-embedded DOCX."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import uuid
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "output/project-documents/中国电信2026年理想公司线上营销智能体平台建设研发工程项目-测试报告.docx"
FONT_PATH = ROOT / "ui/src/styles/font/AlibabaPuHuiTi-3-55-Regular.ttf"
FONT_NAME = "Alibaba PuHuiTi 3.0"

NAVY = "12355B"
BLUE = "1F5FAF"
LIGHT_BLUE = "EAF2FB"
PALE_BLUE = "F5F8FC"
GRAY = "667085"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D0D5DD"
DARK = "1D2939"
GREEN = "16865C"
AMBER = "B76E00"
RED = "B42318"
WHITE = "FFFFFF"


def set_run_font(run, size=10.5, bold=False, color=DARK, font_name=FONT_NAME):
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size=10.5, bold=False, color=DARK):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_table_borders(table, color=MID_GRAY, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_text(cell, text, *, bold=False, color=DARK, size=9.3, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document, headers, rows, widths=None, *, font_size=9.0, header_font_size=9.2):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, (cell, value) in enumerate(zip(header.cells, headers, strict=True)):
        set_cell_text(
            cell,
            value,
            bold=True,
            color=WHITE,
            size=header_font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(cell, NAVY)
        if widths:
            cell.width = Cm(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, (cell, value) in enumerate(zip(row.cells, values, strict=True)):
            align = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, value, size=font_size, align=align)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            if widths:
                cell.width = Cm(widths[col_index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_key_value_table(document, rows):
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for index, (label, value) in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        set_cell_text(row.cells[0], label, bold=True, color=NAVY, size=9.4)
        set_cell_text(row.cells[1], value, size=9.3)
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(13.4)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        if index % 2 == 1:
            set_cell_shading(row.cells[1], PALE_BLUE)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(11 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    run = paragraph.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True, color=NAVY if level == 1 else BLUE)
    if level == 1:
        p_pr = paragraph._p.get_or_add_pPr()
        bottom = OxmlElement("w:pBdr")
        border = OxmlElement("w:bottom")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), BLUE)
        bottom.append(border)
        p_pr.append(bottom)
    return paragraph


def add_body(document, text, *, bold_lead=None, color=DARK):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.45
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, size=10.5, bold=True, color=NAVY)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, size=10.5, color=color)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=color)
    return paragraph


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.left_indent = Cm(0.74 + 0.55 * level)
    paragraph.paragraph_format.first_line_indent = Cm(-0.32)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(text)
    set_run_font(run, size=10.2)
    return paragraph


def add_status_box(document, label, text, color):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=color, size=8)
    set_cell_text(table.cell(0, 0), label, bold=True, color=WHITE, size=10.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), text, bold=True, color=color, size=10.2)
    set_cell_shading(table.cell(0, 0), color)
    set_cell_shading(table.cell(0, 1), "FFFFFF")
    table.cell(0, 0).width = Cm(3.2)
    table.cell(0, 1).width = Cm(13.2)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=8.5, color=GRAY)


def add_hyperlink(paragraph, text, anchor):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    run_pr.append(color)
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def embed_font(docx_path):
    font_key = uuid.uuid4()
    key_bytes = font_key.bytes_le
    font_bytes = bytearray(FONT_PATH.read_bytes())
    for index in range(min(32, len(font_bytes))):
        font_bytes[index] ^= key_bytes[index % 16]

    with ZipFile(docx_path, "r") as source:
        parts = {name: source.read(name) for name in source.namelist()}

    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"

    font_table = etree.fromstring(parts["word/fontTable.xml"])
    font_node = etree.SubElement(font_table, f"{{{w_ns}}}font")
    font_node.set(f"{{{w_ns}}}name", FONT_NAME)
    embed_node = etree.SubElement(font_node, f"{{{w_ns}}}embedRegular")
    embed_node.set(f"{{{r_ns}}}id", "rIdEmbeddedProjectFont")
    embed_node.set(f"{{{w_ns}}}fontKey", "{" + str(font_key).upper() + "}")
    parts["word/fontTable.xml"] = etree.tostring(
        font_table, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    rel_path = "word/_rels/fontTable.xml.rels"
    if rel_path in parts:
        relationships = etree.fromstring(parts[rel_path])
    else:
        relationships = etree.Element(f"{{{rel_ns}}}Relationships", nsmap={None: rel_ns})
    relationship = etree.SubElement(relationships, f"{{{rel_ns}}}Relationship")
    relationship.set("Id", "rIdEmbeddedProjectFont")
    relationship.set("Type", f"{r_ns}/font")
    relationship.set("Target", "fonts/project-report-font.odttf")
    parts[rel_path] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    content_types = etree.fromstring(parts["[Content_Types].xml"])
    if not content_types.xpath("ct:Default[@Extension='odttf']", namespaces={"ct": content_ns}):
        default = etree.SubElement(content_types, f"{{{content_ns}}}Default")
        default.set("Extension", "odttf")
        default.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")
    parts["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    parts["word/fonts/project-report-font.odttf"] = bytes(font_bytes)

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)
    docx_path.write_bytes(output.getvalue())


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(4)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet", "List Bullet 2"):
        style = styles[name]
        style.font.name = FONT_NAME
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("线上营销智能体平台建设研发工程项目  |  测试报告")
    set_run_font(header_run, size=8.5, color=GRAY)
    p_pr = header._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:space"), "4")
    border.set(qn("w:color"), MID_GRAY)
    p_bdr.append(border)
    p_pr.append(p_bdr)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document):
    for _ in range(5):
        document.add_paragraph()
    badge = document.add_paragraph()
    badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = badge.add_run("中国电信  |  研发交付文档")
    set_run_font(run, size=11, bold=True, color=BLUE)
    badge.paragraph_format.space_after = Pt(22)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    title_run = title.add_run("中国电信2026年理想公司\n线上营销智能体平台建设研发工程项目")
    set_run_font(title_run, size=22, bold=True, color=NAVY)

    rule_table = document.add_table(rows=1, cols=3)
    rule_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, cell in enumerate(rule_table.rows[0].cells):
        cell.text = ""
        set_cell_shading(cell, BLUE if index == 1 else LIGHT_BLUE)
        cell.width = Cm(2.6 if index == 1 else 4.5)
        cell.height = Cm(0.12)
    document.add_paragraph()

    report_title = document.add_paragraph()
    report_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    report_title.paragraph_format.space_before = Pt(16)
    report_title.paragraph_format.space_after = Pt(12)
    run = report_title.add_run("测 试 报 告")
    set_run_font(run, size=28, bold=True, color=NAVY)

    version = document.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = version.add_run("版本 V1.0  |  开发阶段综合验收")
    set_run_font(run, size=11, color=GRAY)

    for _ in range(3):
        document.add_paragraph()
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = author.add_run("项目研发组")
    set_run_font(run, size=12, bold=True, color=NAVY)
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date.add_run("2026年9月")
    set_run_font(run, size=11, color=GRAY)
    document.add_page_break()


def add_contents(document):
    add_heading(document, "文档目录", 1)
    items = [
        ("一、测试报告概要", "sec1"),
        ("二、测试执行概况", "sec2"),
        ("三、缺陷与问题分析", "sec3"),
        ("四、各类型测试结果详情", "sec4"),
        ("五、风险评估与建议", "sec5"),
        ("六、测试结论", "sec6"),
    ]
    for text, anchor in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.left_indent = Cm(0.7)
        add_hyperlink(paragraph, text, anchor)
        set_paragraph_font(paragraph, size=11, color=BLUE)
    document.add_paragraph()
    add_status_box(
        document,
        "总体结论",
        "V1 粗粒度功能与开发环境部署验收通过；生产发布条件尚未全部满足。",
        BLUE,
    )
    document.add_page_break()


def build_report():
    document = Document()
    configure_document(document)
    properties = document.core_properties
    properties.title = "中国电信2026年理想公司线上营销智能体平台建设研发工程项目测试报告"
    properties.subject = "开发阶段综合测试与验收结论"
    properties.author = "项目研发组"
    properties.keywords = "线上营销, 智能体, 测试报告, V1, Docker"
    properties.comments = "依据项目现有自动化回归、专项验收、部署验收与安全扫描证据生成。"

    add_cover(document)
    add_contents(document)

    heading = add_heading(document, "一、测试报告概要", 1)
    add_bookmark(heading, "sec1", 1)
    add_key_value_table(
        document,
        [
            ("项目名称", "中国电信2026年理想公司线上营销智能体平台建设研发工程项目"),
            ("测试周期", "2026年9月10日至2026年9月23日；报告编制日期为2026年9月24日"),
            (
                "测试范围",
                "建设方案 X01-X38 粗粒度 V1 功能、服务助手与知识库、推荐与精排、指标告警、AI 报告、事件调度、模型与知识版本、API 编排、故障与场景预演、营销运营、权限审计、Docker 单业务镜像及数据库恢复演练。",
            ),
            (
                "测试类型",
                "功能测试、API 与跨模块集成测试、权限与审计测试、异常与幂等测试、浏览器专项验收、容器部署冒烟、备份恢复演练、安全扫描基线。性能压测按项目当前开发阶段安排后置。",
            ),
            (
                "测试环境",
                "开发主机与 Docker Desktop；业务镜像 telecom-ai-agent:local，仅暴露 8080；PostgreSQL/pgvector、Redis 为外部独立服务；Ollama 为宿主机或独立服务；镜像内运行 Django/Gunicorn、Celery、管理端与对话端静态资源及本地向量服务。",
            ),
            (
                "参考文档",
                "建设方案 v3.0、功能对比表 X01-X38、V1 实施计划、项目详细设计说明书、M02-M12 专项验收文档、全量串联回归台账、Docker 手动打包与运行命令。",
            ),
            (
                "测试结论",
                "核心业务自动化回归与开发环境单业务镜像验收通过，粗粒度 V1 主流程已闭合。由于第三方依赖漏洞、安全镜像扫描缺证、500 并发压测未执行、生产灾备及真实 IVR/晓搭/外部 MLflow 联调未完成，本报告不构成生产上线批准。",
            ),
        ],
    )
    add_status_box(document, "验收判定", "开发阶段通过；进入预上线整改与补充测试。", GREEN)

    document.add_page_break()
    heading = add_heading(document, "二、测试执行概况", 1)
    add_bookmark(heading, "sec2", 2)
    add_heading(document, "（一）测试用例执行统计", 2)
    add_body(
        document,
        "统计采用两个独立自动化批次：M12 跨模块主回归 113 项，以及单业务镜像部署专项 19 项。静态检查、前端构建、浏览器专项操作、安全扫描与恢复演练作为证据批次单独列示，不与自动化用例数重复相加。",
    )
    add_table(
        document,
        ["测试批次", "计划数", "执行数", "通过数", "未通过数", "通过率"],
        [
            ["M12 跨模块主回归", "113", "113", "113", "0", "100.0%"],
            ["单业务镜像部署专项", "19", "19", "19", "0", "100.0%"],
            ["自动化用例合计", "132", "132", "132", "0", "100.0%"],
            ["安全扫描发布门禁", "1 批", "1 批", "0", "1 批阻断", "不适用"],
            ["性能压测", "后置", "0", "0", "未执行", "不适用"],
        ],
        [4.7, 2.1, 2.1, 2.1, 2.5, 2.4],
        font_size=8.8,
    )
    add_body(
        document,
        "核心回归覆盖推荐、服务助手、指标与监控、事件编排与动态路径、训练流水线、知识 ETL、AI 报告、营销运营、IVR 契约、变更控制、晓搭 API 编排及预演。Django 系统检查、迁移差异检查、目标代码静态检查、前端类型检查与生产构建均已通过。",
    )

    add_heading(document, "（二）核心模块执行详情", 2)
    add_table(
        document,
        ["模块", "主要覆盖", "结果", "证据", "边界说明"],
        [
            ["服务助手与知识库", "导入、分段、向量化、混合检索、重排、RAG、引用、合规审核与人工复核", "通过", "M02 / X06-X08、X14", "语义效果与规则覆盖率留待专项优化"],
            ["推荐与精排引擎", "双塔训练、多路召回、TopN 精排、反馈增量、效果监控", "通过", "M07 / X01-X05", "验收反馈为模拟数据，不代表商业转化"],
            ["指标、告警与报告", "动态基线、派单、铃铛通知、归因、订阅、跨模块来源快照", "通过", "M03、M04 / X11-X13、X36-X38", "仅推荐复核待办已接自动采集"],
            ["事件调度", "标准事件、DAG、资源协调、人工节点、动态路径、重试与幂等", "通过", "M06 / X15-X17", "外部生产事件源仍需接入"],
            ["模型与数据管道", "模型版本、训练评估、索引发布回滚、全维监控", "通过", "M08 / X32-X35", "外部 MLflow 待真实联调"],
            ["API 编排与晓搭", "Schema、鉴权、版本、发布、回滚、Mock 调用与审计", "条件通过", "M05 / X10", "真实晓搭网关与凭证未接入"],
            ["变更管控与 IVR", "风险灯、人工确认、通信审计、IVR 消息契约回放", "条件通过", "M09 / X09", "无真实电话线路和 WebSocket 网关"],
            ["故障与场景预演", "申告、处置、数据、告警、回溯、能力沙盒、故障上报与辅助报告", "通过", "M10 / X18-X31", "均为脱敏预演事实"],
            ["营销运营闭环", "场景、活动、任务、客户、触达、转化、质检、全局通知", "通过", "既有营销域专项验收", "未执行真实外呼"],
            ["部署与恢复", "单业务镜像、8080 单端口、健康检查、迁移检查、向量请求、PostgreSQL 隔离恢复", "通过", "M11-05、M11-06", "仅开发环境，不代表生产灾备"],
            ["安全发布门禁", "依赖审计、镜像扫描、安全三同步检查表", "未通过", "M11-04", "依赖发现项和镜像扫描证据未闭合"],
        ],
        [3.0, 5.1, 2.0, 2.8, 3.5],
        font_size=7.6,
        header_font_size=8.0,
    )

    document.add_page_break()
    heading = add_heading(document, "三、缺陷与问题分析", 1)
    add_bookmark(heading, "sec3", 3)
    add_heading(document, "（一）缺陷统计与分布", 2)
    add_body(
        document,
        "在最终核心回归和容器验收批次中，未发现未关闭的 P0/P1 功能缺陷。过程中的测试基础设施与探针问题已完成修复并复验；安全扫描发现项、性能未测项和外部联调缺口按发布风险单独管理，不与功能缺陷混淆。",
    )
    add_table(
        document,
        ["级别", "未关闭功能缺陷", "过程已修复问题", "发布风险项", "结论"],
        [
            ["P0 致命", "0", "0", "0", "无"],
            ["P1 严重", "0", "0", "3", "安全依赖、镜像扫描、性能门禁待闭合"],
            ["P2 一般", "0", "4", "2", "外部联调、生产灾备与兼容性证据待补"],
            ["P3 优化", "0", "0", "2", "前端包体和既有 eval 使用待治理"],
        ],
        [2.4, 3.0, 3.0, 3.0, 5.0],
        font_size=8.4,
    )

    add_heading(document, "（二）关键问题与发布缺口", 2)
    add_table(
        document,
        ["编号", "类别", "问题描述", "状态", "完成标准"],
        [
            ["GAP-001", "依赖安全", "Python 扫描发现 7 个包共 42 条记录；前端发现 25 条记录，其中高危 13 条。", "待整改", "升级或形成限期风险接受，复扫无未接受高危项"],
            ["GAP-002", "镜像安全", "当前源码对应的新镜像缺少可追溯漏洞扫描报告；历史 Docker Scout 因凭据缺失返回错误。", "受阻", "建立源码到镜像对应关系并完成经批准的镜像扫描"],
            ["GAP-003", "性能", "500 并发、接口响应、长稳和容量压测按开发阶段决定后置。", "未执行", "在预上线环境完成性能基线并对照建设方案指标"],
            ["GAP-004", "外部联调", "IVR、晓搭、外部 MLflow、DCOOS/EOP 当前为协议适配或 Mock 证据。", "待外部条件", "使用真实地址与凭证完成成功、失败和审计链路"],
            ["GAP-005", "灾备兼容", "已完成本地隔离恢复，但未验证跨主机灾备、异地备份、Firefox/Edge 与移动终端。", "待补证", "完成生产灾备演练和目标终端兼容性矩阵"],
        ],
        [2.0, 2.4, 5.6, 2.3, 4.1],
        font_size=7.8,
        header_font_size=8.0,
    )

    add_heading(document, "（三）问题趋势分析", 2)
    add_bullet(document, "M12 首轮合并回归曾因测试模块导入时启动后台调度线程导致计数串扰，关闭测试进程后台调度后，113 项核心回归稳定通过。")
    add_bullet(document, "数据库恢复演练首次因 pg_restore 目标参数不完整失败，修正后源库、临时恢复库的表、大对象、文件记录和字节指纹一致。")
    add_bullet(document, "跨模块验收探针出现过反向关系名和精排任务关系读取错误，均已修复并写入错误经验文档，重复执行保持幂等。")
    add_bullet(document, "当前剩余问题集中在安全、性能、生产灾备和外部系统条件，不是核心业务状态机的集中性失败。")

    document.add_page_break()
    heading = add_heading(document, "四、各类型测试结果详情", 1)
    add_bookmark(heading, "sec4", 4)
    add_heading(document, "（一）功能与集成测试结果", 2)
    add_body(
        document,
        "功能测试以建设方案 X01-X38 和 Word 补充要求为范围，采用专项服务测试、API 测试、真实数据库事实、浏览器关键交互以及同一工作空间的跨模块链路进行验证。以下结果均为粗粒度 V1 结论。",
    )
    add_table(
        document,
        ["建设能力组", "编号", "验证链路", "结果"],
        [
            ["推荐召回与精排", "X01-X05", "训练到发布、推理到融合精排、曝光反馈、增量策略与监控", "通过"],
            ["知识库与服务助手", "X06-X08、X14", "资料导入到索引、检索生成、引用、兜底、拦截与人工复核", "通过"],
            ["变更、IVR 与 API 适配", "X09-X10", "风险确认、通信审计、契约回放、Schema/鉴权/版本与 Mock 调用", "条件通过"],
            ["指标预警与配置", "X11-X13", "指标导入、动态阈值、告警、值班派单、处理关闭与趋势下钻", "通过"],
            ["事件编排与动态路径", "X15-X17", "事件校验、DAG、资源门禁、人工节点、重试降级和轨迹", "通过"],
            ["故障与场景预演", "X18-X31", "数据生成、实例执行、时序同步、记录查询、两类回溯和报告导出", "通过"],
            ["模型、训练与 ETL", "X32-X35", "生命周期、训练评估、索引版本发布回滚及监控告警", "条件通过"],
            ["AI 分析报告", "X36-X38", "数据计算、归因、模板、报告、订阅、全局通知及多格式导出", "通过"],
        ],
        [4.1, 3.0, 7.1, 2.2],
        font_size=8.2,
    )
    add_body(
        document,
        "全量串联使用同一工作空间和营销任务形成推荐、触达反馈、事件、指标、派单、全局铃铛、报告来源、人工复核、模型版本、知识版本、IVR 契约、晓搭 Mock 与预演实例的可追踪事实链。所有验收模拟均保留明确标识，不作为真实客户转化或外部系统已接通的证据。",
    )

    add_heading(document, "（二）性能测试结果", 2)
    add_table(
        document,
        ["测试场景", "建设目标/关注点", "本轮实际结果", "判定", "说明"],
        [
            ["500 并发访问", "建设方案并发与响应指标", "未执行", "待测", "按开发阶段决定后置到预上线"],
            ["接口与页面响应", "接口、查询、页面加载时延", "仅完成冒烟与功能可达", "不判定", "不得由 HTTP 200 推断性能达标"],
            ["长稳运行", "7×24 小时稳定性", "未执行", "待测", "需独立长稳环境和监控基线"],
            ["本地备份恢复", "数据可恢复性", "备份约 0.536 秒；恢复约 12.003 秒", "通过", "仅 57 MB 本地源库，不是生产 RTO/RPO 承诺"],
        ],
        [3.1, 3.8, 3.7, 2.0, 4.8],
        font_size=8.0,
    )

    add_heading(document, "（三）安全测试结果", 2)
    add_bullet(document, "生产配置预检能够拒绝缺失密钥、开发默认值、通配域名和 DEBUG；敏感配置不写入镜像和日志。")
    add_bullet(document, "未认证访问管理 API 返回 HTTP 401；工作空间隔离、权限校验、关键状态审计和幂等约束已有自动化覆盖。")
    add_bullet(document, "用户与营销客户直接读取接口接入追加式读取审计；查询值、联系方式、凭证和业务正文不进入通用日志。")
    add_bullet(document, "安全扫描能力已经建立，但因依赖漏洞发现项、镜像扫描缺证和安全三同步未全部完成，发布安全门禁为未通过。")

    add_heading(document, "（四）兼容性测试结果", 2)
    add_table(
        document,
        ["测试场景", "已覆盖对象", "执行结果"],
        [
            ["浏览器交互", "桌面 Chromium / Codex 内置浏览器", "登录、导航、长页面滚动、抽屉、告警铃铛和关键表单专项通过"],
            ["容器运行", "Docker Desktop 上的 Linux 业务镜像", "健康检查、管理端、对话端、Worker 和向量服务可用，仅暴露 8080"],
            ["外部依赖配置", "PostgreSQL/pgvector、Redis、Ollama", "通过环境变量接入；容器健康检查与 Ollama 模型列表验证通过"],
            ["其他浏览器与终端", "Firefox、Edge、平板、移动端", "本轮未形成正式兼容性证据，预上线补测"],
        ],
        [3.4, 5.2, 7.8],
        font_size=8.2,
    )

    add_heading(document, "（五）可靠性与恢复测试结果", 2)
    add_table(
        document,
        ["测试场景", "测试方法", "实际结果", "结论"],
        [
            ["容器健康检查", "检查 PostgreSQL、Redis、向量端口及 Web 路由", "容器 running/healthy，/healthz 返回 200", "通过"],
            ["数据库恢复", "pg_dump 含大对象，恢复到唯一临时数据库并逐项核对", "194 张表、13 个大对象、18 条文件记录及 1,887,340 字节一致", "通过"],
            ["关键链路幂等", "重复执行推荐事件、指标采集、报告、处置和模拟探针", "复用既有事实，不重复创建事件、工单或报告", "通过"],
            ["生产灾备", "跨主机、异地备份、持续写入、故障切换", "未执行", "待测"],
        ],
        [3.0, 5.2, 6.1, 2.1],
        font_size=7.6,
        header_font_size=8.0,
    )

    document.add_page_break()
    heading = add_heading(document, "五、风险评估与建议", 1)
    add_bookmark(heading, "sec5", 5)
    add_heading(document, "（一）残留风险", 2)
    add_table(
        document,
        ["风险项", "风险等级", "影响范围", "应对建议"],
        [
            ["第三方依赖存在未处置发现项", "高", "后端、前端、文档解析与模型相关组件", "逐条核查可达性与补丁版本，分批升级回归；例外必须有期限和风险接受人"],
            ["当前源码镜像缺少漏洞扫描证据", "高", "生产发布与供应链追踪", "重建当前源码镜像，固定镜像摘要，使用已批准扫描工具生成并归档报告"],
            ["性能和长稳未验证", "高", "并发、响应时间、容量和稳定性", "在接近生产配置的预上线环境完成 500 并发、长稳和关键接口基线"],
            ["外部平台仅完成 Mock 或协议回放", "中", "IVR、晓搭、MLflow、DCOOS/EOP", "准备真实地址、账号与网络白名单，执行成功、失败、超时、重试和审计场景"],
            ["生产灾备与异地备份未验证", "中", "数据恢复与连续性", "明确 RPO/RTO、保留周期和责任人，执行跨主机恢复及业务可读性检查"],
            ["浏览器与移动端覆盖不足", "中", "终端兼容性和易用性", "补齐 Chrome、Edge、Firefox 及目标分辨率测试矩阵"],
            ["前端包体较大且存在既有 eval 使用", "中", "加载性能和安全审查", "拆包并治理动态表单实现，完成构建、功能回归和安全复扫"],
        ],
        [5.0, 2.3, 4.0, 5.2],
        font_size=8.1,
    )

    add_heading(document, "（二）优化与后续测试建议", 2)
    add_bullet(document, "先关闭发布门禁：依赖整改、当前镜像扫描、安全三同步证据和生产配置预检应优先于功能纵向优化。")
    add_bullet(document, "在预上线环境建立可重复的性能基线，保留数据规模、并发模型、硬件配置、采样窗口和原始报告，避免仅记录结论。")
    add_bullet(document, "外部系统联调按接口逐项验收鉴权、Schema、超时、重试、幂等、脱敏审计与回滚，不以 Mock 成功替代真实接通。")
    add_bullet(document, "将 X01-X38 与测试用例、实现模块、证据文档建立追踪矩阵，后续每次版本报告只更新受影响项和新增证据。")

    document.add_page_break()
    heading = add_heading(document, "六、测试结论", 1)
    add_bookmark(heading, "sec6", 6)
    add_status_box(document, "阶段结论", "开发阶段 V1 综合验收通过，生产上线暂不批准。", BLUE)
    document.add_paragraph()
    add_body(
        document,
        "功能层面：建设方案 X01-X38 已形成粗粒度 V1 能力，113 项跨模块主回归和 19 项部署专项自动化用例全部通过；推荐、知识问答、指标告警、事件编排、预演、模型与知识版本、AI 报告和营销运营已完成本地事实链串联。",
        bold_lead="功能层面：",
    )
    add_body(
        document,
        "部署层面：单业务镜像在开发环境真实运行通过，仅暴露 8080，PostgreSQL、Redis 和 Ollama 均按外部服务配置接入；健康检查、管理端、对话端、Celery、本地向量请求及迁移状态正常。",
        bold_lead="部署层面：",
    )
    add_body(
        document,
        "质量层面：权限、审计、异常、幂等和恢复已有可重复证据；本地 PostgreSQL 含大对象隔离恢复演练通过。安全扫描已能执行，但发现项尚未整改，镜像扫描和安全三同步尚未闭合。",
        bold_lead="质量层面：",
    )
    add_body(
        document,
        "边界层面：500 并发与长稳测试尚未执行；真实 IVR、晓搭、外部 MLflow、DCOOS/EOP、跨主机灾备及完整浏览器兼容性仍待外部条件或预上线环境验证。上述项目不得被解释为已经通过。",
        bold_lead="边界层面：",
    )
    add_body(
        document,
        "整体结论：系统已满足粗粒度 V1 开发阶段演示、联调和继续部署准备的条件，建议进入预上线整改与补充测试。完成安全高风险项处置、当前镜像扫描、性能基线、生产灾备和必要外部联调后，再组织生产上线评审。",
        bold_lead="整体结论：",
    )

    document.add_page_break()
    add_heading(document, "主要证据索引", 2)
    add_body(
        document,
        "下表所列 M11/M12 验收文档均位于 docs/测试/；资产清单和详细设计说明书使用表内给出的项目相对位置。",
    )
    add_table(
        document,
        ["证据", "位置", "用途"],
        [
            ["全量串联回归台账", "全量串联M12-01回归台账.md", "113 项主回归与 M12-01 至 M12-10 跨模块事实链"],
            ["单业务镜像部署验收", "单业务镜像部署M11-06验收文档.md", "19 项部署专项、容器健康与 8080 单端口"],
            ["安全扫描与三同步", "安全扫描与三同步M11-04验收文档.md", "依赖发现项、镜像扫描缺口与门禁结论"],
            ["数据库恢复演练", "数据库备份恢复演练M11-05验收文档.md", "PostgreSQL 与大对象本地隔离恢复证据"],
            ["资产与接口清单", "output/m11-02/asset-inventory.json", "模型、接口分层、敏感字段提示与外部接入状态"],
            ["详细设计说明书", "output/project-documents/详设说明书.docx", "系统设计与模块说明；实际文件名见项目文档目录"],
        ],
        [4.0, 7.6, 5.0],
        font_size=7.7,
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    embed_font(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_report())
