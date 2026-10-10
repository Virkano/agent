from datetime import date, datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
WHITE_BOLD_FONT = Font(color="FFFFFF", bold=True)
BOLD_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)

CAMPAIGN_STATUS_LABELS = {
    "DRAFT": "草稿",
    "ACTIVE": "已上线",
    "PAUSED": "已暂停",
    "ARCHIVED": "已归档",
}
APPROVAL_STATUS_LABELS = {
    "DRAFT": "未提交",
    "PENDING": "待审批",
    "APPROVED": "已通过",
    "REJECTED": "已驳回",
}
RECOMMENDATION_LABELS = {
    "CONTINUE_GRAY": "继续灰度",
    "EXPAND": "扩大触达",
    "PAUSE": "暂停活动",
}


def safe_text(value):
    if value is None:
        return ""
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _excel_value(value):
    if isinstance(value, (datetime, date, int, float)):
        return value
    if isinstance(value, (list, tuple)):
        return safe_text("、".join(str(item) for item in value))
    return safe_text(value)


def _apply_border_and_wrap(sheet, min_row, max_row, min_col, max_col):
    for row in sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_section(sheet, row, title):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = sheet.cell(row=row, column=1, value=title)
    cell.fill = SECTION_FILL
    cell.font = BOLD_FONT


def _write_pair_row(sheet, row, first_label, first_value, second_label, second_value):
    values = (first_label, first_value, second_label, second_value)
    for column, value in enumerate(values, 1):
        sheet.cell(row=row, column=column, value=_excel_value(value))
    sheet.cell(row=row, column=1).font = BOLD_FONT
    sheet.cell(row=row, column=3).font = BOLD_FONT


def _write_summary(workbook, effect):
    sheet = workbook.active
    sheet.title = "复盘汇总"
    campaign = effect["campaign"]
    metrics = effect["metrics"]
    recommendation = effect["recommendation"]

    sheet.merge_cells("A1:D1")
    title = sheet["A1"]
    title.value = safe_text(f"{campaign.get('name', '')} - 活动效果复盘")
    title.fill = HEADER_FILL
    title.font = Font(color="FFFFFF", bold=True, size=16)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    now = timezone.localtime().replace(tzinfo=None)
    _write_pair_row(sheet, 2, "导出时间", now, "统计口径", effect.get("scope", ""))
    sheet["B2"].number_format = "yyyy-mm-dd hh:mm:ss"

    _write_section(sheet, 3, "活动信息")
    _write_pair_row(sheet, 4, "活动编码", campaign.get("code", ""), "活动版本", campaign.get("version", 1))
    _write_pair_row(
        sheet,
        5,
        "活动状态",
        CAMPAIGN_STATUS_LABELS.get(campaign.get("status"), campaign.get("status", "")),
        "审批状态",
        APPROVAL_STATUS_LABELS.get(
            campaign.get("approval_status"),
            campaign.get("approval_status", ""),
        ),
    )

    _write_section(sheet, 6, "核心指标")
    _write_pair_row(sheet, 7, "关联任务数", metrics.get("task_count", 0), "归因会话数", metrics.get("chat_count", 0))
    _write_pair_row(sheet, 8, "客户数", metrics.get("customer_count", 0), "意向客户数", metrics.get("intent_count", 0))
    _write_pair_row(sheet, 9, "已办理数", metrics.get("converted_count", 0), "转人工数", metrics.get("manual_count", 0))
    _write_pair_row(
        sheet, 10, "无意向数", metrics.get("no_intent_count", 0), "意向率", metrics.get("intent_rate", 0) / 100
    )
    _write_pair_row(sheet, 11, "转人工率", metrics.get("manual_rate", 0) / 100, "", "")
    sheet["D10"].number_format = "0.0%"
    sheet["B11"].number_format = "0.0%"

    _write_section(sheet, 12, "系统建议")
    _write_pair_row(
        sheet,
        13,
        "建议动作",
        RECOMMENDATION_LABELS.get(recommendation.get("action"), recommendation.get("action", "")),
        "建议下一批规模",
        recommendation.get("suggested_next_batch_size", 0),
    )
    sheet["A14"] = "判断依据"
    sheet["A14"].font = BOLD_FONT
    sheet.merge_cells("B14:D14")
    sheet["B14"] = safe_text(recommendation.get("reason", ""))

    _apply_border_and_wrap(sheet, 1, 14, 1, 4)
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 54
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False


def _write_segments(workbook, title, rows):
    sheet = workbook.create_sheet(title)
    headers = [
        "分组",
        "客户数",
        "意向数",
        "已办理数",
        "转人工数",
        "无意向数",
        "意向率",
        "办理率",
    ]
    sheet.append(headers)
    for item in rows:
        sheet.append(
            [
                safe_text(item.get("name", "")),
                item.get("customer_count", 0),
                item.get("intent_count", 0),
                item.get("converted_count", 0),
                item.get("manual_count", 0),
                item.get("no_intent_count", 0),
                item.get("intent_rate", 0) / 100,
                item.get("conversion_rate", 0) / 100,
            ]
        )

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if sheet.max_row > 1:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=7).number_format = "0.0%"
            sheet.cell(row=row, column=8).number_format = "0.0%"
    _apply_border_and_wrap(sheet, 1, sheet.max_row, 1, 8)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{sheet.max_row}"
    sheet.column_dimensions["A"].width = 24
    for column in ("B", "C", "D", "E", "F", "G", "H"):
        sheet.column_dimensions[column].width = 14
    sheet.sheet_view.showGridLines = False


def build_campaign_effect_workbook(effect):
    workbook = Workbook()
    _write_summary(workbook, effect)
    breakdowns = effect.get("breakdowns") or {}
    _write_segments(workbook, "省份效果", breakdowns.get("province") or [])
    _write_segments(workbook, "套餐效果", breakdowns.get("current_plan") or [])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
