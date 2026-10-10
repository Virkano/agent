"""Validation and parsing helpers for task-scoped marketing customer imports."""

from __future__ import annotations

import csv
import json
import re
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook

from common.exception.app_exception import AppApiException


IMPORT_FIELDS = (
    "name",
    "phone",
    "province",
    "current_plan",
    "monthly_data_usage",
    "monthly_call_usage",
    "intent",
    "remark",
)
REQUIRED_IMPORT_FIELDS = {"phone"}
FIELD_LABELS = {
    "name": "客户名称",
    "phone": "手机号/客户标识",
    "province": "省份",
    "current_plan": "当前套餐",
    "monthly_data_usage": "近月流量",
    "monthly_call_usage": "近月通话",
    "intent": "初始意向",
    "remark": "备注",
}
HEADER_ALIASES = {
    "name": ("客户名称", "姓名", "客户", "name"),
    "phone": ("手机号", "手机号码", "客户标识", "客户id", "客户ID", "phone", "customer_id"),
    "province": ("省份", "归属省", "province"),
    "current_plan": ("当前套餐", "套餐", "current_plan", "plan"),
    "monthly_data_usage": ("近月流量", "月流量", "monthly_data_usage", "data_usage"),
    "monthly_call_usage": ("近月通话", "月通话", "monthly_call_usage", "call_usage"),
    "intent": ("初始意向", "意向", "intent"),
    "remark": ("备注", "remark", "note"),
}
FIELD_LIMITS = {
    "name": 128,
    "phone": 64,
    "province": 64,
    "current_plan": 128,
    "monthly_data_usage": 64,
    "monthly_call_usage": 64,
    "intent": 64,
    "remark": 2000,
}
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000


def mask_identifier(value: str) -> str:
    if not value:
        return ""
    if len(value) >= 7:
        return f"{value[:3]}****{value[-4:]}"
    return value


def normalize_identifier(value: Any) -> str:
    return re.sub(r"\s+", "", stringify_cell(value))


def stringify_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_mapping(raw_mapping: Any) -> dict[str, str]:
    if raw_mapping in (None, ""):
        return {}
    if isinstance(raw_mapping, str):
        try:
            raw_mapping = json.loads(raw_mapping)
        except json.JSONDecodeError as exc:
            raise AppApiException(400, "字段映射格式不正确") from exc
    if not isinstance(raw_mapping, dict):
        raise AppApiException(400, "字段映射格式不正确")
    return {field: stringify_cell(column) for field, column in raw_mapping.items() if field in IMPORT_FIELDS and column}


def _read_csv(content: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return list(csv.reader(StringIO(content.decode(encoding))))
        except UnicodeDecodeError:
            continue
    raise AppApiException(400, "CSV 文件编码不支持，请使用 UTF-8 或 GB18030")


def _read_xlsx(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise AppApiException(400, "Excel 文件无法读取，请确认文件未损坏且为 .xlsx 格式") from exc


def parse_customer_import_file(uploaded_file) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    if uploaded_file is None:
        raise AppApiException(400, "请选择要导入的 CSV 或 XLSX 文件")
    file_name = stringify_cell(getattr(uploaded_file, "name", ""))
    suffix = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if suffix not in {"csv", "xlsx"}:
        raise AppApiException(400, "仅支持 CSV 和 XLSX 客户名单文件")
    if getattr(uploaded_file, "size", 0) > MAX_FILE_SIZE:
        raise AppApiException(400, "导入文件不能超过 5MB")
    content = uploaded_file.read()
    if not content:
        raise AppApiException(400, "导入文件为空")
    rows = _read_csv(content) if suffix == "csv" else _read_xlsx(content)
    non_empty_rows = [row for row in rows if any(stringify_cell(value) for value in row)]
    if len(non_empty_rows) < 2:
        raise AppApiException(400, "导入文件至少需要表头和一条客户数据")
    headers = [stringify_cell(value) for value in non_empty_rows[0]]
    if not any(headers):
        raise AppApiException(400, "导入文件缺少表头")
    body_rows = [
        (index, row) for index, row in enumerate(rows[1:], start=2) if any(stringify_cell(value) for value in row)
    ]
    if len(body_rows) > MAX_IMPORT_ROWS:
        raise AppApiException(400, f"一次最多导入 {MAX_IMPORT_ROWS} 条客户数据")
    return headers, body_rows


def resolve_mapping(headers: list[str], raw_mapping: Any = None) -> dict[str, str]:
    mapping = parse_mapping(raw_mapping)
    header_lookup = {header.casefold(): header for header in headers if header}
    for field in IMPORT_FIELDS:
        if mapping.get(field):
            continue
        for alias in HEADER_ALIASES[field]:
            column = header_lookup.get(alias.casefold())
            if column:
                mapping[field] = column
                break
    invalid_columns = [column for column in mapping.values() if column not in headers]
    if invalid_columns:
        raise AppApiException(400, f"字段映射包含文件中不存在的列：{invalid_columns[0]}")
    if not mapping.get("phone"):
        raise AppApiException(400, "请将“手机号/客户标识”映射到导入文件中的一列")
    return mapping


def _row_value(row: list[Any], header_index: dict[str, int], column: str) -> str:
    index = header_index.get(column)
    return stringify_cell(row[index]) if index is not None and index < len(row) else ""


def validate_import_rows(
    headers: list[str],
    source_rows: list[tuple[int, list[Any]]],
    mapping: dict[str, str],
    existing_identifiers: set[str],
) -> list[dict[str, Any]]:
    header_index = {header: index for index, header in enumerate(headers)}
    seen_identifiers: set[str] = set()
    results = []
    for row_number, source_row in source_rows:
        data = {
            field: _row_value(source_row, header_index, mapping[field]) if mapping.get(field) else ""
            for field in IMPORT_FIELDS
        }
        data["phone"] = normalize_identifier(data["phone"])
        errors = []
        for field in REQUIRED_IMPORT_FIELDS:
            if not data[field]:
                errors.append(f"{FIELD_LABELS[field]}不能为空")
        for field, limit in FIELD_LIMITS.items():
            if len(data[field]) > limit:
                errors.append(f"{FIELD_LABELS[field]}不能超过 {limit} 个字符")
        identifier = data["phone"]
        status = "READY"
        reason = ""
        if errors:
            status = "ERROR"
            reason = "；".join(errors)
        elif identifier in existing_identifiers:
            status = "SKIPPED"
            reason = "该客户标识已存在于当前营销任务"
        elif identifier in seen_identifiers:
            status = "SKIPPED"
            reason = "与本文件前序行客户标识重复"
        else:
            seen_identifiers.add(identifier)
        results.append(
            {
                "row_number": row_number,
                "data": data,
                "phone_masked": mask_identifier(identifier),
                "status": status,
                "reason": reason,
            }
        )
    return results
