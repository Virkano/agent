from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from concurrent.futures import TimeoutError as FutureTimeoutError
from datetime import date, timedelta
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from application.models import Application
from application.models.application_chat import Chat, ChatRecord
from application.chat_pipeline.step.chat_step.impl.base_chat_step import get_marketing_runtime_fallback
from application.chat_pipeline.step.search_dataset_step.impl.base_search_dataset_step import BaseSearchDatasetStep
from common.exception.app_exception import AppApiException
from knowledge.models import File, Knowledge
from marketing.insight import analyze_customer_message, update_customer_conversation_insight
from marketing.models import (
    MarketingAlert,
    MarketingAlertAction,
    MarketingAlertActionRecord,
    MarketingAlertAttachment,
    MarketingAlertHandlingRound,
    MarketingAlertRoundEndReason,
    MarketingAlertRule,
    MarketingAlertRuleChange,
    MarketingAlertStatus,
    MarketingAlertSubscriber,
    MarketingNotification,
    MarketingNotificationStatus,
    MarketingAnswerManualReviewAction,
    MarketingAnswerManualReviewStatus,
    MarketingAnswerPreflightRecord,
    MarketingAnswerQualityReviewActionRecord,
    MarketingCampaign,
    MarketingCampaignApprovalStatus,
    MarketingCampaignStatus,
    MarketingCampaignVersion,
    MarketingAnswerQualityReview,
    MarketingAnswerQualityStatus,
    MarketingAnalysisReport,
    MarketingAttributionMethod,
    MarketingReportAttributionDetail,
    MarketingReportDelivery,
    MarketingReportGenerationRun,
    MarketingReportStatus,
    MarketingReportSubscription,
    MarketingReportTemplate,
    MarketingPreflightAction,
    MarketingQualityPolicy,
    MarketingQualityPolicyMode,
    MarketingConversionRecord,
    MarketingConversionOutcome,
    MarketingCustomer,
    MarketingCustomerImportBatch,
    MarketingCustomerStatus,
    MarketingScenario,
    MarketingTask,
    MarketingTaskStatus,
    MarketingTouchBatch,
    MarketingTouchBatchApprovalStatus,
    MarketingTouchBatchStatus,
    MarketingTouchExecution,
    MarketingTouchExecutionResult,
)
from marketing.attribution import build_marketing_source
from marketing.exporting import build_campaign_effect_workbook
from marketing.report_exports import build_docx_export, build_html_export, build_pdf_export
from marketing.recommendation import (
    get_customer_recommendation_answer,
    select_recommendation,
    validate_recommendation_policy,
)
from marketing.quality import evaluate_marketing_answer
from marketing.preflight import (
    INTERCEPTED_CUSTOMER_ANSWER,
    prepare_marketing_preflight,
    save_marketing_preflight,
    should_buffer_marketing_preflight,
)
from marketing.serializers import (
    MarketingAlertService,
    MarketingAnalysisReportService,
    MarketingReportSubscriptionService,
    MarketingNotificationService,
    MarketingCampaignService,
    MarketingCustomerService,
    MarketingOperationsService,
    MarketingQualityReviewService,
    MarketingQualityPolicyService,
    MarketingQualityRuleEffectService,
    MarketingScenarioService,
    MarketingTaskService,
    MarketingTouchBatchService,
    scan_all_marketing_alerts,
)
from marketing.serializers.report_subscription import calculate_next_run
from models_provider.models import Model
from users.models import User


class MarketingAnalysisReportServiceTests(TestCase):
    def setUp(self):
        self.workspace_id = "report-workspace"
        self.other_workspace_id = "other-workspace"
        self.user = User.objects.create(
            username="report_tester",
            nick_name="报告测试员",
            password="x",
            role="USER",
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id,
            name="报告测试场景",
            user=self.user,
        )
        self.task = MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="报告测试任务",
            scenario=self.scenario,
            user=self.user,
            channel="ONLINE",
            target_count=10,
        )
        self.customer_a = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="客户A",
            phone="13800000001",
        )
        self.customer_b = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="客户B",
            phone="13800000002",
        )
        self.batch = MarketingTouchBatch.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            user=self.user,
            name="首轮触达",
            selected_count=2,
        )
        MarketingTouchExecution.objects.create(
            workspace_id=self.workspace_id,
            batch=self.batch,
            customer=self.customer_a,
            user=self.user,
            result=MarketingTouchExecutionResult.CONTACTED,
        )
        MarketingTouchExecution.objects.create(
            workspace_id=self.workspace_id,
            batch=self.batch,
            customer=self.customer_b,
            user=self.user,
            result=MarketingTouchExecutionResult.NO_ANSWER,
        )
        MarketingConversionRecord.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            customer=self.customer_a,
            batch=self.batch,
            user=self.user,
            outcome=MarketingConversionOutcome.CONVERTED,
        )
        self.service = MarketingAnalysisReportService(self.workspace_id, self.user)

    def _payload(self, **overrides):
        payload = {
            "task_id": str(self.task.id),
            "period_start": timezone.localdate().isoformat(),
            "period_end": timezone.localdate().isoformat(),
            "attributed_revenue": "1000.00",
            "marketing_cost": "200.00",
            "variable_cost": "300.00",
            "use_llm": False,
        }
        payload.update(overrides)
        return payload

    def test_rule_report_generates_metrics_html_and_audit_run(self):
        report = self.service.generate(self._payload())

        self.assertEqual(report["status"], MarketingReportStatus.SUCCESS)
        self.assertEqual(report["metrics"]["exposure_count"], 2)
        self.assertEqual(report["metrics"]["reached_customer_count"], 1)
        self.assertEqual(report["metrics"]["touch_rate"], 50.0)
        self.assertEqual(report["metrics"]["converted_customer_count"], 1)
        self.assertEqual(report["metrics"]["conversion_rate"], 100.0)
        self.assertEqual(report["metrics"]["marginal_contribution"], 700.0)
        self.assertEqual(report["metrics"]["roi"], 400.0)
        self.assertIn("核心经营指标", report["html_content"])
        lineage = report["input_snapshot"]["source_lineage"]
        self.assertEqual(lineage["scope"], "TASK_PERIOD_ASSOCIATION")
        self.assertEqual(lineage["task_id"], str(self.task.id))
        self.assertEqual(
            set(lineage["touch_executions"]["ids"]),
            {str(row.id) for row in MarketingTouchExecution.objects.filter(batch=self.batch)},
        )
        self.assertEqual(lineage["touch_executions"]["total_count"], 2)
        self.assertEqual(lineage["recommendation_rankings"]["total_count"], 0)
        self.assertEqual(MarketingReportTemplate.objects.filter(workspace_id=self.workspace_id).count(), 1)
        self.assertEqual(MarketingReportGenerationRun.objects.get(report_id=report["id"]).attempt, 1)

    def test_source_lineage_is_workspace_scoped_and_frozen(self):
        foreign = MarketingTouchExecution.objects.create(
            workspace_id=self.other_workspace_id,
            batch=self.batch,
            customer=self.customer_a,
            user=self.user,
            result=MarketingTouchExecutionResult.CONTACTED,
        )
        report = self.service.generate(self._payload())
        original_ids = report["input_snapshot"]["source_lineage"]["touch_executions"]["ids"]
        self.assertNotIn(str(foreign.id), original_ids)
        later = MarketingTouchExecution.objects.create(
            workspace_id=self.workspace_id,
            batch=self.batch,
            customer=self.customer_a,
            user=self.user,
            result=MarketingTouchExecutionResult.CONTACTED,
        )
        persisted = MarketingAnalysisReport.objects.get(id=report["id"])
        self.assertEqual(persisted.input_snapshot["source_lineage"]["touch_executions"]["ids"], original_ids)
        self.assertNotIn(str(later.id), original_ids)
        self.assertEqual(
            MarketingReportGenerationRun.objects.get(report=persisted).input_snapshot,
            persisted.input_snapshot,
        )

    def test_llm_failure_is_saved_and_can_be_retried(self):
        model = Model.objects.create(
            workspace_id=self.workspace_id,
            user=self.user,
            name="本地分析模型",
            status="SUCCESS",
            model_type="LLM",
            model_name="qwen3:8b",
            provider="OLLAMA",
            credential="encrypted-placeholder",
        )
        llm = Mock()
        llm.invoke.side_effect = [RuntimeError("local model unavailable"), SimpleNamespace(content="建议聚焦高转化客群。")]
        payload = self._payload(use_llm=True, model_id=str(model.id))

        with patch(
            "marketing.serializers.report.get_model_instance_by_model_workspace_id",
            return_value=llm,
        ):
            failed = self.service.generate(payload)
            retried = self.service.retry(failed["id"])

        self.assertEqual(failed["status"], MarketingReportStatus.FAILED)
        self.assertEqual(failed["error_message"], "智能解读模型调用失败，请检查模型服务和配置后重试。")
        self.assertEqual(retried["status"], MarketingReportStatus.SUCCESS)
        self.assertEqual(retried["retry_count"], 1)
        self.assertIn("高转化客群", retried["narrative"])
        attempts = list(
            MarketingReportGenerationRun.objects.filter(report_id=failed["id"]).values_list("attempt", flat=True)
        )
        self.assertEqual(sorted(attempts), [1, 2])

    def test_foreign_workspace_task_is_rejected(self):
        other_scenario = MarketingScenario.objects.create(
            workspace_id=self.other_workspace_id,
            name="其它空间场景",
            user=self.user,
        )
        other_task = MarketingTask.objects.create(
            workspace_id=self.other_workspace_id,
            name="其它空间任务",
            scenario=other_scenario,
            user=self.user,
        )

        with self.assertRaises(AppApiException) as context:
            self.service.generate(self._payload(task_id=str(other_task.id)))

        self.assertEqual(context.exception.code, 400)
        self.assertFalse(MarketingAnalysisReport.objects.exists())

    def test_report_page_is_workspace_isolated(self):
        current = self.service.generate(self._payload())
        other_service = MarketingAnalysisReportService(self.other_workspace_id, self.user)
        other_service.generate(
            {
                "period_start": timezone.localdate().isoformat(),
                "period_end": timezone.localdate().isoformat(),
            }
        )

        page = self.service.page(1, 20, {})

        self.assertEqual(page["total"], 1)
        self.assertEqual(page["records"][0]["id"], current["id"])

    def test_last_touch_and_linear_attribution_are_persisted(self):
        second_batch = MarketingTouchBatch.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            user=self.user,
            name="二次触达",
            selected_count=1,
        )
        MarketingTouchExecution.objects.create(
            workspace_id=self.workspace_id,
            batch=second_batch,
            customer=self.customer_a,
            user=self.user,
            result=MarketingTouchExecutionResult.CONTACTED,
        )
        MarketingConversionRecord.objects.all().delete()
        conversion = MarketingConversionRecord.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            customer=self.customer_a,
            batch=second_batch,
            user=self.user,
            outcome=MarketingConversionOutcome.CONVERTED,
        )

        last_touch = self.service.generate(self._payload(attribution_method=MarketingAttributionMethod.LAST_TOUCH))
        linear = self.service.generate(self._payload(attribution_method=MarketingAttributionMethod.LINEAR))

        last_rows = MarketingReportAttributionDetail.objects.filter(report_id=last_touch["id"])
        linear_rows = MarketingReportAttributionDetail.objects.filter(report_id=linear["id"])
        self.assertEqual(last_rows.count(), 1)
        self.assertEqual(last_rows.get().batch_id, second_batch.id)
        self.assertEqual(linear_rows.count(), 2)
        self.assertEqual(sum(row.weight for row in linear_rows), 1)
        self.assertEqual(sum(row.attributed_revenue for row in linear_rows), 1000)
        self.assertTrue(linear_rows.filter(conversion=conversion).exists())

    def test_report_export_formats_are_valid(self):
        generated = self.service.generate(self._payload())
        report = MarketingAnalysisReport.objects.get(id=generated["id"])

        html_bytes = build_html_export(report)
        pdf_bytes = build_pdf_export(report)
        docx_bytes = build_docx_export(report)

        self.assertIn("<!doctype html>", html_bytes.decode("utf-8"))
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        self.assertTrue(docx_bytes.startswith(b"PK"))

    def test_subscription_can_be_saved_and_run_into_global_notification(self):
        template = self.service._ensure_default_template()
        subscription_service = MarketingReportSubscriptionService(self.workspace_id, self.user)
        saved = subscription_service.save(
            {
                "name": "每日报告",
                "template_id": str(template.id),
                "task_id": str(self.task.id),
                "frequency": "DAILY",
                "run_time": "09:00:00",
                "period_days": 7,
                "attribution_method": "LAST_TOUCH",
                "use_llm": False,
                "email_enabled": False,
                "recipient_emails": [],
                "enabled": True,
            }
        )

        result = subscription_service.run_now(saved["id"])

        subscription = MarketingReportSubscription.objects.get(id=saved["id"])
        self.assertEqual(result["report_status"], MarketingReportStatus.SUCCESS)
        self.assertEqual(str(subscription.last_report_id), result["report_id"])
        self.assertIsNotNone(subscription.next_run_at)
        self.assertTrue(MarketingNotification.objects.filter(source_type="REPORT_GENERATED").exists())
        self.assertTrue(MarketingReportDelivery.objects.filter(report_id=result["report_id"]).exists())
        self.assertGreater(calculate_next_run(subscription), timezone.now())


class MarketingRuntimeFallbackTests(SimpleTestCase):
    def test_runtime_fallback_never_exposes_technical_exception(self):
        answer = get_marketing_runtime_fallback(
            "升级会不会有额外费用？",
            {"marketing_task_id": "task-1", "marketing_customer_id": "customer-1"},
        )

        self.assertIn("当地实际可办理", answer)
        self.assertNotIn("Exception", answer)
        self.assertNotIn("Expecting value", answer)

    @patch("application.chat_pipeline.step.search_dataset_step.impl.base_search_dataset_step.ModelManage.get_model")
    @patch(
        "application.chat_pipeline.step.search_dataset_step.impl.base_search_dataset_step.get_model_default_params",
        return_value={},
    )
    @patch("application.chat_pipeline.step.search_dataset_step.impl.base_search_dataset_step.get_model_by_id")
    @patch(
        "application.chat_pipeline.step.search_dataset_step.impl.base_search_dataset_step.get_embedding_id",
        return_value="embedding-1",
    )
    def test_marketing_retrieval_failure_degrades_to_empty_references(
        self,
        _get_embedding_id,
        get_model_by_id,
        _get_default_params,
        get_cached_model,
    ):
        get_model_by_id.return_value = SimpleNamespace(model_type="EMBEDDING", name="local embedding")
        embedding = Mock()
        embedding.embed_query.side_effect = ValueError("invalid local embedding response")
        get_cached_model.return_value = embedding
        manage = SimpleNamespace(context={"source": {"marketing_task_id": "task-1"}})
        step = BaseSearchDatasetStep()

        result = step.execute(
            "套餐怎么升级？",
            ["knowledge-1"],
            [],
            [],
            5,
            0.6,
            search_mode="embedding",
            workspace_id="workspace-1",
            manage=manage,
        )

        self.assertEqual(result, [])
        self.assertTrue(manage.context["marketing_retrieval_degraded"])
        self.assertEqual(step.context["retrieval_warning"], "MARKETING_KNOWLEDGE_RETRIEVAL_UNAVAILABLE")


class MarketingScenarioServiceTests(TestCase):
    def setUp(self):
        self.workspace_id = "workspace-a"
        self.other_workspace_id = "workspace-b"
        self.user = User.objects.create(
            username="marketing_tester",
            nick_name="marketing_tester",
            password="x",
            role="USER",
        )
        self.application = Application.objects.create(
            workspace_id=self.workspace_id,
            name="存量客户套餐升级助手",
            user=self.user,
        )
        self.knowledge = Knowledge.objects.create(
            workspace_id=self.workspace_id,
            name="套餐政策知识库",
            desc="套餐政策",
            user=self.user,
        )
        self.other_application = Application.objects.create(
            workspace_id=self.other_workspace_id,
            name="其它空间智能体",
            user=self.user,
        )

    def test_create_list_update_delete_scenario(self):
        service = MarketingScenarioService(self.workspace_id, self.user)

        created = service.create(
            {
                "name": "199升299营销",
                "target_audience": "199套餐高价值用户",
                "objective": "推荐升级到299融合套餐",
                "channel": "ONLINE",
                "status": "ACTIVE",
                "application_id": str(self.application.id),
                "knowledge_id": str(self.knowledge.id),
                "description": "优先使用合规话术",
            }
        )

        self.assertEqual(created["application_name"], self.application.name)
        self.assertEqual(created["knowledge_name"], self.knowledge.name)

        page = service.page(1, 20, {"name": "199"})
        self.assertEqual(page["total"], 1)

        updated = service.update(created["id"], {**created, "status": "PAUSED", "objective": "复盘话术效果"})
        self.assertEqual(updated["status"], "PAUSED")
        self.assertEqual(updated["objective"], "复盘话术效果")

        self.assertTrue(service.delete(created["id"]))
        self.assertEqual(service.page(1, 20, {})["total"], 0)

    def test_reject_related_resource_from_another_workspace(self):
        service = MarketingScenarioService(self.workspace_id, self.user)

        with self.assertRaises(AppApiException) as ctx:
            service.create(
                {
                    "name": "跨空间关联",
                    "application_id": str(self.other_application.id),
                    "status": "DRAFT",
                }
            )

        self.assertEqual(ctx.exception.code, 400)


class MarketingCustomerInsightTests(TestCase):
    def setUp(self):
        self.workspace_id = "marketing-insight-workspace"
        self.user = User.objects.create(
            username="marketing_insight_tester",
            nick_name="营销测试员",
            password="x",
            role="USER",
        )
        self.application = Application.objects.create(
            workspace_id=self.workspace_id,
            name="套餐升级智能体",
            user=self.user,
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id,
            name="套餐升级场景",
            application=self.application,
            user=self.user,
        )
        self.task = MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="套餐升级任务",
            scenario=self.scenario,
            application=self.application,
            user=self.user,
        )
        self.customer = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="张女士",
            phone="13800000000",
            remark="原始人工备注",
        )

    def test_message_insight_is_deterministic_and_preserves_manual_state(self):
        self.assertEqual(analyze_customer_message("我先考虑一下，费用有点贵")["status"], "CONSULTED")
        source = {
            "marketing_task_id": str(self.task.id),
            "marketing_customer_id": str(self.customer.id),
        }
        insight = update_customer_conversation_insight(source, "我先考虑一下，费用有点贵", "chat-1", "record-1")
        self.customer.refresh_from_db()
        self.assertEqual(insight["objection_tags"], ["费用"])
        self.assertEqual(self.customer.status, MarketingCustomerStatus.CONSULTED)
        self.assertEqual(self.customer.intent, "待跟进")
        self.assertEqual(self.customer.remark, "原始人工备注")

        update_customer_conversation_insight(source, "我要办理升级", "chat-1", "record-2")
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.status, MarketingCustomerStatus.INTENT)

        update_customer_conversation_insight(source, "我要投诉扣费问题", "chat-1", "record-3")
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.status, MarketingCustomerStatus.MANUAL)
        update_customer_conversation_insight(source, "我不想办了", "chat-1", "record-4")
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.status, MarketingCustomerStatus.MANUAL)
        self.assertEqual(len(self.customer.config["marketing_insight_history"]), 4)

    def test_manual_follow_up_is_recorded_and_effect_is_aggregated(self):
        customer_service = MarketingCustomerService(self.workspace_id, self.user)
        saved = customer_service.follow_up(
            str(self.task.id),
            str(self.customer.id),
            {"status": "INTENT", "intent": "费用确认后办理", "note": "约定明天回访确认当地方案"},
        )
        self.assertEqual(saved["status"], MarketingCustomerStatus.INTENT)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.remark, "原始人工备注")
        self.assertEqual(len(self.customer.config["marketing_follow_up_history"]), 1)

        effect = MarketingTaskService(self.workspace_id, self.user).effect(str(self.task.id))
        self.assertEqual(effect["metrics"]["intent_count"], 1)
        self.assertEqual(effect["metrics"]["intent_rate"], 100.0)
        self.assertEqual(effect["metrics"]["follow_up_record_count"], 1)


class MarketingAnswerQualityTests(TestCase):
    def setUp(self):
        self.workspace_id = "marketing-quality-workspace"
        self.user = User.objects.create(
            username="marketing_quality_tester",
            nick_name="质检测试员",
            password="x",
            role="USER",
        )
        self.application = Application.objects.create(
            workspace_id=self.workspace_id,
            name="营销质检智能体",
            user=self.user,
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id,
            name="营销质检场景",
            application=self.application,
            user=self.user,
        )
        self.task = MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="营销质检任务",
            scenario=self.scenario,
            application=self.application,
            user=self.user,
        )
        self.customer = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="张女士",
            phone="13800000000",
            province="广东",
            current_plan="199元套餐",
        )
        self.source = {
            "marketing_task_id": str(self.task.id),
            "marketing_customer_id": str(self.customer.id),
            "marketing_customer_plan": self.customer.current_plan,
            "marketing_customer_province": self.customer.province,
            "marketing_recommendation": {
                "status": "MATCHED",
                "offer_id": "offer-199",
                "offer_name": "199元套餐同价扩容活动",
                "benefit": "可优先核实同价扩容权益。",
                "fee_note": "费用以当地实际可办理方案为准。",
                "contract_note": "办理前确认优惠期限和合约期。",
            },
        }
        self.chat = Chat.objects.create(
            application=self.application,
            abstract="套餐升级",
            source=self.source,
        )

    def _seed_eligible_rule(self, risk_type="INTERNAL_EXPOSURE"):
        analysis_chat = Chat.objects.create(
            application=self.application,
            abstract=f"{risk_type} 准入样本",
            source={},
        )
        reviews = []
        for index in range(20):
            record = ChatRecord.objects.create(
                chat=analysis_chat,
                problem_text=f"{risk_type} 准入问题 {index}",
                answer_text=f"{risk_type} 准入回答 {index}",
                index=index,
                source={},
            )
            reviews.append(
                MarketingAnswerQualityReview.objects.create(
                    workspace_id=self.workspace_id,
                    task=self.task,
                    customer=self.customer,
                    record=record,
                    status=MarketingAnswerQualityStatus.HIGH_RISK,
                    risk_types=[risk_type],
                    manual_status=(
                        MarketingAnswerManualReviewStatus.CONFIRMED_PASS
                        if index == 19
                        else MarketingAnswerManualReviewStatus.CONFIRMED_ISSUE
                    ),
                )
            )
        return reviews

    def test_quality_result_is_saved_and_aggregated_in_task_effect(self):
        normal = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="我的199元套餐流量不够用，有什么活动？",
            answer_text=(
                "当前可核实199元套餐同价扩容活动及同价扩容权益。"
                "费用以当地实际可办理方案为准，办理前确认优惠期限和合约期。"
            ),
            index=0,
            source=self.source,
        )
        risky = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="那具体怎么办理？",
            answer_text="建议坐席话术：先安装新版本MySQL，再执行数据库升级。",
            index=1,
            source=self.source,
        )

        normal_review = MarketingAnswerQualityReview.objects.get(record=normal)
        risky_review = MarketingAnswerQualityReview.objects.get(record=risky)
        self.assertEqual(normal_review.status, MarketingAnswerQualityStatus.PASS)
        self.assertEqual(risky_review.status, MarketingAnswerQualityStatus.HIGH_RISK)
        self.assertIn("INTERNAL_EXPOSURE", risky_review.risk_types)
        self.assertIn("OFF_TOPIC", risky_review.risk_types)

        effect = MarketingTaskService(self.workspace_id, self.user).effect(str(self.task.id))
        self.assertEqual(effect["metrics"]["quality_pass_count"], 1)
        self.assertEqual(effect["metrics"]["quality_high_risk_count"], 1)
        self.assertEqual(effect["metrics"]["quality_pass_rate"], 50.0)
        records = {row["id"]: row for row in effect["records"]["records"]}
        self.assertEqual(records[str(risky.id)]["quality_review"]["status"], "HIGH_RISK")
        self.assertTrue(records[str(risky.id)]["quality_review"]["reasons"])

    def test_unverified_price_and_commitment_are_high_risk(self):
        result = evaluate_marketing_answer(
            "这个套餐多少钱？",
            "每月299元，保证办理并且立即生效。",
            source={"marketing_task_id": str(self.task.id)},
        )

        self.assertEqual(result["status"], MarketingAnswerQualityStatus.HIGH_RISK)
        self.assertIn("UNVERIFIED_PRICE", result["risk_types"])
        self.assertIn("UNVERIFIED_COMMITMENT", result["risk_types"])
        self.assertEqual(result["evidence"]["unverified_amounts"], ["299"])

    def test_repeated_answer_requires_review(self):
        answer = "当前可优先核实199元套餐同价扩容活动和同价扩容权益，办理前会确认具体条件。"
        result = evaluate_marketing_answer(
            "那贵不贵？",
            answer,
            source=self.source,
            previous_answer=answer,
        )

        self.assertEqual(result["status"], MarketingAnswerQualityStatus.REVIEW)
        self.assertIn("REPEATED_ANSWER", result["risk_types"])

    def test_manual_review_revision_preserves_original_answer_and_audit_history(self):
        original_answer = "建议坐席话术：向客户承诺299元套餐立即生效。"
        record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="这个套餐升级怎么收费？",
            answer_text=original_answer,
            index=2,
            source=self.source,
        )
        review = MarketingAnswerQualityReview.objects.get(record=record)
        service = MarketingQualityReviewService(self.workspace_id, self.user)

        page = service.page(1, 20, {"manual_status": "PENDING"})
        self.assertEqual(page["total"], 1)
        self.assertEqual(page["summary"]["pending_count"], 1)
        self.assertEqual(page["records"][0]["original_answer"], original_answer)

        revised_answer = "套餐费用以广东当地实际可办理方案为准，办理前会向您确认月费和合约期。"
        handled = service.handle(
            str(review.id),
            {
                "action": MarketingAnswerManualReviewAction.REVISE,
                "note": "移除内部话术和未核实资费承诺",
                "revised_answer": revised_answer,
            },
        )
        record.refresh_from_db()
        review.refresh_from_db()
        self.assertEqual(record.answer_text, original_answer)
        self.assertEqual(review.manual_status, MarketingAnswerManualReviewStatus.REVISED)
        self.assertEqual(review.revised_answer, revised_answer)
        self.assertEqual(review.manual_revision, 1)
        action = MarketingAnswerQualityReviewActionRecord.objects.get(review=review)
        self.assertEqual(action.original_answer, original_answer)
        self.assertEqual(action.revised_answer, revised_answer)
        self.assertEqual(handled["manual_actions"][0]["revision"], 1)

        effect = MarketingTaskService(self.workspace_id, self.user).effect(str(self.task.id))
        self.assertEqual(effect["metrics"]["quality_manual_pending_count"], 0)
        self.assertEqual(effect["metrics"]["quality_manual_reviewed_count"], 1)
        self.assertEqual(effect["metrics"]["quality_manual_revised_count"], 1)

    def test_manual_review_rejects_cross_workspace_record(self):
        record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="这个套餐怎么升级？",
            answer_text="建议坐席话术：直接承诺办理。",
            index=3,
            source=self.source,
        )
        review = MarketingAnswerQualityReview.objects.get(record=record)

        with self.assertRaises(AppApiException) as ctx:
            MarketingQualityReviewService("another-workspace", self.user).handle(
                str(review.id),
                {"action": MarketingAnswerManualReviewAction.CONFIRM_PASS},
            )

        self.assertEqual(ctx.exception.code, 404)

    def test_manual_review_validation_returns_business_400_without_mutation(self):
        record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="套餐升级是否收费？",
            answer_text="建议坐席话术：直接承诺免费办理。",
            index=4,
            source=self.source,
        )
        review = MarketingAnswerQualityReview.objects.get(record=record)

        with self.assertRaises(AppApiException) as ctx:
            MarketingQualityReviewService(self.workspace_id, self.user).handle(
                str(review.id),
                {"action": MarketingAnswerManualReviewAction.CONFIRM_ISSUE, "note": ""},
            )

        review.refresh_from_db()
        self.assertEqual(ctx.exception.code, 400)
        self.assertEqual(review.manual_status, MarketingAnswerManualReviewStatus.PENDING)
        self.assertEqual(review.manual_revision, 0)

    def test_shadow_preflight_preserves_generated_and_final_answers(self):
        generated_answer = "这个活动每月299元并且立即生效。"
        final_answer = "费用以广东当地实际可办理方案为准，办理前会确认月费变化和合约期。"
        snapshot = prepare_marketing_preflight(
            "升级需要多少钱？",
            generated_answer,
            final_answer,
            self.source,
            chat_id=self.chat.id,
            paragraph_count=1,
        )
        record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="升级需要多少钱？",
            answer_text=final_answer,
            index=5,
            source=self.source,
        )
        preflight = save_marketing_preflight(record.id, snapshot)

        self.assertIsNotNone(preflight)
        self.assertEqual(preflight.policy_mode, MarketingQualityPolicyMode.SHADOW)
        self.assertEqual(preflight.action, MarketingPreflightAction.SAFEGUARD_REWRITTEN)
        self.assertEqual(preflight.generated_answer, generated_answer)
        self.assertEqual(preflight.final_answer, final_answer)
        self.assertEqual(preflight.evidence["knowledge_reference_count"], 1)
        self.assertEqual(record.answer_text, final_answer)

        page = MarketingQualityPolicyService(self.workspace_id, self.user).page(1, 20, {})
        self.assertEqual(page["total"], 1)
        self.assertEqual(page["summary"]["safeguard_rewritten_count"], 1)
        self.assertEqual(page["records"][0]["generated_answer"], generated_answer)

    def test_shadow_preflight_logs_high_risk_without_rewriting(self):
        unsafe_answer = "建议坐席话术：直接承诺299元套餐立即生效。"
        snapshot = prepare_marketing_preflight(
            "升级需要多少钱？",
            unsafe_answer,
            unsafe_answer,
            self.source,
            chat_id=self.chat.id,
        )
        record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="升级需要多少钱？",
            answer_text=unsafe_answer,
            index=6,
            source=self.source,
        )
        preflight = save_marketing_preflight(record.id, snapshot)

        self.assertEqual(preflight.status, MarketingAnswerQualityStatus.HIGH_RISK)
        self.assertEqual(preflight.action, MarketingPreflightAction.OBSERVED)
        self.assertIn("INTERNAL_EXPOSURE", preflight.risk_types)
        self.assertEqual(record.answer_text, unsafe_answer)

    def test_quality_policy_can_disable_shadow_preflight(self):
        service = MarketingQualityPolicyService(self.workspace_id, self.user)
        updated = service.update({"mode": MarketingQualityPolicyMode.OFF})
        snapshot = prepare_marketing_preflight(
            "套餐怎么升级？",
            "建议坐席话术：直接办理。",
            "建议坐席话术：直接办理。",
            self.source,
            chat_id=self.chat.id,
        )

        self.assertEqual(updated["mode"], MarketingQualityPolicyMode.OFF)
        self.assertIsNone(snapshot)
        self.assertEqual(MarketingAnswerPreflightRecord.objects.count(), 0)
        self.assertEqual(MarketingQualityPolicy.objects.get(workspace_id=self.workspace_id).updated_by, self.user)

    def test_quality_policy_rejects_interception_without_eligible_rule(self):
        with self.assertRaises(AppApiException) as ctx:
            MarketingQualityPolicyService(self.workspace_id, self.user).update(
                {
                    "mode": MarketingQualityPolicyMode.INTERCEPT,
                    "intercept_rule_types": ["INTERNAL_EXPOSURE"],
                }
            )

        self.assertEqual(ctx.exception.code, 400)
        self.assertFalse(MarketingQualityPolicy.objects.filter(workspace_id=self.workspace_id).exists())

    def test_eligible_rule_can_enable_interception_and_buffer_delivery(self):
        self._seed_eligible_rule()

        updated = MarketingQualityPolicyService(self.workspace_id, self.user).update(
            {
                "mode": MarketingQualityPolicyMode.INTERCEPT,
                "intercept_rule_types": ["INTERNAL_EXPOSURE"],
            }
        )

        self.assertEqual(updated["mode"], MarketingQualityPolicyMode.INTERCEPT)
        self.assertEqual(updated["intercept_rule_types"], ["INTERNAL_EXPOSURE"])
        self.assertTrue(updated["interception_available"])
        self.assertTrue(should_buffer_marketing_preflight(self.source))

    def test_interception_replaces_only_configured_risk_and_preserves_audit_answers(self):
        self._seed_eligible_rule()
        MarketingQualityPolicyService(self.workspace_id, self.user).update(
            {
                "mode": MarketingQualityPolicyMode.INTERCEPT,
                "intercept_rule_types": ["INTERNAL_EXPOSURE"],
            }
        )
        unsafe_answer = "建议坐席话术：直接承诺299元套餐立即生效。"

        snapshot = prepare_marketing_preflight(
            "升级需要多少钱？",
            unsafe_answer,
            unsafe_answer,
            self.source,
            chat_id=self.chat.id,
        )
        record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="升级需要多少钱？",
            answer_text=snapshot["final_answer"],
            index=21,
            source=self.source,
        )
        preflight = save_marketing_preflight(record.id, snapshot)

        self.assertEqual(snapshot["action"], MarketingPreflightAction.INTERCEPTED)
        self.assertEqual(snapshot["generated_answer"], unsafe_answer)
        self.assertEqual(snapshot["final_answer"], INTERCEPTED_CUSTOMER_ANSWER)
        self.assertEqual(snapshot["evidence"]["intercepted_risk_types"], ["INTERNAL_EXPOSURE"])
        self.assertEqual(preflight.generated_answer, unsafe_answer)
        self.assertEqual(preflight.final_answer, INTERCEPTED_CUSTOMER_ANSWER)
        self.assertEqual(record.answer_text, INTERCEPTED_CUSTOMER_ANSWER)

    def test_interception_does_not_replace_unconfigured_risk(self):
        self._seed_eligible_rule()
        MarketingQualityPolicyService(self.workspace_id, self.user).update(
            {
                "mode": MarketingQualityPolicyMode.INTERCEPT,
                "intercept_rule_types": ["INTERNAL_EXPOSURE"],
            }
        )
        off_topic_answer = "请先安装最新版 MySQL，然后执行数据库升级。"

        snapshot = prepare_marketing_preflight(
            "套餐怎么升级？",
            off_topic_answer,
            off_topic_answer,
            self.source,
            chat_id=self.chat.id,
        )

        self.assertEqual(snapshot["action"], MarketingPreflightAction.OBSERVED)
        self.assertIn("OFF_TOPIC", snapshot["risk_types"])
        self.assertNotIn("INTERNAL_EXPOSURE", snapshot["risk_types"])
        self.assertEqual(snapshot["final_answer"], off_topic_answer)

    @patch("marketing.preflight._evaluate_with_timeout", side_effect=FutureTimeoutError())
    def test_interception_timeout_bypasses_without_changing_customer_answer(self, _evaluate):
        self._seed_eligible_rule()
        MarketingQualityPolicyService(self.workspace_id, self.user).update(
            {
                "mode": MarketingQualityPolicyMode.INTERCEPT,
                "intercept_rule_types": ["INTERNAL_EXPOSURE"],
            }
        )
        original_answer = "费用以当地实际可办理方案为准。"

        snapshot = prepare_marketing_preflight(
            "升级是否收费？",
            original_answer,
            original_answer,
            self.source,
            chat_id=self.chat.id,
        )

        self.assertEqual(snapshot["action"], MarketingPreflightAction.BYPASSED)
        self.assertEqual(snapshot["evidence"]["bypass_reason"], "TIMEOUT")
        self.assertEqual(snapshot["final_answer"], original_answer)

    @patch("marketing.preflight._evaluate_with_timeout", side_effect=RuntimeError("quality unavailable"))
    def test_interception_error_bypasses_without_changing_customer_answer(self, _evaluate):
        self._seed_eligible_rule()
        MarketingQualityPolicyService(self.workspace_id, self.user).update(
            {
                "mode": MarketingQualityPolicyMode.INTERCEPT,
                "intercept_rule_types": ["INTERNAL_EXPOSURE"],
            }
        )
        original_answer = "费用以当地实际可办理方案为准。"

        snapshot = prepare_marketing_preflight(
            "升级是否收费？",
            original_answer,
            original_answer,
            self.source,
            chat_id=self.chat.id,
        )

        self.assertEqual(snapshot["action"], MarketingPreflightAction.BYPASSED)
        self.assertEqual(snapshot["evidence"]["bypass_reason"], "ERROR")
        self.assertEqual(snapshot["final_answer"], original_answer)

    def test_manual_false_positive_degrades_rule_and_automatically_falls_back_to_shadow(self):
        reviews = self._seed_eligible_rule()
        MarketingQualityPolicyService(self.workspace_id, self.user).update(
            {
                "mode": MarketingQualityPolicyMode.INTERCEPT,
                "intercept_rule_types": ["INTERNAL_EXPOSURE"],
            }
        )

        MarketingQualityReviewService(self.workspace_id, self.user).handle(
            str(reviews[0].id),
            {"action": MarketingAnswerManualReviewAction.CONFIRM_PASS},
        )

        policy = MarketingQualityPolicy.objects.get(workspace_id=self.workspace_id)
        self.assertEqual(policy.mode, MarketingQualityPolicyMode.SHADOW)
        self.assertEqual(policy.intercept_rule_types, [])
        self.assertEqual(policy.updated_by_name, "系统自动回退")
        self.assertFalse(should_buffer_marketing_preflight(self.source))

    def test_quality_rule_effect_calculates_manual_accuracy_and_gate(self):
        analysis_chat = Chat.objects.create(
            application=self.application,
            abstract="规则效果样本",
            source={},
        )
        manual_statuses = (
            [MarketingAnswerManualReviewStatus.REVISED] * 5
            + [MarketingAnswerManualReviewStatus.CONFIRMED_ISSUE] * 14
            + [MarketingAnswerManualReviewStatus.CONFIRMED_PASS]
        )
        records = []
        for index, manual_status in enumerate(manual_statuses):
            record = ChatRecord.objects.create(
                chat=analysis_chat,
                problem_text=f"规则效果问题 {index}",
                answer_text=f"规则效果回答 {index}",
                index=index,
                source={},
            )
            records.append(record)
            MarketingAnswerQualityReview.objects.create(
                workspace_id=self.workspace_id,
                task=self.task,
                customer=self.customer,
                record=record,
                status=MarketingAnswerQualityStatus.HIGH_RISK,
                risk_types=["INTERNAL_EXPOSURE"],
                manual_status=manual_status,
            )
        pending_record = ChatRecord.objects.create(
            chat=analysis_chat,
            problem_text="待处理规则效果问题",
            answer_text="待处理规则效果回答",
            index=20,
            source={},
        )
        MarketingAnswerQualityReview.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            customer=self.customer,
            record=pending_record,
            status=MarketingAnswerQualityStatus.HIGH_RISK,
            risk_types=["INTERNAL_EXPOSURE"],
        )
        MarketingAnswerPreflightRecord.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            customer=self.customer,
            record=records[0],
            policy_mode=MarketingQualityPolicyMode.SHADOW,
            action=MarketingPreflightAction.OBSERVED,
            status=MarketingAnswerQualityStatus.HIGH_RISK,
            risk_types=["INTERNAL_EXPOSURE"],
        )

        result = MarketingQualityRuleEffectService(self.workspace_id, self.user).analyze(
            {"risk_type": "INTERNAL_EXPOSURE"}
        )

        self.assertEqual(result["summary"]["hit_record_count"], 21)
        self.assertEqual(result["summary"]["manually_reviewed_record_count"], 20)
        self.assertEqual(result["summary"]["pending_record_count"], 1)
        self.assertEqual(result["summary"]["eligible_rule_count"], 1)
        rule = result["rules"][0]
        self.assertEqual(rule["hit_count"], 21)
        self.assertEqual(rule["shadow_hit_count"], 1)
        self.assertEqual(rule["confirmed_issue_count"], 19)
        self.assertEqual(rule["false_positive_count"], 1)
        self.assertEqual(rule["revised_count"], 5)
        self.assertEqual(rule["issue_confirmation_rate"], 95.0)
        self.assertEqual(rule["false_positive_rate"], 5.0)
        self.assertTrue(rule["eligible"])
        self.assertEqual(rule["gate_status"], "ELIGIBLE")

    def test_quality_rule_effect_marks_review_only_rule_as_not_eligible(self):
        result = MarketingQualityRuleEffectService(self.workspace_id, self.user).analyze(
            {"risk_type": "REPEATED_ANSWER"}
        )

        self.assertEqual(len(result["rules"]), 1)
        self.assertFalse(result["rules"][0]["eligible"])
        self.assertEqual(result["rules"][0]["gate_status"], "NOT_HIGH_RISK")

    def test_quality_rule_effect_rejects_inverted_date_range(self):
        with self.assertRaises(AppApiException) as ctx:
            MarketingQualityRuleEffectService(self.workspace_id, self.user).analyze(
                {"date_from": "2026-09-08", "date_to": "2026-09-01"}
            )

        self.assertEqual(ctx.exception.code, 400)


class MarketingOperationsTests(TestCase):
    def setUp(self):
        self.workspace_id = "marketing-operations-workspace"
        self.user = User.objects.create(
            username="marketing_operations_tester",
            nick_name="运营监控测试员",
            password="x",
            role="USER",
        )
        self.application = Application.objects.create(
            workspace_id=self.workspace_id,
            name="运营监控智能体",
            user=self.user,
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id,
            name="运营监控场景",
            application=self.application,
            user=self.user,
        )
        self.task = MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="运营监控任务",
            scenario=self.scenario,
            application=self.application,
            user=self.user,
            status=MarketingTaskStatus.RUNNING,
        )
        self.customer = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="王女士",
            phone="13700000000",
            status=MarketingCustomerStatus.CONVERTED,
        )
        self.source = {
            "marketing_task_id": str(self.task.id),
            "marketing_customer_id": str(self.customer.id),
        }
        self.chat = Chat.objects.create(
            application=self.application,
            abstract="运营监控对话",
            source=self.source,
        )

    def _create_record(self, index, run_time, status, risk_types):
        record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text=f"监控问题 {index}",
            answer_text="费用以当地实际可办理方案为准，办理前会确认优惠期限和合约期。",
            index=index,
            run_time=run_time,
            message_tokens=10,
            answer_tokens=20,
            source=self.source,
        )
        review = MarketingAnswerQualityReview.objects.get(record=record)
        review.status = status
        review.risk_types = risk_types
        review.save(update_fields=["status", "risk_types"])
        return record

    def test_overview_aggregates_health_trend_and_task_drilldown(self):
        healthy_record = self._create_record(0, 2, MarketingAnswerQualityStatus.PASS, [])
        error_record = self._create_record(
            1,
            12,
            MarketingAnswerQualityStatus.HIGH_RISK,
            ["RUNTIME_ERROR"],
        )
        MarketingAnswerPreflightRecord.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            customer=self.customer,
            record=error_record,
            policy_mode=MarketingQualityPolicyMode.INTERCEPT,
            action=MarketingPreflightAction.BYPASSED,
            status=MarketingAnswerQualityStatus.REVIEW,
            evidence={"bypass_reason": "TIMEOUT"},
        )
        MarketingConversionRecord.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            customer=self.customer,
            user=self.user,
            outcome="CONVERTED",
            note="客户确认办理",
        )
        MarketingTouchBatch.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            user=self.user,
            name="逾期测试批次",
            deadline=timezone.now() - timedelta(hours=1),
            status=MarketingTouchBatchStatus.RUNNING,
            approval_status=MarketingTouchBatchApprovalStatus.APPROVED,
        )

        overview = MarketingOperationsService(self.workspace_id, self.user).overview({"days": 7})

        self.assertEqual(overview["health"]["status"], "CRITICAL")
        self.assertEqual(overview["metrics"]["conversation_count"], 2)
        self.assertEqual(overview["metrics"]["conversation_success_rate"], 50.0)
        self.assertEqual(overview["metrics"]["avg_response_seconds"], 7.0)
        self.assertEqual(overview["metrics"]["slow_conversation_count"], 1)
        self.assertEqual(overview["metrics"]["total_tokens"], 60)
        self.assertEqual(overview["metrics"]["quality_pass_rate"], 50.0)
        self.assertEqual(overview["metrics"]["runtime_error_count"], 1)
        self.assertEqual(overview["metrics"]["bypassed_count"], 1)
        self.assertEqual(overview["metrics"]["overdue_batch_count"], 1)
        self.assertEqual(overview["metrics"]["conversion_rate"], 100.0)
        self.assertEqual(overview["trend"][-1]["conversation_count"], 2)
        self.assertEqual(overview["trend"][-1]["runtime_error_count"], 1)
        self.assertEqual(overview["trend"][-1]["bypassed_count"], 1)
        self.assertEqual(overview["trend"][-1]["conversion_count"], 1)
        self.assertEqual(overview["task_health"][0]["health_status"], "ERROR")
        self.assertEqual(overview["task_health"][0]["conversation_count"], 2)
        self.assertEqual(overview["task_health"][0]["last_activity_at"], error_record.create_time)
        self.assertEqual(healthy_record.chat_id, self.chat.id)
        self.assertEqual(
            [item["type"] for item in overview["issues"][:3]],
            ["RUNTIME_ERROR", "QUALITY_BYPASS", "OVERDUE_BATCH"],
        )

    def test_overview_is_workspace_scoped_and_validates_filters(self):
        self._create_record(0, 1, MarketingAnswerQualityStatus.PASS, [])
        other_scenario = MarketingScenario.objects.create(workspace_id="other-workspace", name="其它场景")
        MarketingTask.objects.create(
            workspace_id="other-workspace",
            name="其它任务",
            scenario=other_scenario,
            status=MarketingTaskStatus.RUNNING,
        )

        overview = MarketingOperationsService(self.workspace_id, self.user).overview(
            {"days": 30, "task_id": str(self.task.id)}
        )

        self.assertEqual(overview["metrics"]["task_count"], 1)
        self.assertEqual(overview["metrics"]["conversation_count"], 1)
        self.assertEqual(overview["range"]["task_id"], str(self.task.id))
        with self.assertRaises(AppApiException) as invalid_days:
            MarketingOperationsService(self.workspace_id, self.user).overview({"days": 14})
        self.assertEqual(invalid_days.exception.code, 400)
        with self.assertRaises(AppApiException) as other_task:
            MarketingOperationsService(self.workspace_id, self.user).overview(
                {"task_id": str(MarketingTask.objects.get(workspace_id="other-workspace").id)}
            )
        self.assertEqual(other_task.exception.code, 404)


class MarketingAlertTests(TestCase):
    def setUp(self):
        self.workspace_id = "marketing-alert-workspace"
        self.user = User.objects.create(
            username="marketing_alert_tester",
            nick_name="告警测试员",
            password="x",
            role="USER",
        )
        self.application = Application.objects.create(
            workspace_id=self.workspace_id,
            name="告警测试智能体",
            user=self.user,
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id,
            name="告警测试场景",
            application=self.application,
            user=self.user,
        )
        self.task = MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="告警测试任务",
            scenario=self.scenario,
            application=self.application,
            user=self.user,
            status=MarketingTaskStatus.RUNNING,
        )
        self.customer = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="告警客户",
            phone="13600000000",
        )
        self.source = {
            "marketing_task_id": str(self.task.id),
            "marketing_customer_id": str(self.customer.id),
        }
        self.chat = Chat.objects.create(application=self.application, abstract="告警测试对话", source=self.source)
        self.record = ChatRecord.objects.create(
            chat=self.chat,
            problem_text="告警测试问题",
            answer_text="系统异常",
            index=0,
            run_time=12,
            source=self.source,
        )
        self.review = MarketingAnswerQualityReview.objects.get(record=self.record)
        self.review.status = MarketingAnswerQualityStatus.HIGH_RISK
        self.review.risk_types = ["RUNTIME_ERROR"]
        self.review.save(update_fields=["status", "risk_types"])
        MarketingAnswerPreflightRecord.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            customer=self.customer,
            record=self.record,
            policy_mode=MarketingQualityPolicyMode.INTERCEPT,
            action=MarketingPreflightAction.BYPASSED,
            status=MarketingAnswerQualityStatus.HIGH_RISK,
        )
        MarketingTouchBatch.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            user=self.user,
            name="告警逾期批次",
            deadline=timezone.now() - timedelta(hours=1),
            status=MarketingTouchBatchStatus.RUNNING,
            approval_status=MarketingTouchBatchApprovalStatus.APPROVED,
        )
        self.service = MarketingAlertService(self.workspace_id, self.user)
        self.notification_service = MarketingNotificationService(self.workspace_id, self.user)

    def test_sync_creates_rules_and_deduplicated_alerts(self):
        first_summary = self.service.sync({"days": 7})

        self.assertEqual(len(self.service.rules()), 5)
        self.assertEqual(first_summary["total_count"], 5)
        self.assertEqual(first_summary["critical_count"], 2)
        self.assertEqual(MarketingAlertActionRecord.objects.count(), 5)

        second_summary = self.service.sync({"days": 7})

        self.assertEqual(second_summary["total_count"], 5)
        self.assertEqual(MarketingAlert.objects.count(), 5)
        self.assertEqual(MarketingAlertActionRecord.objects.count(), 5)

    def test_alert_can_be_claimed_started_closed_and_manually_reopened(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")

        claimed = self.service.handle(str(alert.id), {"action": "CLAIM"})
        self.assertEqual(claimed["status"], MarketingAlertStatus.CLAIMED)
        self.assertEqual(claimed["assignee_name"], self.user.nick_name)

        processing = self.service.handle(str(alert.id), {"action": "START", "note": "开始排查模型调用"})
        self.assertEqual(processing["status"], MarketingAlertStatus.PROCESSING)

        closed = self.service.handle(
            str(alert.id),
            {
                "action": "CLOSE",
                "note": "已核实并恢复",
                "root_cause_category": "MODEL",
                "resolution_type": "CONFIG_CHANGE",
                "impact_scope": "营销对话响应",
            },
        )
        self.assertEqual(closed["status"], MarketingAlertStatus.CLOSED)
        self.assertEqual(closed["close_note"], "已核实并恢复")
        self.assertEqual(closed["root_cause_category"], "MODEL")
        self.assertEqual(closed["resolution_type"], "CONFIG_CHANGE")
        self.assertEqual(closed["actions"][0]["snapshot"]["impact_scope"], "营销对话响应")

        self.service.sync({"days": 7})
        alert.refresh_from_db()
        self.assertEqual(alert.status, MarketingAlertStatus.CLOSED)

        reopened = self.service.handle(str(alert.id), {"action": "REOPEN", "note": "继续观察"})
        self.assertEqual(reopened["status"], MarketingAlertStatus.OPEN)
        self.assertEqual(
            [item["action"] for item in reversed(reopened["actions"])],
            [
                MarketingAlertAction.CREATED,
                MarketingAlertAction.CLAIMED,
                MarketingAlertAction.STARTED,
                MarketingAlertAction.CLOSED,
                MarketingAlertAction.REOPENED,
            ],
        )

    def test_closed_alert_reopens_only_after_source_recovers_and_recurs(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        self.service.handle(
            str(alert.id),
            {
                "action": "CLOSE",
                "note": "阶段性关闭",
                "root_cause_category": "MODEL",
                "resolution_type": "RETRY_RECOVERY",
            },
        )

        self.review.status = MarketingAnswerQualityStatus.PASS
        self.review.risk_types = []
        self.review.save(update_fields=["status", "risk_types"])
        self.service.sync({"days": 7})
        alert.refresh_from_db()
        self.assertFalse(alert.source_active)
        self.assertEqual(alert.status, MarketingAlertStatus.CLOSED)

        self.review.status = MarketingAnswerQualityStatus.HIGH_RISK
        self.review.risk_types = ["RUNTIME_ERROR"]
        self.review.save(update_fields=["status", "risk_types"])
        self.service.sync({"days": 7})
        alert.refresh_from_db()
        self.assertTrue(alert.source_active)
        self.assertEqual(alert.status, MarketingAlertStatus.OPEN)

    def test_rule_threshold_suppresses_alert_and_workspace_scope_is_enforced(self):
        runtime_rule = next(rule for rule in self.service.rules() if rule["alert_type"] == "RUNTIME_ERROR")
        updated_rule = self.service.update_rule(runtime_rule["id"], {"threshold": 2, "severity": "WARNING"})
        self.assertEqual(updated_rule["threshold"], 2)

        self.service.sync({"days": 7})
        self.assertFalse(
            MarketingAlert.objects.filter(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR").exists()
        )
        other_service = MarketingAlertService("another-workspace", self.user)
        existing_alert = MarketingAlert.objects.filter(workspace_id=self.workspace_id).first()
        with self.assertRaises(AppApiException) as missing:
            other_service.handle(str(existing_alert.id), {"action": "CLAIM"})
        self.assertEqual(missing.exception.code, 404)

    def test_close_requires_note_and_invalid_transition_is_rejected(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")

        with self.assertRaises(AppApiException) as missing_note:
            self.service.handle(str(alert.id), {"action": "CLOSE", "note": ""})
        self.assertEqual(missing_note.exception.code, 400)
        self.service.handle(str(alert.id), {"action": "CLAIM"})
        with self.assertRaises(AppApiException) as duplicate_claim:
            self.service.handle(str(alert.id), {"action": "CLAIM"})
        self.assertEqual(duplicate_claim.exception.code, 400)

    def test_sync_emits_deduplicated_in_app_notifications(self):
        first = self.notification_service.sync_and_notify({"days": 7})

        self.assertEqual(first["notification_count"], 5)
        self.assertEqual(self.notification_service.unread_count(), 5)
        self.assertEqual(MarketingNotification.objects.count(), 5)

        second = self.notification_service.sync_and_notify({"days": 7})

        self.assertEqual(second["notification_count"], 0)
        self.assertEqual(MarketingNotification.objects.count(), 5)

    def test_notification_can_be_read_individually_or_in_bulk(self):
        self.notification_service.sync_and_notify({"days": 7})
        notification = MarketingNotification.objects.filter(user=self.user).first()

        read = self.notification_service.mark_read(str(notification.id))
        self.assertEqual(read["status"], MarketingNotificationStatus.READ)
        self.assertIsNotNone(read["read_at"])
        self.assertEqual(self.notification_service.unread_count(), 4)

        result = self.notification_service.mark_all_read()
        self.assertEqual(result["updated_count"], 4)
        self.assertEqual(result["unread_count"], 0)
        self.assertEqual(
            MarketingNotification.objects.filter(status=MarketingNotificationStatus.UNREAD).count(),
            0,
        )

    def test_global_notification_page_supports_severity_filter_and_event_metadata(self):
        self.notification_service.sync_and_notify({"days": 7})

        page = self.notification_service.page(
            1,
            20,
            {"status": "UNREAD", "severity": "CRITICAL", "action": "CREATED"},
        )

        self.assertEqual(page["total"], 2)
        self.assertEqual(page["unread_count"], 5)
        self.assertTrue(all(item["severity"] == "CRITICAL" for item in page["records"]))
        self.assertTrue(all(item["action"] == MarketingAlertAction.CREATED for item in page["records"]))
        self.assertTrue(all(item["alert_type"] for item in page["records"]))

    def test_notification_is_user_and_workspace_scoped(self):
        self.notification_service.sync_and_notify({"days": 7})
        notification = MarketingNotification.objects.filter(user=self.user).first()
        other_user = User.objects.create(
            username="other_notice_user",
            nick_name="其它通知用户",
            password="x",
            role="USER",
        )

        with self.assertRaises(AppApiException) as other_user_missing:
            MarketingNotificationService(self.workspace_id, other_user).mark_read(str(notification.id))
        self.assertEqual(other_user_missing.exception.code, 404)
        with self.assertRaises(AppApiException) as other_workspace_missing:
            MarketingNotificationService("other-workspace", self.user).mark_read(str(notification.id))
        self.assertEqual(other_workspace_missing.exception.code, 404)

    def test_periodic_scan_processes_workspace_and_notifies_task_owner(self):
        result = scan_all_marketing_alerts(7)

        row = next(item for item in result["scanned"] if item["workspace_id"] == self.workspace_id)
        self.assertEqual(row["active_count"], 5)
        self.assertEqual(row["notification_count"], 5)
        self.assertEqual(MarketingNotification.objects.filter(user=self.user).count(), 5)

    def test_sla_deadlines_are_created_and_claim_records_first_response(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        rule = MarketingAlertRule.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")

        self.assertEqual(alert.response_due_at, alert.sla_started_at + timedelta(minutes=rule.response_sla_minutes))
        self.assertEqual(
            alert.resolution_due_at,
            alert.sla_started_at + timedelta(minutes=rule.resolution_sla_minutes),
        )
        self.assertIsNone(alert.first_responded_at)

        self.service.handle(str(alert.id), {"action": "CLAIM"})
        alert.refresh_from_db()
        self.assertIsNotNone(alert.first_responded_at)

    def test_response_overdue_notifies_only_configured_escalation_user_once(self):
        escalation_user = User.objects.create(
            username="marketing_escalation_user",
            nick_name="值班负责人",
            password="x",
            role="USER",
        )
        MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="升级接收人关联任务",
            scenario=self.scenario,
            application=self.application,
            user=escalation_user,
            status=MarketingTaskStatus.RUNNING,
        )
        self.service.ensure_rules()
        rule = MarketingAlertRule.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        self.service.update_rule(
            str(rule.id),
            {
                "response_sla_minutes": 1,
                "resolution_sla_minutes": 60,
                "escalation_user_id": str(escalation_user.id),
            },
        )
        self.notification_service.sync_and_notify({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        alert.response_due_at = timezone.now() - timedelta(minutes=1)
        alert.resolution_due_at = timezone.now() + timedelta(minutes=30)
        alert.save(update_fields=["response_due_at", "resolution_due_at", "update_time"])

        first = self.notification_service.sync_and_notify({"days": 7})
        second = self.notification_service.sync_and_notify({"days": 7})

        self.assertEqual(first["notification_count"], 1)
        self.assertEqual(second["notification_count"], 0)
        action = MarketingAlertActionRecord.objects.get(
            alert=alert,
            action=MarketingAlertAction.RESPONSE_OVERDUE,
        )
        notification = MarketingNotification.objects.get(action_record=action)
        self.assertEqual(notification.user, escalation_user)
        self.assertIn("首次响应超时", notification.title)
        self.assertFalse(MarketingNotification.objects.filter(action_record=action, user=self.user).exists())

    def test_resolution_overdue_escalates_after_response_without_duplicate(self):
        self.notification_service.sync_and_notify({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        now = timezone.now()
        alert.first_responded_at = now - timedelta(minutes=10)
        alert.response_due_at = now - timedelta(minutes=9)
        alert.resolution_due_at = now - timedelta(minutes=1)
        alert.save(
            update_fields=[
                "first_responded_at",
                "response_due_at",
                "resolution_due_at",
                "update_time",
            ]
        )

        first = self.notification_service.sync_and_notify({"days": 7})
        second = self.notification_service.sync_and_notify({"days": 7})

        self.assertEqual(first["notification_count"], 1)
        self.assertEqual(second["notification_count"], 0)
        self.assertFalse(
            MarketingAlertActionRecord.objects.filter(
                alert=alert,
                action=MarketingAlertAction.RESPONSE_OVERDUE,
            ).exists()
        )
        self.assertEqual(
            MarketingAlertActionRecord.objects.filter(
                alert=alert,
                action=MarketingAlertAction.RESOLUTION_OVERDUE,
            ).count(),
            1,
        )

    def test_sla_rule_rejects_invalid_deadline_and_outside_workspace_user(self):
        self.service.ensure_rules()
        rule = MarketingAlertRule.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        outside_user = User.objects.create(
            username="outside_escalation_user",
            nick_name="其他空间用户",
            password="x",
            role="USER",
        )

        with self.assertRaises(AppApiException) as invalid_deadline:
            self.service.update_rule(
                str(rule.id),
                {"response_sla_minutes": 60, "resolution_sla_minutes": 30},
            )
        self.assertEqual(invalid_deadline.exception.code, 400)
        with self.assertRaises(AppApiException) as invalid_user:
            self.service.update_rule(str(rule.id), {"escalation_user_id": str(outside_user.id)})
        self.assertEqual(invalid_user.exception.code, 400)
        self.assertEqual(self.service.members()[0]["id"], str(self.user.id))

    def test_owner_performance_aggregates_workload_duration_and_sla(self):
        self.service.sync({"days": 7})
        runtime_alert = MarketingAlert.objects.get(
            workspace_id=self.workspace_id,
            alert_type="RUNTIME_ERROR",
        )
        slow_alert = MarketingAlert.objects.get(
            workspace_id=self.workspace_id,
            alert_type="SLOW_RESPONSE",
        )
        self.service.handle(str(runtime_alert.id), {"action": "CLAIM"})
        self.service.handle(
            str(runtime_alert.id),
            {
                "action": "CLOSE",
                "note": "已恢复",
                "root_cause_category": "SYSTEM",
                "resolution_type": "RETRY_RECOVERY",
            },
        )

        other_user = User.objects.create(
            username="marketing_alert_owner",
            nick_name="二线负责人",
            password="x",
            role="USER",
        )
        MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="二线负责人任务",
            scenario=self.scenario,
            application=self.application,
            user=other_user,
            status=MarketingTaskStatus.RUNNING,
        )
        MarketingAlertService(self.workspace_id, other_user).handle(str(slow_alert.id), {"action": "CLAIM"})

        now = timezone.now()
        MarketingAlert.objects.filter(id=runtime_alert.id).update(
            sla_started_at=now - timedelta(minutes=60),
            response_due_at=now - timedelta(minutes=30),
            first_responded_at=now - timedelta(minutes=45),
            resolution_due_at=now - timedelta(minutes=20),
            closed_at=now - timedelta(minutes=10),
        )
        MarketingAlert.objects.filter(id=slow_alert.id).update(
            sla_started_at=now - timedelta(minutes=30),
            response_due_at=now - timedelta(minutes=20),
            first_responded_at=now - timedelta(minutes=15),
            resolution_due_at=now + timedelta(minutes=30),
        )

        result = self.service.performance({"days": 7})

        self.assertEqual(result["summary"]["total_count"], 5)
        self.assertEqual(result["summary"]["active_count"], 4)
        self.assertEqual(result["summary"]["unassigned_count"], 3)
        self.assertEqual(result["summary"]["closed_count"], 1)
        self.assertEqual(result["summary"]["response_sla_rate"], 50)
        self.assertEqual(result["summary"]["resolution_sla_rate"], 0)
        self.assertEqual(result["summary"]["sla_attainment_rate"], 33.33)
        self.assertEqual(result["workload"][0]["assignee_name"], "未认领")
        owner = next(item for item in result["workload"] if item["assignee_id"] == str(self.user.id))
        self.assertEqual(owner["avg_response_seconds"], 900)
        self.assertEqual(owner["avg_resolution_seconds"], 3000)
        self.assertEqual(owner["sla_attainment_rate"], 50)

    def test_alert_page_filters_owner_unassigned_and_active_work(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        self.service.handle(str(alert.id), {"action": "CLAIM"})

        assigned = self.service.page(
            1,
            20,
            {"assignee_id": str(self.user.id), "active_only": True},
        )
        unassigned = self.service.page(1, 20, {"unassigned": True, "active_only": True})

        self.assertEqual(assigned["total"], 1)
        self.assertEqual(assigned["records"][0]["assignee_id"], str(self.user.id))
        self.assertEqual(unassigned["total"], 4)
        self.assertTrue(all(row["assignee_id"] is None for row in unassigned["records"]))
        with self.assertRaises(AppApiException) as invalid_filter:
            self.service.page(
                1,
                20,
                {"assignee_id": str(self.user.id), "unassigned": True},
            )
        self.assertEqual(invalid_filter.exception.code, 400)

    def test_retrospective_aggregates_root_causes_recurrence_and_experiences(self):
        self.service.sync({"days": 7})
        runtime_alert = MarketingAlert.objects.get(
            workspace_id=self.workspace_id,
            alert_type="RUNTIME_ERROR",
        )
        slow_alert = MarketingAlert.objects.get(
            workspace_id=self.workspace_id,
            alert_type="SLOW_RESPONSE",
        )
        model_close = {
            "action": "CLOSE",
            "note": "调整模型超时配置后恢复",
            "root_cause_category": "MODEL",
            "resolution_type": "CONFIG_CHANGE",
            "impact_scope": "在线营销对话",
        }
        self.service.handle(str(runtime_alert.id), model_close)
        self.service.handle(str(runtime_alert.id), {"action": "REOPEN", "note": "问题复发"})
        self.service.handle(str(runtime_alert.id), model_close)
        self.service.handle(
            str(slow_alert.id),
            {
                "action": "CLOSE",
                "note": "重试后恢复",
                "root_cause_category": "MODEL",
                "resolution_type": "RETRY_RECOVERY",
                "impact_scope": "部分慢响应会话",
            },
        )

        result = self.service.retrospective({"days": 30})

        self.assertEqual(result["summary"]["closed_count"], 2)
        self.assertEqual(result["summary"]["classified_count"], 2)
        self.assertEqual(result["summary"]["recurrence_count"], 1)
        self.assertEqual(result["root_causes"][0]["value"], "MODEL")
        self.assertEqual(result["root_causes"][0]["count"], 2)
        self.assertEqual(len(result["resolution_types"]), 2)
        self.assertEqual(result["recurring_alerts"][0]["id"], str(runtime_alert.id))
        self.assertEqual(len(result["experiences"]), 2)
        self.assertEqual(len(self.service.retrospective_export_rows({"days": 30})), 2)

    def test_noise_analysis_identifies_repeats_merge_candidates_and_reuses_history(self):
        self.service.sync({"days": 7})
        runtime_alert = MarketingAlert.objects.get(
            workspace_id=self.workspace_id,
            alert_type="RUNTIME_ERROR",
        )
        slow_alert = MarketingAlert.objects.get(
            workspace_id=self.workspace_id,
            alert_type="SLOW_RESPONSE",
        )
        self.service.handle(
            str(runtime_alert.id),
            {
                "action": "CLOSE",
                "note": "调整模型超时配置后恢复",
                "root_cause_category": "MODEL",
                "resolution_type": "CONFIG_CHANGE",
                "impact_scope": "在线营销对话",
            },
        )
        self.service.handle(str(runtime_alert.id), {"action": "REOPEN", "note": "继续观察"})
        runtime_alert.refresh_from_db()
        slow_alert.target_route = runtime_alert.target_route
        slow_alert.save(update_fields=["target_route", "update_time"])
        for index in range(4):
            runtime_alert.current_count += 1
            self.service._append_action(
                runtime_alert,
                MarketingAlertAction.REFRESHED,
                runtime_alert.status,
                None,
                "系统",
                f"第 {index + 1} 次指标刷新",
            )
            runtime_alert.save(update_fields=["current_count", "revision", "update_time"])

        result = self.service.noise_analysis({"days": 7})

        runtime_noise = next(item for item in result["noise_alerts"] if item["id"] == str(runtime_alert.id))
        self.assertEqual(result["summary"]["deduplicated_event_count"], 4)
        self.assertEqual(result["summary"]["repeated_alert_count"], 1)
        self.assertEqual(result["summary"]["high_noise_count"], 1)
        self.assertEqual(runtime_noise["refresh_count"], 4)
        self.assertEqual(runtime_noise["reopen_count"], 1)
        self.assertEqual(runtime_noise["noise_level"], "HIGH")
        self.assertEqual(runtime_noise["suggested_root_cause_category"], "MODEL")
        self.assertEqual(runtime_noise["suggested_resolution_type"], "CONFIG_CHANGE")
        merge_group = next(
            item
            for item in result["merge_suggestions"]
            if str(runtime_alert.id) in {candidate["id"] for candidate in item["candidates"]}
        )
        self.assertIn(str(slow_alert.id), {candidate["id"] for candidate in merge_group["candidates"]})
        detail = self.service.get(str(runtime_alert.id))
        self.assertEqual(detail["related_experiences"][0]["close_note"], "调整模型超时配置后恢复")

    def test_rule_effect_recommends_threshold_and_history_can_rollback(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(
            workspace_id=self.workspace_id,
            alert_type="RUNTIME_ERROR",
        )
        rule = MarketingAlertRule.objects.get(
            workspace_id=self.workspace_id,
            alert_type="RUNTIME_ERROR",
        )
        for index in range(5):
            alert.current_count += 1
            self.service._append_action(
                alert,
                MarketingAlertAction.REFRESHED,
                alert.status,
                None,
                "系统",
                f"第 {index + 1} 次规则效果样本",
            )
            alert.save(update_fields=["current_count", "revision", "update_time"])

        effect = self.service.rule_effect({"days": 7})
        runtime_effect = next(
            item for item in effect["rules"] if item["rule"]["alert_type"] == "RUNTIME_ERROR"
        )
        self.assertEqual(runtime_effect["metrics"]["signal_event_count"], 6)
        self.assertEqual(runtime_effect["metrics"]["refresh_count"], 5)
        self.assertTrue(runtime_effect["recommendation"]["changed"])
        self.assertEqual(runtime_effect["recommendation"]["threshold"], 5)
        self.assertTrue(
            any(item["recommended"] for item in runtime_effect["simulations"])
        )

        updated = self.service.update_rule(
            str(rule.id),
            {"threshold": 5, "change_reason": "应用规则效果推荐值"},
        )
        self.assertEqual(updated["threshold"], 5)
        original_change = MarketingAlertRuleChange.objects.get(rule=rule, version=1)
        self.assertEqual(original_change.snapshot["threshold"], 1)
        self.assertEqual(original_change.reason, "应用规则效果推荐值")

        after_update = self.service.rule_effect({"days": 7})
        updated_effect = next(
            item for item in after_update["rules"] if item["rule"]["alert_type"] == "RUNTIME_ERROR"
        )
        self.assertEqual(updated_effect["history"][0]["version"], 1)

        restored = self.service.rollback_rule(
            str(rule.id),
            {"change_id": str(original_change.id), "reason": "验证历史版本回滚"},
        )
        self.assertEqual(restored["threshold"], 1)
        rollback_change = MarketingAlertRuleChange.objects.get(rule=rule, version=2)
        self.assertEqual(rollback_change.change_type, "ROLLED_BACK")
        self.assertEqual(rollback_change.snapshot["threshold"], 5)

    def test_manager_assigns_owner_and_transfer_keeps_collaboration_history(self):
        target_user = User.objects.create(
            username="marketing_assigned_owner",
            nick_name="一线负责人",
            password="x",
            role="USER",
        )
        MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="一线负责人关联任务",
            scenario=self.scenario,
            application=self.application,
            user=target_user,
            status=MarketingTaskStatus.RUNNING,
        )
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        original_resolution_due_at = alert.resolution_due_at

        with patch.object(MarketingAlertService, "_can_manage", return_value=True):
            assigned = self.service.collaborate(
                str(alert.id),
                {"action": "ASSIGN", "assignee_id": str(target_user.id), "note": "请一线先排查"},
            )

        self.assertEqual(assigned["status"], MarketingAlertStatus.CLAIMED)
        self.assertEqual(assigned["assignee_id"], str(target_user.id))
        self.assertIsNone(assigned["handling_rounds"][0]["first_responded_at"])

        target_service = MarketingAlertService(self.workspace_id, target_user)
        started = target_service.handle(str(alert.id), {"action": "START", "note": "开始检查调用链"})
        self.assertIsNotNone(started["handling_rounds"][0]["first_responded_at"])
        target_service.collaborate(str(alert.id), {"action": "COMMENT", "note": "已定位到上游超时"})
        transferred = target_service.collaborate(
            str(alert.id),
            {"action": "TRANSFER", "assignee_id": str(self.user.id), "note": "需要二线调整模型配置"},
        )

        self.assertEqual(transferred["status"], MarketingAlertStatus.CLAIMED)
        self.assertEqual(transferred["assignee_id"], str(self.user.id))
        self.assertEqual(len(transferred["handling_rounds"]), 2)
        latest_round, previous_round = transferred["handling_rounds"]
        self.assertIsNone(latest_round["ended_at"])
        self.assertEqual(latest_round["resolution_due_at"], original_resolution_due_at)
        self.assertEqual(previous_round["end_reason"], MarketingAlertRoundEndReason.TRANSFERRED)
        self.assertEqual(previous_round["end_note"], "需要二线调整模型配置")
        self.service.handle(str(alert.id), {"action": "START", "note": "二线接手"})
        now = timezone.now()
        MarketingAlertHandlingRound.objects.filter(alert=alert, sequence=1).update(
            assigned_at=now - timedelta(minutes=60),
            response_due_at=now - timedelta(minutes=45),
            first_responded_at=now - timedelta(minutes=50),
            ended_at=now - timedelta(minutes=30),
        )
        MarketingAlertHandlingRound.objects.filter(alert=alert, sequence=2).update(
            assigned_at=now - timedelta(minutes=30),
            response_due_at=now - timedelta(minutes=20),
            first_responded_at=now - timedelta(minutes=15),
        )
        performance = self.service.performance({"days": 7})
        first_owner = next(
            item for item in performance["workload"] if item["assignee_id"] == str(target_user.id)
        )
        second_owner = next(
            item for item in performance["workload"] if item["assignee_id"] == str(self.user.id)
        )
        self.assertEqual(first_owner["avg_response_seconds"], 600)
        self.assertEqual(first_owner["response_sla_rate"], 100)
        self.assertEqual(second_owner["avg_response_seconds"], 900)
        self.assertEqual(second_owner["response_sla_rate"], 0)
        actions = list(
            MarketingAlertActionRecord.objects.filter(alert=alert).values_list("action", flat=True)
        )
        self.assertIn(MarketingAlertAction.ASSIGNED, actions)
        self.assertIn(MarketingAlertAction.COMMENTED, actions)
        self.assertIn(MarketingAlertAction.TRANSFERRED, actions)

    def test_assignment_requires_manager_and_workspace_member(self):
        target_user = User.objects.create(
            username="marketing_assignment_target",
            nick_name="空间内成员",
            password="x",
            role="USER",
        )
        MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="空间内成员任务",
            scenario=self.scenario,
            application=self.application,
            user=target_user,
            status=MarketingTaskStatus.RUNNING,
        )
        outside_user = User.objects.create(
            username="marketing_assignment_outside",
            nick_name="空间外成员",
            password="x",
            role="USER",
        )
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")

        with patch.object(MarketingAlertService, "_can_manage", return_value=False):
            with self.assertRaises(AppApiException) as forbidden:
                self.service.collaborate(
                    str(alert.id),
                    {"action": "ASSIGN", "assignee_id": str(target_user.id)},
                )
        self.assertEqual(forbidden.exception.code, 403)

        with patch.object(MarketingAlertService, "_can_manage", return_value=True):
            with self.assertRaises(AppApiException) as outside_workspace:
                self.service.collaborate(
                    str(alert.id),
                    {"action": "ASSIGN", "assignee_id": str(outside_user.id)},
                )
        self.assertEqual(outside_workspace.exception.code, 400)
        self.assertFalse(MarketingAlertHandlingRound.objects.filter(alert=alert).exists())

    def test_transfer_requires_reason_and_different_owner(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        self.service.handle(str(alert.id), {"action": "CLAIM"})

        with self.assertRaises(AppApiException) as missing_reason:
            self.service.collaborate(
                str(alert.id),
                {"action": "TRANSFER", "assignee_id": str(self.user.id), "note": ""},
            )
        self.assertEqual(missing_reason.exception.code, 400)
        with self.assertRaises(AppApiException) as same_owner:
            self.service.collaborate(
                str(alert.id),
                {"action": "TRANSFER", "assignee_id": str(self.user.id), "note": "转给自己"},
            )
        self.assertEqual(same_owner.exception.code, 400)

    def test_alert_subscribers_are_workspace_scoped_and_recorded_in_timeline(self):
        subscriber = User.objects.create(
            username="marketing_alert_subscriber",
            nick_name="协作关注人",
            password="x",
            role="USER",
        )
        MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="关注人关联任务",
            scenario=self.scenario,
            application=self.application,
            user=subscriber,
            status=MarketingTaskStatus.RUNNING,
        )
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")

        added = self.service.update_subscriber(
            str(alert.id),
            {"action": "ADD", "user_id": str(subscriber.id)},
        )

        self.assertEqual(added["subscribers"][0]["user_name"], "协作关注人")
        self.assertEqual(MarketingAlertSubscriber.objects.filter(alert=alert).count(), 1)
        self.assertEqual(added["actions"][0]["action"], MarketingAlertAction.SUBSCRIBED)
        with self.assertRaises(AppApiException) as duplicate:
            self.service.update_subscriber(
                str(alert.id),
                {"action": "ADD", "user_id": str(subscriber.id)},
            )
        self.assertEqual(duplicate.exception.code, 400)

        removed = self.service.update_subscriber(
            str(alert.id),
            {"action": "REMOVE", "user_id": str(subscriber.id)},
        )
        self.assertEqual(removed["subscribers"], [])
        self.assertEqual(removed["actions"][0]["action"], MarketingAlertAction.UNSUBSCRIBED)

    def test_alert_attachment_validates_type_size_and_private_repository_source(self):
        self.service.sync({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")
        upload = SimpleUploadedFile("runtime.log", b"upstream timeout", content_type="text/plain")

        result = self.service.upload_attachment(str(alert.id), upload)

        self.assertEqual(result["attachments"][0]["file_name"], "runtime.log")
        self.assertIn("/attachment/", result["attachments"][0]["download_url"])
        attachment = MarketingAlertAttachment.objects.get(alert=alert)
        stored_file = File.objects.get(id=attachment.file_id)
        self.assertEqual(stored_file.source_type, "MARKETING")
        self.assertEqual(stored_file.get_bytes(), b"upstream timeout")
        self.assertEqual(
            self.service.attachment_file(str(alert.id), str(attachment.id))[0],
            attachment,
        )
        with self.assertRaises(AppApiException) as invalid_type:
            self.service.upload_attachment(
                str(alert.id),
                SimpleUploadedFile("payload.exe", b"x", content_type="application/octet-stream"),
            )
        self.assertEqual(invalid_type.exception.code, 400)
        with self.assertRaises(AppApiException) as too_large:
            self.service.upload_attachment(
                str(alert.id),
                SimpleUploadedFile("huge.log", b"x" * (10 * 1024 * 1024 + 1), content_type="text/plain"),
            )
        self.assertEqual(too_large.exception.code, 400)

    def test_assignment_and_collaboration_notifications_are_directed(self):
        target = User.objects.create(
            username="marketing_direct_notice_target",
            nick_name="定向通知人",
            password="x",
            role="USER",
        )
        MarketingTask.objects.create(
            workspace_id=self.workspace_id,
            name="定向通知关联任务",
            scenario=self.scenario,
            application=self.application,
            user=target,
            status=MarketingTaskStatus.RUNNING,
        )
        self.notification_service.sync_and_notify({"days": 7})
        alert = MarketingAlert.objects.get(workspace_id=self.workspace_id, alert_type="RUNTIME_ERROR")

        with patch.object(MarketingAlertService, "_can_manage", return_value=True):
            self.service.collaborate(
                str(alert.id),
                {"action": "ASSIGN", "assignee_id": str(target.id), "note": "请立即排查"},
            )
        created = self.notification_service.dispatch()

        self.assertEqual(created, 1)
        assignment_action = MarketingAlertActionRecord.objects.get(
            alert=alert,
            action=MarketingAlertAction.ASSIGNED,
        )
        notification = MarketingNotification.objects.get(action_record=assignment_action)
        self.assertEqual(notification.user, target)
        self.assertIn("告警已指派给你", notification.title)

        target_service = MarketingAlertService(self.workspace_id, target)
        target_service.update_subscriber(str(alert.id), {"action": "ADD", "user_id": str(self.user.id)})
        self.notification_service.dispatch()
        target_service.collaborate(str(alert.id), {"action": "COMMENT", "note": "补充排查结果"})
        self.assertEqual(self.notification_service.dispatch(), 1)
        comment_action = MarketingAlertActionRecord.objects.get(
            alert=alert,
            action=MarketingAlertAction.COMMENTED,
        )
        self.assertEqual(MarketingNotification.objects.get(action_record=comment_action).user, self.user)


class MarketingConversionTests(TestCase):
    def setUp(self):
        self.workspace_id = "marketing-conversion-workspace"
        self.user = User.objects.create(
            username="marketing_conversion_tester", nick_name="转化测试员", password="x", role="USER"
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id, name="套餐转化场景", user=self.user
        )
        self.task = MarketingTask.objects.create(
            workspace_id=self.workspace_id, name="套餐转化任务", scenario=self.scenario, user=self.user
        )
        self.customer = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="李女士",
            phone="13900000000",
        )
        self.batch = MarketingTouchBatch.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            user=self.user,
            name="套餐升档批次",
            approval_status=MarketingTouchBatchApprovalStatus.APPROVED,
            customer_snapshot=[{"customer_id": str(self.customer.id)}],
            selected_count=1,
        )

    def test_conversion_history_updates_customer_and_dashboard_attribution(self):
        service = MarketingCustomerService(self.workspace_id, self.user)
        service.record_conversion(
            str(self.task.id),
            str(self.customer.id),
            {
                "outcome": "CONVERTED",
                "product_name": "199 元套餐升档",
                "note": "客户确认办理",
                "batch_id": str(self.batch.id),
            },
        )
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.status, MarketingCustomerStatus.CONVERTED)
        detail = MarketingTouchBatchService(self.workspace_id, self.user).get(str(self.task.id), str(self.batch.id))
        self.assertEqual(detail["metrics"]["converted_count"], 1)
        self.assertEqual(detail["metrics"]["conversion_rate"], 100.0)

        service.record_conversion(
            str(self.task.id),
            str(self.customer.id),
            {"outcome": "LOST", "note": "客户确认暂不办理", "batch_id": str(self.batch.id)},
        )
        self.customer.refresh_from_db()
        dashboard = MarketingTaskService(self.workspace_id, self.user).dashboard(str(self.task.id))
        self.assertEqual(MarketingConversionRecord.objects.filter(customer=self.customer).count(), 2)
        self.assertEqual(self.customer.status, MarketingCustomerStatus.NO_INTENT)
        self.assertEqual(dashboard["funnel"]["lost_count"], 1)


class MarketingCustomerImportTests(TestCase):
    def setUp(self):
        self.workspace_id = "marketing-import-workspace"
        self.user = User.objects.create(
            username="marketing_import_tester", nick_name="导入测试员", password="x", role="USER"
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id, name="名单导入场景", user=self.user
        )
        self.task = MarketingTask.objects.create(
            workspace_id=self.workspace_id, name="名单导入任务", scenario=self.scenario, user=self.user
        )
        self.service = MarketingCustomerService(self.workspace_id, self.user)
        MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task=self.task,
            name="已有客户",
            phone="13700000000",
        )

    def test_preview_confirm_and_history_keep_existing_customers_unchanged(self):
        uploaded_file = SimpleUploadedFile(
            "客户名单.csv",
            "客户名称,手机号,省份,当前套餐,近月流量,近月通话\n张女士,13800000000,广东,199元套餐,80GB,300分钟\n已有客户,13700000000,广东,199元套餐,60GB,200分钟\n重复客户,13800000000,广东,199元套餐,50GB,100分钟\n缺标识,,广东,129元套餐,20GB,50分钟\n".encode(),
            content_type="text/csv",
        )
        preview = self.service.preview_import(str(self.task.id), uploaded_file)
        self.assertEqual(preview["summary"], {"total_count": 4, "ready_count": 1, "skipped_count": 2, "error_count": 1})
        self.assertEqual(preview["mapping"]["phone"], "手机号")

        result = self.service.confirm_import(
            str(self.task.id),
            {"file_name": preview["file_name"], "mapping": preview["mapping"], "rows": preview["rows"]},
        )
        batch = result["batch"]
        self.assertEqual(batch["imported_count"], 1)
        self.assertEqual(batch["skipped_count"], 2)
        self.assertEqual(batch["error_count"], 1)
        self.assertEqual(MarketingCustomer.objects.filter(task=self.task).count(), 2)
        imported = MarketingCustomer.objects.get(task=self.task, phone="13800000000")
        self.assertEqual(imported.status, MarketingCustomerStatus.PENDING)
        self.assertEqual(imported.config["marketing_import_batch_id"], batch["id"])
        self.assertEqual(MarketingCustomerImportBatch.objects.filter(task=self.task).count(), 1)
        self.assertEqual(self.service.import_history(str(self.task.id))[0]["id"], batch["id"])

    def test_preview_supports_manual_column_mapping(self):
        uploaded_file = SimpleUploadedFile(
            "custom-columns.csv",
            "客户,客户编号,归属省\n王先生,CUSTOMER-001,浙江\n".encode(),
            content_type="text/csv",
        )
        preview = self.service.preview_import(
            str(self.task.id),
            uploaded_file,
            '{"name":"客户","phone":"客户编号","province":"归属省"}',
        )
        self.assertEqual(preview["summary"]["ready_count"], 1)
        self.assertEqual(preview["rows"][0]["data"]["phone"], "CUSTOMER-001")

    def test_preview_supports_xlsx(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["客户名称", "手机号", "当前套餐"])
        worksheet.append(["赵女士", "13600000000", "129元套餐"])
        content = BytesIO()
        workbook.save(content)
        uploaded_file = SimpleUploadedFile(
            "客户名单.xlsx",
            content.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        preview = self.service.preview_import(str(self.task.id), uploaded_file)
        self.assertEqual(preview["summary"]["ready_count"], 1)
        self.assertEqual(preview["rows"][0]["data"]["current_plan"], "129元套餐")


class MarketingRecommendationTests(SimpleTestCase):
    policy = {
        "enabled": True,
        "offers": [
            {
                "id": "same-price-data-expansion",
                "name": "199 套餐同价扩容核实活动",
                "benefit": "优先核实同价扩容权益。",
                "plan_keywords": ["199"],
                "provinces": ["广东"],
                "min_data_usage_gb": 60,
                "min_call_usage_minutes": 200,
                "fee_note": "费用以当地实际可办理方案为准。",
                "contract_note": "办理前会确认优惠期限和合约期。",
            }
        ],
    }

    def test_selects_only_when_all_customer_conditions_match(self):
        recommendation = select_recommendation(
            self.policy,
            {"plan": "199元套餐", "province": "广东", "data_usage": "80GB", "call_usage": "300分钟"},
        )
        self.assertEqual(recommendation["status"], "MATCHED")
        self.assertEqual(recommendation["offer_id"], "same-price-data-expansion")
        unmatched = select_recommendation(
            self.policy,
            {"plan": "129元套餐", "province": "广东", "data_usage": "80GB", "call_usage": "300分钟"},
        )
        self.assertEqual(unmatched["status"], "NO_MATCH")

    def test_returns_customer_visible_configured_offer_for_recommendation_question(self):
        recommendation = select_recommendation(
            self.policy,
            {"plan": "199元套餐", "province": "广东", "data_usage": "80GB", "call_usage": "300分钟"},
        )
        answer = get_customer_recommendation_answer(
            "我流量经常不够，推荐什么套餐？",
            {"marketing_customer_id": "customer-1", "marketing_recommendation": recommendation},
            "模型原始回答",
        )
        self.assertIn("199 套餐同价扩容核实活动", answer)
        self.assertNotIn("。。", answer)

    def test_priority_and_effective_dates_choose_only_the_active_top_offer(self):
        policy = {
            "enabled": True,
            "offers": [
                {**self.policy["offers"][0], "id": "regular", "priority": 20},
                {
                    **self.policy["offers"][0],
                    "id": "priority",
                    "name": "优先扩容活动",
                    "priority": 10,
                    "effective_start": "2026-01-01",
                    "effective_end": "2026-12-31",
                },
                {
                    **self.policy["offers"][0],
                    "id": "expired",
                    "name": "过期活动",
                    "priority": 1,
                    "effective_end": "2025-12-31",
                },
            ],
        }
        result = select_recommendation(
            policy,
            {"plan": "199元套餐", "province": "广东", "data_usage": "80GB", "call_usage": "300分钟"},
            on_date=date(2026, 9, 1),
        )
        self.assertEqual(result["offer_id"], "priority")
        self.assertEqual(result["priority"], 10)

    def test_rejects_same_priority_in_one_exclusive_group(self):
        policy = {
            "enabled": True,
            "offers": [
                {**self.policy["offers"][0], "priority": 10, "exclusive_group": "199升级"},
                {
                    **self.policy["offers"][0],
                    "id": "other",
                    "name": "另一活动",
                    "priority": 10,
                    "exclusive_group": "199升级",
                },
            ],
        }
        with self.assertRaises(AppApiException):
            validate_recommendation_policy(policy)


class MarketingCampaignSnapshotTests(TestCase):
    def setUp(self):
        self.workspace_id = "marketing-campaign-workspace"
        self.user = User.objects.create(
            username="marketing_campaign_tester",
            nick_name="活动测试员",
            password="x",
            role="USER",
        )
        self.application = Application.objects.create(
            workspace_id=self.workspace_id,
            name="活动快照测试智能体",
            user=self.user,
        )
        self.scenario = MarketingScenario.objects.create(
            workspace_id=self.workspace_id,
            name="活动快照测试场景",
            application=self.application,
            user=self.user,
            config={
                "recommendation_policy": {
                    "enabled": True,
                    "offers": [
                        {
                            "id": "legacy-offer",
                            "name": "不应命中的旧场景活动",
                            "benefit": "旧场景权益",
                            "fee_note": "以当地实际可办理方案为准。",
                            "contract_note": "办理前确认合约期。",
                        }
                    ],
                }
            },
        )
        self.campaign = MarketingCampaign.objects.create(
            workspace_id=self.workspace_id,
            user=self.user,
            name="199 同价扩容活动",
            code="CAMPAIGN-199-DATA",
            status=MarketingCampaignStatus.ACTIVE,
            priority=10,
            benefit="优先核实 199 元套餐同价扩容权益。",
            fee_note="费用以当地实际可办理方案为准。",
            contract_note="办理前会确认优惠期限和合约期。",
            plan_keywords=["199"],
            provinces=["广东"],
            min_data_usage_gb=60,
            min_call_usage_minutes=200,
        )

    def test_task_keeps_activity_snapshot_after_activity_changes(self):
        task = MarketingTaskService(self.workspace_id, self.user).create(
            {
                "name": "199 客户触达任务",
                "scenario_id": str(self.scenario.id),
                "application_id": str(self.application.id),
                "campaign_ids": [str(self.campaign.id)],
            }
        )
        self.assertEqual(task["campaign_ids"], [str(self.campaign.id)])
        self.assertEqual(task["campaign_names"], [self.campaign.name])

        customer = MarketingCustomer.objects.create(
            workspace_id=self.workspace_id,
            task_id=task["id"],
            name="张女士",
            phone="13800000000",
            province="广东",
            current_plan="199 元套餐",
            monthly_data_usage="80GB",
            monthly_call_usage="300 分钟",
        )
        updated = MarketingCampaignService(self.workspace_id, self.user).update(
            str(self.campaign.id),
            {
                "name": self.campaign.name,
                "code": self.campaign.code,
                "status": MarketingCampaignStatus.PAUSED,
                "priority": 10,
                "benefit": "不应覆盖历史任务的新权益。",
                "fee_note": self.campaign.fee_note,
                "contract_note": self.campaign.contract_note,
                "plan_keywords": ["199"],
                "provinces": ["广东"],
                "min_data_usage_gb": 60,
                "min_call_usage_minutes": 200,
            },
        )
        self.assertEqual(updated["version"], 2)

        source = build_marketing_source(self.application.id, task["id"], customer.id)
        recommendation = source["marketing_recommendation"]
        self.assertEqual(recommendation["status"], "MATCHED")
        self.assertEqual(recommendation["offer_id"], str(self.campaign.id))
        self.assertEqual(recommendation["benefit"], "优先核实 199 元套餐同价扩容权益。")

    def test_new_campaign_must_be_submitted_before_approval(self):
        service = MarketingCampaignService(self.workspace_id, self.user)
        campaign = service.create(
            {
                "name": "待审批活动",
                "code": "PENDING-APPROVAL-001",
                "status": "ACTIVE",
                "benefit": "客户可见的待审批权益。",
            }
        )
        self.assertEqual(campaign["status"], "DRAFT")
        self.assertEqual(campaign["approval_status"], "DRAFT")
        submitted = service.submit(campaign["id"])
        self.assertEqual(submitted["status"], "DRAFT")
        self.assertEqual(submitted["approval_status"], "PENDING")

    def test_expand_action_creates_a_new_pending_touch_batch(self):
        task = MarketingTaskService(self.workspace_id, self.user).create(
            {
                "name": "活动放量任务",
                "scenario_id": str(self.scenario.id),
                "application_id": str(self.application.id),
                "campaign_ids": [str(self.campaign.id)],
            }
        )
        for index in range(2):
            MarketingCustomer.objects.create(
                workspace_id=self.workspace_id,
                task_id=task["id"],
                name=f"放量客户{index}",
                phone=f"1390000000{index}",
                province="广东",
                current_plan="199 元套餐",
                monthly_data_usage="80GB",
                monthly_call_usage="300 分钟",
            )

        result = MarketingCampaignService(self.workspace_id, self.user).execute_recommendation(
            str(self.campaign.id),
            {"action": "EXPAND", "task_id": task["id"]},
        )

        self.assertEqual(result["action"], "EXPAND")
        self.assertEqual(result["batch"]["approval_status"], "PENDING")
        self.assertEqual(result["batch"]["selected_count"], 2)

    def test_pause_action_stops_future_activity_selection(self):
        result = MarketingCampaignService(self.workspace_id, self.user).execute_recommendation(
            str(self.campaign.id),
            {"action": "PAUSE"},
        )
        self.assertEqual(result["campaign"]["status"], "PAUSED")

    def test_approved_history_can_be_restored_as_a_new_draft_version(self):
        service = MarketingCampaignService(self.workspace_id, self.user)
        original_benefit = self.campaign.benefit
        task = MarketingTaskService(self.workspace_id, self.user).create(
            {
                "name": "版本恢复快照任务",
                "scenario_id": str(self.scenario.id),
                "application_id": str(self.application.id),
                "campaign_ids": [str(self.campaign.id)],
            }
        )
        service.update(
            str(self.campaign.id),
            {
                "name": self.campaign.name,
                "code": self.campaign.code,
                "benefit": "待审批的新权益内容。",
                "priority": 20,
                "plan_keywords": ["199"],
                "provinces": ["广东"],
            },
        )

        history = service.versions(str(self.campaign.id))["records"]
        self.assertEqual([item["version"] for item in history], [2, 1])
        self.assertEqual(history[1]["approval_status"], MarketingCampaignApprovalStatus.APPROVED)
        self.assertTrue(history[1]["can_restore"])

        restored = service.restore(str(self.campaign.id), 1)
        self.assertEqual(restored["version"], 3)
        self.assertEqual(restored["benefit"], original_benefit)
        self.assertEqual(restored["status"], MarketingCampaignStatus.PAUSED)
        self.assertEqual(restored["approval_status"], MarketingCampaignApprovalStatus.DRAFT)
        restored_version = MarketingCampaignVersion.objects.get(campaign=self.campaign, version=3)
        self.assertEqual(restored_version.change_type, "RESTORE")
        self.assertEqual(restored_version.source_version, 1)

        task_row = MarketingTask.objects.get(id=task["id"])
        task_snapshot = task_row.config["marketing_campaign_snapshot"][0]
        self.assertEqual(task_snapshot["campaign_version"], 1)
        self.assertEqual(task_snapshot["benefit"], original_benefit)

    def test_draft_version_cannot_be_restored(self):
        service = MarketingCampaignService(self.workspace_id, self.user)
        service.update(
            str(self.campaign.id),
            {
                "name": self.campaign.name,
                "code": self.campaign.code,
                "benefit": "未审批版本。",
            },
        )

        with self.assertRaises(AppApiException) as context:
            service.restore(str(self.campaign.id), 2)

        self.assertEqual(context.exception.code, 400)

    def test_effect_contains_province_and_plan_breakdowns(self):
        task = MarketingTaskService(self.workspace_id, self.user).create(
            {
                "name": "活动分层复盘任务",
                "scenario_id": str(self.scenario.id),
                "application_id": str(self.application.id),
                "campaign_ids": [str(self.campaign.id)],
            }
        )
        customer_specs = [
            ("广东客户", "13600000001", "广东", "199 元套餐", MarketingCustomerStatus.INTENT),
            ("浙江客户", "13600000002", "浙江", "299 元套餐", MarketingCustomerStatus.CONVERTED),
        ]
        for name, phone, province, current_plan, status in customer_specs:
            customer = MarketingCustomer.objects.create(
                workspace_id=self.workspace_id,
                task_id=task["id"],
                name=name,
                phone=phone,
                province=province,
                current_plan=current_plan,
                status=status,
            )
            Chat.objects.create(
                application=self.application,
                abstract=f"{name}活动会话",
                source={
                    "marketing_task_id": task["id"],
                    "marketing_customer_id": str(customer.id),
                    "marketing_recommendation": {"offer_id": str(self.campaign.id)},
                },
            )

        effect = MarketingCampaignService(self.workspace_id, self.user).effect(str(self.campaign.id))

        self.assertEqual(effect["metrics"]["customer_count"], 2)
        self.assertEqual([item["name"] for item in effect["breakdowns"]["province"]], ["广东", "浙江"])
        self.assertEqual(
            [item["name"] for item in effect["breakdowns"]["current_plan"]],
            ["199 元套餐", "299 元套餐"],
        )
        self.assertEqual(effect["breakdowns"]["province"][0]["intent_rate"], 100.0)

        workbook = load_workbook(build_campaign_effect_workbook(effect))
        self.assertEqual(workbook.sheetnames, ["复盘汇总", "省份效果", "套餐效果"])
        self.assertEqual(workbook["复盘汇总"]["B8"].value, 2)
        self.assertEqual(workbook["省份效果"]["A2"].value, "广东")
        self.assertEqual(workbook["省份效果"]["G2"].value, 1)
        self.assertEqual(workbook["省份效果"]["G2"].number_format, "0.0%")
