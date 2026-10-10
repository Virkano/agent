# coding=utf-8
from io import BytesIO
from html import escape
from pathlib import Path
import uuid
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


EXPORT_CONTENT_TYPES = {
    "html": "text/html; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
REPORT_FONT_NAME = "AlibabaPuHuiTi"
DOCX_FONT_NAME = "Alibaba PuHuiTi 3.0"
REPORT_FONT_PATH = Path(__file__).resolve().parents[2] / "ui/src/styles/font/AlibabaPuHuiTi-3-55-Regular.ttf"


def _embed_docx_font(docx_bytes):
    """Embed the project CJK font so Word exports render consistently on font-light hosts."""
    font_key = uuid.uuid4()
    key_bytes = font_key.bytes_le
    font_bytes = bytearray(REPORT_FONT_PATH.read_bytes())
    for index in range(min(32, len(font_bytes))):
        font_bytes[index] ^= key_bytes[index % 16]

    with ZipFile(BytesIO(docx_bytes), "r") as source:
        parts = {name: source.read(name) for name in source.namelist()}

    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    font_table = etree.fromstring(parts["word/fontTable.xml"])
    font_node = etree.SubElement(font_table, f"{{{w_ns}}}font")
    font_node.set(f"{{{w_ns}}}name", DOCX_FONT_NAME)
    embed_node = etree.SubElement(font_node, f"{{{w_ns}}}embedRegular")
    embed_node.set(f"{{{r_ns}}}id", "rIdEmbeddedReportFont")
    embed_node.set(f"{{{w_ns}}}fontKey", "{" + str(font_key).upper() + "}")
    parts["word/fontTable.xml"] = etree.tostring(
        font_table, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    rel_path = "word/_rels/fontTable.xml.rels"
    if rel_path in parts:
        relationships = etree.fromstring(parts[rel_path])
    else:
        relationships = etree.Element(f"{{{rel_ns}}}Relationships", nsmap={None: rel_ns})
    relationship = etree.SubElement(relationships, f"{{{rel_ns}}}Relationship")
    relationship.set("Id", "rIdEmbeddedReportFont")
    relationship.set("Type", f"{r_ns}/font")
    relationship.set("Target", "fonts/report-font.odttf")
    parts[rel_path] = etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True)

    content_types = etree.fromstring(parts["[Content_Types].xml"])
    if not content_types.xpath("ct:Default[@Extension='odttf']", namespaces={"ct": content_ns}):
        default = etree.SubElement(content_types, f"{{{content_ns}}}Default")
        default.set("Extension", "odttf")
        default.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")
    parts["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    parts["word/fonts/report-font.odttf"] = bytes(font_bytes)

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)
    return output.getvalue()


def safe_report_filename(report, extension):
    safe_title = "".join(char if char.isalnum() or char in "-_" else "_" for char in report.title).strip("_")
    return f"{safe_title[:80] or 'marketing-report'}.{extension}"


def _metric_rows(report):
    metrics = report.metrics or {}
    return [
        ("统计周期", f"{report.period_start.isoformat()} 至 {report.period_end.isoformat()}"),
        ("归因方式", report.get_attribution_method_display()),
        ("曝光人数", metrics.get("exposure_count", 0)),
        ("实际触达人数", metrics.get("reached_customer_count", 0)),
        ("触达率", f"{metrics.get('touch_rate', 0)}%"),
        ("确认转化人数", metrics.get("converted_customer_count", 0)),
        ("触达后转化率", f"{metrics.get('conversion_rate', 0)}%"),
        ("归因收入", f"¥{metrics.get('attributed_revenue', 0):.2f}"),
        ("营销投入", f"¥{metrics.get('marketing_cost', 0):.2f}"),
        ("边际贡献", f"¥{metrics.get('marginal_contribution', 0):.2f}"),
        ("ROI", f"{metrics.get('roi', 0)}%"),
    ]


def build_html_export(report):
    style = """
<style>
body{margin:0;background:#eef2f1;color:#27323b;font-family:Arial,'Noto Sans SC',sans-serif}
.export-page{max-width:900px;margin:30px auto;padding:52px 60px;border-top:5px solid #0d7c7c;background:#fbfaf6;box-shadow:0 18px 42px #1f343c26}
table{width:100%;border-collapse:collapse}th,td{padding:10px 12px;border-bottom:1px solid #dde2df;text-align:left}th{width:40%;color:#68767b}
.attribution-bar{margin:12px 0}.attribution-bar div{display:flex;justify-content:space-between}.attribution-bar i{display:block;height:8px;margin-top:6px;border-radius:8px;background:#0d7c7c}
</style>
"""
    return (
        "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\"><title>"
        + escape(report.title)
        + "</title>"
        + style
        + "</head><body><main class=\"export-page\">"
        + report.html_content
        + "</main></body></html>"
    ).encode("utf-8")


def build_docx_export(report):
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = DOCX_FONT_NAME
    normal_style.element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_NAME)
    title = document.add_heading(report.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        f"{report.period_start.isoformat()} 至 {report.period_end.isoformat()} · {report.get_attribution_method_display()}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("核心经营指标", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "结果"
    for label, value in _metric_rows(report):
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)

    document.add_heading("经营解读", level=1)
    document.add_paragraph(report.narrative or "暂无经营解读。")
    document.add_heading("渠道归因分布", level=1)
    distribution = (report.metrics or {}).get("attribution_distribution", [])
    if distribution:
        attribution_table = document.add_table(rows=1, cols=3)
        attribution_table.style = "Light Shading Accent 1"
        for cell, text in zip(attribution_table.rows[0].cells, ("渠道", "归因收入", "收入占比"), strict=True):
            cell.text = text
        for item in distribution:
            cells = attribution_table.add_row().cells
            cells[0].text = str(item["channel"])
            cells[1].text = f"¥{item['attributed_revenue']:.2f}"
            cells[2].text = f"{item['revenue_share']}%"
    else:
        document.add_paragraph("当前周期暂无可归因转化记录。")
    paragraphs = list(document.paragraphs)
    for table_row in document.tables:
        for row in table_row.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.name = DOCX_FONT_NAME
            run._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_NAME)
            run.font.size = Pt(10.5)
    output = BytesIO()
    document.save(output)
    return _embed_docx_font(output.getvalue())


def build_pdf_export(report):
    output = BytesIO()
    if REPORT_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(REPORT_FONT_NAME, REPORT_FONT_PATH))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=styles["Title"],
        fontName=REPORT_FONT_NAME,
        fontSize=20,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#13202D"),
        spaceAfter=8,
    )
    body_style = ParagraphStyle(
        "ChineseBody",
        parent=styles["BodyText"],
        fontName=REPORT_FONT_NAME,
        fontSize=10.5,
        leading=18,
        textColor=colors.HexColor("#34434A"),
    )
    heading_style = ParagraphStyle(
        "ChineseHeading",
        parent=body_style,
        fontSize=13,
        leading=20,
        spaceBefore=10,
        textColor=colors.HexColor("#0D6D70"),
    )
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=22 * mm,
        leftMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=report.title,
        author="AI 营销分析报告中心",
    )
    story = [
        Paragraph(report.title, title_style),
        Paragraph(
            f"{report.period_start.isoformat()} 至 {report.period_end.isoformat()} · {report.get_attribution_method_display()}",
            body_style,
        ),
        Spacer(1, 8 * mm),
        Paragraph("核心经营指标", heading_style),
    ]
    table_data = [[Paragraph("指标", body_style), Paragraph("结果", body_style)]]
    table_data.extend([[Paragraph(str(label), body_style), Paragraph(str(value), body_style)] for label, value in _metric_rows(report)])
    metrics_table = Table(table_data, colWidths=[65 * mm, 80 * mm], repeatRows=1)
    metrics_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEDEC")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CAD5D4")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend([metrics_table, Paragraph("经营解读", heading_style), Paragraph(report.narrative or "暂无经营解读。", body_style)])
    story.append(Paragraph("渠道归因分布", heading_style))
    distribution = (report.metrics or {}).get("attribution_distribution", [])
    if distribution:
        channel_data = [[Paragraph(text, body_style) for text in ("渠道", "归因收入", "收入占比")]]
        channel_data.extend(
            [
                Paragraph(str(item["channel"]), body_style),
                Paragraph(f"¥{item['attributed_revenue']:.2f}", body_style),
                Paragraph(f"{item['revenue_share']}%", body_style),
            ]
            for item in distribution
        )
        channel_table = Table(channel_data, colWidths=[55 * mm, 50 * mm, 40 * mm], repeatRows=1)
        channel_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEDEC")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CAD5D4")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("PADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(channel_table)
    else:
        story.append(Paragraph("当前周期暂无可归因转化记录。", body_style))
    doc.build(story)
    return output.getvalue()


def build_report_export(report, export_format):
    builders = {
        "html": build_html_export,
        "pdf": build_pdf_export,
        "docx": build_docx_export,
    }
    if export_format not in builders:
        raise ValueError("不支持的报告格式")
    return builders[export_format](report)
