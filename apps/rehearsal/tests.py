import json
from datetime import timedelta
from io import BytesIO
from types import SimpleNamespace

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework.exceptions import ValidationError
from openpyxl import load_workbook
from pypdf import PdfReader

from common.exception.app_exception import AppApiException
from users.models import User

from .exports import (
    build_batch_validation_csv,
    build_complaint_csv,
    build_data_simulation_csv,
    build_import_template_csv,
    build_process_run_csv,
    build_simulation_record_csv,
    build_traceback_csv,
)
from .auxiliary_exports import build_auxiliary_json, build_auxiliary_pdf, build_auxiliary_xlsx
from .models import (
    AlertStatus,
    AuxiliarySimulationStatus,
    CapabilityCallStatus,
    CapabilityMockKind,
    ComplaintBatchStatus,
    ComplaintValidationStatus,
    DataModelBuildStrategy,
    DataModelQualityStatus,
    DataSimulationKind,
    FaultConfigurationAction,
    FaultConfigurationStatus,
    FaultReportAction,
    FaultReportLevel,
    FaultReportStatus,
    RehearsalComplaint,
    RehearsalComplaintBatch,
    RehearsalComplaintTemplate,
    RehearsalDataBatch,
    RehearsalDataModel,
    RehearsalDataModelLoad,
    RehearsalDataModelVersion,
    RehearsalDataPoint,
    RehearsalDataTemplate,
    RehearsalInstance,
    RehearsalInstanceEvent,
    RehearsalInstanceStatus,
    RehearsalAlert,
    RehearsalAlertEvent,
    RehearsalAlertTemplate,
    RehearsalAuxiliarySimulation,
    RehearsalCapabilityCall,
    RehearsalCapabilityMock,
    RehearsalFaultConfiguration,
    RehearsalFaultConfigurationEvent,
    RehearsalFaultReport,
    RehearsalFaultReportEvent,
    RehearsalFaultReportTemplate,
    RehearsalProcessEvent,
    RehearsalProcessRun,
    RehearsalProcessStep,
    RehearsalProcessTemplate,
    RehearsalSceneEvent,
    RehearsalSceneLibrary,
    RehearsalSceneLibraryEvent,
    RehearsalSceneReplay,
    RehearsalSceneRun,
    RehearsalSceneScript,
    RehearsalSceneStep,
    RehearsalTracebackSnapshot,
    RehearsalWorkOrder,
    RehearsalWorkOrderEvent,
    ProcessRunStatus,
    ProcessStepStatus,
    SceneLibraryStatus,
    SceneReplayStatus,
    SceneRunStatus,
    SceneStepStatus,
    TracebackScope,
    TracebackStatus,
    WorkOrderStatus,
)
from .serializers import RehearsalService


class RehearsalFixture(TestCase):
    def setUp(self):
        self.workspace_id = "rehearsal-test"
        self.user = User.objects.create(
            username="rehearsal_admin",
            nick_name="预演管理员",
            password="x",
            role="WORKSPACE_MANAGE",
        )
        self.service = RehearsalService(self.workspace_id, self.user)
        self.scenario, self.templates = self.service.ensure_default_catalog()

    def create_instance(self, suffix="001"):
        return self.service.create_instance(
            {
                "request_id": f"instance:{suffix}",
                "scenario_id": str(self.scenario.id),
                "name": f"客户申告预演 {suffix}",
            }
        )["instance"]

    def create_ready_instance(self, suffix="process"):
        instance = self.create_instance(suffix)
        template = next(item for item in self.templates if item.customer_type == "PUBLIC")
        self.service.generate_complaints(
            instance["id"],
            {
                "request_id": f"batch:{suffix}",
                "template_id": str(template.id),
                "count": 3,
                "start_at": timezone.now().isoformat(),
                "interval_minutes": 5,
            },
        )
        self.service.operate_instance(instance["id"], {"action": "PREPARE"})
        return self.service.instance_detail(instance["id"])


class RehearsalServiceTests(RehearsalFixture):
    def test_default_catalog_and_dashboard(self):
        dashboard = self.service.dashboard()

        self.assertEqual(dashboard["summary"]["scenario_count"], 1)
        self.assertEqual(len(dashboard["templates"]), 2)
        self.assertEqual(
            {item["customer_type"] for item in dashboard["templates"]},
            {"PUBLIC", "ENTERPRISE"},
        )
        self.assertEqual(RehearsalComplaintTemplate.objects.count(), 2)
        self.assertEqual(len(dashboard["process_templates"]), 2)
        self.assertEqual(RehearsalProcessTemplate.objects.count(), 2)
        self.assertEqual(len(dashboard["data_templates"]), 2)
        self.assertEqual(RehearsalDataTemplate.objects.count(), 2)
        self.assertEqual(len(dashboard["scene_scripts"]), 1)
        self.assertEqual(RehearsalSceneScript.objects.count(), 1)
        self.assertEqual(len(dashboard["scene_libraries"]), 1)
        self.assertEqual(dashboard["summary"]["scene_library_count"], 1)
        self.assertEqual(dashboard["summary"]["scene_library_ready_count"], 1)
        self.assertEqual(RehearsalSceneLibrary.objects.count(), 1)
        self.assertEqual(RehearsalSceneLibraryEvent.objects.count(), 1)
        self.assertEqual(len(dashboard["alert_templates"]), 1)
        self.assertEqual(RehearsalAlertTemplate.objects.count(), 1)
        self.assertEqual(dashboard["summary"]["fault_configuration_count"], 1)
        self.assertEqual(dashboard["summary"]["fault_configuration_ready_count"], 1)
        self.assertEqual(len(dashboard["fault_configurations"]), 1)
        self.assertEqual(RehearsalFaultConfiguration.objects.count(), 1)

    def test_generate_idempotency_and_instance_lifecycle(self):
        instance = self.create_instance()
        template = next(item for item in self.templates if item.customer_type == "PUBLIC")
        payload = {
            "request_id": "batch:generated:001",
            "template_id": str(template.id),
            "count": 6,
            "start_at": timezone.now().isoformat(),
            "interval_minutes": 10,
        }

        generated = self.service.generate_complaints(instance["id"], payload)
        duplicate = self.service.generate_complaints(instance["id"], payload)
        prepared = self.service.operate_instance(instance["id"], {"action": "PREPARE"})
        started = self.service.operate_instance(instance["id"], {"action": "START"})
        completed = self.service.operate_instance(instance["id"], {"action": "COMPLETE"})

        self.assertTrue(generated["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(RehearsalComplaint.objects.count(), 6)
        self.assertEqual(prepared["instance"]["status"], RehearsalInstanceStatus.READY)
        self.assertEqual(started["instance"]["status"], RehearsalInstanceStatus.RUNNING)
        self.assertEqual(completed["instance"]["status"], RehearsalInstanceStatus.COMPLETED)
        self.assertEqual(RehearsalInstanceEvent.objects.filter(instance_id=instance["id"]).count(), 5)

    def test_import_validation_filters_and_exports(self):
        instance = self.create_instance("import")
        content = (
            "occurred_at,customer_type,customer_ref,customer_name,contact,customer_segment,region,channel,category,severity,title,content\n"
            "2026-09-14T09:00:00+08:00,PUBLIC,PUB-001,模拟客户,138****8000,融合套餐,东区,10000号,宽带中断,HIGH,有效模拟申告,请核查\n"
            "bad-time,ENTERPRISE,ENT-001,模拟企业,13800138001,政企专线,南区,客户经理,专线质量,UNKNOWN,无效模拟申告,请核查\n"
        )
        uploaded = SimpleUploadedFile("complaints.csv", content.encode("utf-8"), content_type="text/csv")

        result = self.service.import_complaints(instance["id"], "batch:import:001", uploaded)
        batch = RehearsalComplaintBatch.objects.get(id=result["batch"]["id"])
        filtered = self.service.list_complaints(
            {"instance_id": instance["id"], "validation_status": ComplaintValidationStatus.VALID}
        )
        validation_csv = build_batch_validation_csv(batch)
        complaint_csv = build_complaint_csv(self.service.complaint_queryset({"instance_id": instance["id"]}))

        self.assertEqual(batch.status, ComplaintBatchStatus.PARTIAL)
        self.assertEqual(batch.valid_count, 1)
        self.assertEqual(batch.invalid_count, 1)
        self.assertEqual(filtered["total"], 1)
        self.assertEqual(RehearsalComplaint.objects.get(sequence_no=2).contact_masked, "138****8001")
        self.assertTrue(build_import_template_csv().startswith(b"\xef\xbb\xbf"))
        self.assertIn("校验错误".encode(), validation_csv)
        self.assertIn("有效模拟申告".encode(), complaint_csv)

    def test_prepare_requires_valid_complaint(self):
        instance = self.create_instance("empty")
        with self.assertRaises(AppApiException):
            self.service.operate_instance(instance["id"], {"action": "PREPARE"})

    def test_process_manual_lifecycle_block_resume_skip_and_export(self):
        instance = self.create_ready_instance()
        template = RehearsalProcessTemplate.objects.get(code="PUBLIC-DISPOSAL-V1")
        payload = {
            "request_id": "process:manual:001",
            "template_id": str(template.id),
            "name": "公众申告人工处置演练",
            "auto_advance": False,
        }

        created = self.service.create_process_run(instance["id"], payload)
        duplicate = self.service.create_process_run(instance["id"], payload)
        run_id = created["run"]["id"]
        self.service.operate_process_run(run_id, {"action": "START"})
        self.service.operate_process_run(run_id, {"action": "COMPLETE_STEP", "note": "受理完成"})
        blocked = self.service.operate_process_run(run_id, {"action": "BLOCK", "reason": "等待线路资源"})
        resumed = self.service.operate_process_run(run_id, {"action": "RESUME", "note": "资源已到位"})
        self.service.operate_process_run(run_id, {"action": "SKIP", "note": "已有检测结论"})
        self.service.operate_process_run(run_id, {"action": "COMPLETE_STEP"})
        completed = self.service.operate_process_run(run_id, {"action": "COMPLETE_STEP"})
        run = RehearsalProcessRun.objects.get(id=run_id)
        csv_content = build_process_run_csv(self.service.process_run_queryset({"status": "COMPLETED"}))

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(blocked["run"]["status"], ProcessRunStatus.BLOCKED)
        self.assertEqual(resumed["run"]["status"], ProcessRunStatus.RUNNING)
        self.assertEqual(completed["run"]["status"], ProcessRunStatus.COMPLETED)
        self.assertEqual(run.progress_percent, 100)
        self.assertEqual(run.completed_steps, 4)
        self.assertEqual(
            RehearsalProcessStep.objects.filter(run=run, status=ProcessStepStatus.SKIPPED).count(),
            1,
        )
        self.assertEqual(RehearsalProcessEvent.objects.filter(run=run).count(), 8)
        self.assertIn("公众申告人工处置演练".encode(), csv_content)
        self.assertIn("已全部完成".encode(), csv_content)

    def test_process_auto_advance_and_custom_template_validation(self):
        instance = self.create_ready_instance("auto")
        template = RehearsalProcessTemplate.objects.get(code="ENTERPRISE-DISPOSAL-V1")
        created = self.service.create_process_run(
            instance["id"],
            {
                "request_id": "process:auto:001",
                "template_id": str(template.id),
            },
        )
        run_id = created["run"]["id"]
        self.service.operate_process_run(run_id, {"action": "START"})
        for _ in range(5):
            result = self.service.operate_process_run(run_id, {"action": "AUTO_ADVANCE"})

        self.assertEqual(result["run"]["status"], ProcessRunStatus.COMPLETED)
        self.assertEqual(result["run"]["progress_percent"], 100)
        self.assertEqual(self.service.list_process_runs({"status": "COMPLETED"})["total"], 1)
        with self.assertRaises(ValidationError):
            self.service.create_process_template(
                {
                    "code": "BROKEN",
                    "name": "重复步骤模板",
                    "customer_scope": "MIXED",
                    "steps": [{"code": "A", "name": "步骤一"}, {"code": "A", "name": "步骤二"}],
                }
            )

    def test_data_simulation_generate_summary_filters_and_export(self):
        instance = self.create_instance("data")
        template = RehearsalDataTemplate.objects.get(code="NETWORK-METRIC-V1")
        payload = {
            "request_id": "data:metric:001",
            "template_id": str(template.id),
            "name": "网络质量指标模拟",
            "count": 12,
            "start_at": timezone.now().isoformat(),
            "interval_minutes": 10,
            "tags": ["验收", "网络质量"],
        }

        created = self.service.generate_simulation_data(instance["id"], payload)
        duplicate = self.service.generate_simulation_data(instance["id"], payload)
        batch = RehearsalDataBatch.objects.get(id=created["batch"]["id"])
        detail = self.service.data_batch_detail(batch.id)
        filtered = self.service.list_data_batches({"kind": DataSimulationKind.METRIC, "tag": "验收"})
        selected = self.service.list_data_batches({"batch_id": str(batch.id)})
        csv_content = build_data_simulation_csv(
            self.service.data_batch_queryset({"instance_id": instance["id"]}, include_points=True)
        )
        instance_detail = self.service.instance_detail(instance["id"])

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(RehearsalDataPoint.objects.filter(batch=batch).count(), 12)
        self.assertEqual(detail["point_count"], 12)
        self.assertEqual(len(detail["points"]), 12)
        self.assertEqual(set(detail["summary"]["fields"]), {"NETWORK_LOAD", "LATENCY_MS", "ALARM_COUNT"})
        self.assertEqual(filtered["total"], 1)
        self.assertEqual(selected["items"][0]["id"], str(batch.id))
        self.assertEqual(instance_detail["data_batch_count"], 1)
        self.assertEqual(len(instance_detail["data_batches"]), 1)
        self.assertTrue(csv_content.startswith(b"\xef\xbb\xbf"))
        self.assertIn("网络负载".encode(), csv_content)
        self.assertEqual(RehearsalInstanceEvent.objects.filter(instance_id=instance["id"]).count(), 2)

    def test_data_template_validation_and_terminal_instance_guard(self):
        with self.assertRaises(ValidationError):
            self.service.create_data_template(
                {
                    "code": "BROKEN-DATA",
                    "name": "错误数据模板",
                    "kind": "METRIC",
                    "default_rules": [
                        {
                            "code": "VALUE",
                            "name": "错误范围",
                            "min_value": 10,
                            "max_value": 1,
                        }
                    ],
                }
            )
        instance = self.create_instance("terminal-data")
        self.service.operate_instance(instance["id"], {"action": "CANCEL"})
        template = RehearsalDataTemplate.objects.get(code="SENTIMENT-TREND-V1")
        with self.assertRaises(AppApiException):
            self.service.generate_simulation_data(
                instance["id"],
                {
                    "request_id": "data:terminal:001",
                    "template_id": str(template.id),
                    "count": 6,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
            )

    def test_scene_run_execution_summary_and_idempotency(self):
        instance = self.create_ready_instance("scene")
        script = RehearsalSceneScript.objects.get(code="NETWORK-FAULT-SCENE-V1")
        payload = {
            "request_id": "scene:run:001",
            "script_id": str(script.id),
            "name": "网络故障全链路预演",
            "complaint_batch_ids": [instance["batches"][0]["id"]],
        }

        created = self.service.create_scene_run(instance["id"], payload)
        duplicate = self.service.create_scene_run(instance["id"], payload)
        run_id = created["run"]["id"]
        self.service.operate_scene_run(run_id, {"action": "START"})
        for index in range(5):
            result = self.service.operate_scene_run(
                run_id,
                {
                    "action": "COMPLETE_STEP",
                    "note": f"步骤 {index + 1} 验收通过",
                    "output": {"accepted": True, "sequence": index + 1},
                },
            )

        run = RehearsalSceneRun.objects.get(id=run_id)
        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(result["run"]["status"], SceneRunStatus.COMPLETED)
        self.assertEqual(run.progress_percent, 100)
        self.assertEqual(run.completed_steps, 5)
        self.assertIn("conclusion", run.summary)
        self.assertEqual(
            RehearsalSceneStep.objects.filter(run=run, status=SceneStepStatus.COMPLETED).count(),
            5,
        )
        self.assertEqual(RehearsalSceneEvent.objects.filter(run=run).count(), 7)
        self.assertEqual(self.service.list_scene_runs({"status": SceneRunStatus.COMPLETED})["total"], 1)
        self.assertEqual(self.service.instance_detail(instance["id"])["scene_run_count"], 1)

    def test_failed_scene_manual_work_order_closed_loop(self):
        instance = self.create_ready_instance("scene-order")
        script = RehearsalSceneScript.objects.get(code="NETWORK-FAULT-SCENE-V1")
        scene = self.service.create_scene_run(
            instance["id"],
            {
                "request_id": "scene:failure:001",
                "script_id": str(script.id),
                "complaint_batch_ids": [instance["batches"][0]["id"]],
            },
        )["run"]
        self.service.operate_scene_run(scene["id"], {"action": "START"})
        failed = self.service.operate_scene_run(
            scene["id"],
            {"action": "FAIL", "reason": "恢复验证未达到阈值", "output": {"latency_ms": 132}},
        )["run"]
        created = self.service.create_work_order(
            scene["id"],
            {
                "request_id": "work-order:001",
                "title": "跟踪网络恢复指标",
                "description": "预演失败后转人工继续处置。",
                "priority": "HIGH",
                "assignee": "网络运维组",
            },
        )
        duplicate = self.service.create_work_order(
            scene["id"],
            {
                "request_id": "work-order:001",
                "title": "重复提交不会新建",
            },
        )
        order_id = created["work_order"]["id"]
        self.service.operate_work_order(order_id, {"action": "ACCEPT"})
        self.service.operate_work_order(
            order_id,
            {"action": "RESOLVE", "resolution": "扩容后指标已恢复"},
        )
        closed = self.service.operate_work_order(order_id, {"action": "CLOSE"})["work_order"]

        self.assertEqual(failed["status"], SceneRunStatus.FAILED)
        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(closed["status"], WorkOrderStatus.CLOSED)
        self.assertEqual(RehearsalWorkOrder.objects.count(), 1)
        self.assertEqual(RehearsalWorkOrderEvent.objects.filter(work_order_id=order_id).count(), 4)
        self.assertEqual(self.service.list_work_orders({"status": WorkOrderStatus.CLOSED})["total"], 1)

    def test_alert_report_acknowledge_recover_and_timeline(self):
        instance = self.create_ready_instance("alert")
        template = RehearsalAlertTemplate.objects.get(code="NETWORK-ALERT-V1")
        occurred_at = timezone.now()
        payload = {
            "request_id": "alert:network:001",
            "template_id": str(template.id),
            "rule_code": "LATENCY_HIGH",
            "title": "核心链路时延超阈值",
            "resource_id": "NE-X21-001",
            "resource_name": "核心路由器-X21",
            "metric_value": 136,
            "occurred_at": occurred_at.isoformat(),
        }

        created = self.service.create_alert(instance["id"], payload)
        duplicate = self.service.create_alert(instance["id"], payload)
        alert_id = created["alert"]["id"]
        reported = self.service.operate_alert(
            alert_id,
            {"action": "REPORT", "event_at": (occurred_at + timedelta(minutes=1)).isoformat()},
        )["alert"]
        acknowledged = self.service.operate_alert(
            alert_id,
            {"action": "ACKNOWLEDGE", "event_at": (occurred_at + timedelta(minutes=3)).isoformat()},
        )["alert"]
        recovered = self.service.operate_alert(
            alert_id,
            {
                "action": "RECOVER",
                "event_at": (occurred_at + timedelta(minutes=12)).isoformat(),
                "metric_value": 72,
                "note": "链路调整后恢复",
            },
        )["alert"]
        timeline = self.service.instance_timeline(instance["id"])

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(reported["status"], AlertStatus.ACTIVE)
        self.assertEqual(acknowledged["status"], AlertStatus.ACKNOWLEDGED)
        self.assertEqual(recovered["status"], AlertStatus.RECOVERED)
        self.assertEqual(recovered["metric_value"], 72)
        self.assertEqual(RehearsalAlert.objects.count(), 1)
        self.assertEqual(RehearsalAlertEvent.objects.filter(alert_id=alert_id).count(), 4)
        self.assertEqual(self.service.list_alerts({"status": AlertStatus.RECOVERED})["total"], 1)
        self.assertEqual(timeline["kind_counts"]["ALERT"], 4)
        self.assertEqual(timeline["kind_counts"]["COMPLAINT"], 3)

    def test_alert_rejects_cross_instance_injection_and_invalid_transition(self):
        first = self.create_ready_instance("alert-first")
        second = self.create_ready_instance("alert-second")
        template = RehearsalAlertTemplate.objects.get(code="NETWORK-ALERT-V1")
        data_template = RehearsalDataTemplate.objects.get(code="NETWORK-METRIC-V1")
        foreign_batch = self.service.generate_simulation_data(
            second["id"],
            {
                "request_id": "alert:data:foreign",
                "template_id": str(data_template.id),
                "count": 6,
                "start_at": timezone.now().isoformat(),
                "interval_minutes": 5,
            },
        )["batch"]
        with self.assertRaises(AppApiException):
            self.service.create_alert(
                first["id"],
                {
                    "request_id": "alert:invalid:foreign",
                    "template_id": str(template.id),
                    "data_batch_id": foreign_batch["id"],
                },
            )
        alert = self.service.create_alert(
            first["id"],
            {
                "request_id": "alert:invalid:transition",
                "template_id": str(template.id),
            },
        )["alert"]
        with self.assertRaises(AppApiException):
            self.service.operate_alert(alert["id"], {"action": "RECOVER"})

    def test_data_model_build_load_versions_and_lineage(self):
        instance = self.create_ready_instance("data-model")
        data_template = RehearsalDataTemplate.objects.get(code="NETWORK-METRIC-V1")
        data_batch = self.service.generate_simulation_data(
            instance["id"],
            {
                "request_id": "dm:data:001",
                "template_id": str(data_template.id),
                "count": 8,
                "start_at": timezone.now().isoformat(),
                "interval_minutes": 5,
                "tags": ["数据模型"],
            },
        )["batch"]
        process_template = RehearsalProcessTemplate.objects.get(code="PUBLIC-DISPOSAL-V1")
        process_run = self.service.create_process_run(
            instance["id"],
            {
                "request_id": "dm:process:001",
                "template_id": str(process_template.id),
            },
        )["run"]
        script = RehearsalSceneScript.objects.get(code="NETWORK-FAULT-SCENE-V1")
        scene = self.service.create_scene_run(
            instance["id"],
            {
                "request_id": "dm:scene:001",
                "script_id": str(script.id),
                "complaint_batch_ids": [instance["batches"][0]["id"]],
                "data_batch_ids": [data_batch["id"]],
                "process_run_ids": [process_run["id"]],
            },
        )["run"]
        alert_template = RehearsalAlertTemplate.objects.get(code="NETWORK-ALERT-V1")
        self.service.create_alert(
            instance["id"],
            {
                "request_id": "dm:alert:001",
                "template_id": str(alert_template.id),
                "rule_code": "LATENCY_HIGH",
                "scene_run_id": scene["id"],
                "data_batch_id": data_batch["id"],
                "metric_value": 128,
            },
        )

        created = self.service.create_data_model(
            instance["id"],
            {
                "request_id": "dm:model:001",
                "name": "故障预演统一数据模型",
                "build_strategy": DataModelBuildStrategy.SCENE_LINKED,
                "transformer_code": "CANONICAL_V1",
                "scene_run_id": scene["id"],
                "source_scope": {"mode": "acceptance"},
            },
        )
        duplicate = self.service.create_data_model(
            instance["id"],
            {
                "request_id": "dm:model:001",
                "build_strategy": DataModelBuildStrategy.SCENE_LINKED,
                "scene_run_id": scene["id"],
            },
        )
        reloaded = self.service.load_data_model(
            created["data_model"]["id"],
            {
                "request_id": "dm:model:load:002",
                "note": "刷新验收版本",
            },
        )
        duplicate_load = self.service.load_data_model(
            created["data_model"]["id"],
            {
                "request_id": "dm:model:load:002",
            },
        )
        detail = self.service.data_model_detail(created["data_model"]["id"])
        page = self.service.list_data_models({"quality_status": DataModelQualityStatus.PASS})

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertTrue(reloaded["created"])
        self.assertFalse(duplicate_load["created"])
        self.assertEqual(RehearsalDataModel.objects.count(), 1)
        self.assertEqual(RehearsalDataModelVersion.objects.count(), 2)
        self.assertEqual(RehearsalDataModelLoad.objects.count(), 2)
        self.assertEqual(detail["current_version"], 2)
        self.assertEqual(detail["quality_status"], DataModelQualityStatus.PASS)
        self.assertEqual(detail["summary"]["source_counts"]["data_points"], 8)
        self.assertGreaterEqual(len(detail["lineage"]["sources"]), 5)
        self.assertEqual(page["total"], 1)

    def test_traceback_snapshot_create_filter_detail_and_export(self):
        instance = self.create_ready_instance("traceback")
        data_template = RehearsalDataTemplate.objects.get(code="NETWORK-METRIC-V1")
        data_batch = self.service.generate_simulation_data(
            instance["id"],
            {
                "request_id": "traceback:data:001",
                "template_id": str(data_template.id),
                "name": "回溯链路指标",
                "count": 6,
                "start_at": timezone.now().isoformat(),
                "interval_minutes": 5,
                "tags": ["X25", "回溯"],
            },
        )["batch"]
        data_model = self.service.create_data_model(
            instance["id"],
            {
                "request_id": "traceback:model:001",
                "name": "回溯统一数据模型",
                "build_strategy": DataModelBuildStrategy.FULL_INSTANCE,
                "source_scope": {"batch_id": data_batch["id"]},
            },
        )["data_model"]
        payload = {
            "request_id": "traceback:create:001",
            "name": "X25 网络故障动态回溯",
            "description": "验收客户申告与指标的统一回放快照",
            "scope": TracebackScope.FULL_INSTANCE,
            "data_model_id": data_model["id"],
            "kind_scope": ["COMPLAINT", "METRIC"],
        }

        created = self.service.create_traceback(instance["id"], payload)
        duplicate = self.service.create_traceback(instance["id"], payload)
        traceback_id = created["traceback"]["id"]
        page = self.service.list_tracebacks(
            {
                "instance_id": instance["id"],
                "status": TracebackStatus.READY,
                "kind": "METRIC",
                "keyword": "X25",
            }
        )
        detail = self.service.traceback_detail(traceback_id)
        instance_detail = self.service.instance_detail(instance["id"])
        csv_content = build_traceback_csv(self.service.traceback_queryset({"keyword": "X25"}))

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(RehearsalTracebackSnapshot.objects.count(), 1)
        self.assertEqual(created["traceback"]["status"], TracebackStatus.READY)
        self.assertEqual(created["traceback"]["total_events"], 9)
        self.assertEqual(created["traceback"]["kind_counts"]["COMPLAINT"], 3)
        self.assertEqual(created["traceback"]["kind_counts"]["METRIC"], 6)
        self.assertEqual(page["total"], 1)
        self.assertEqual(len(detail["nodes"]), 2)
        self.assertGreaterEqual(len(detail["transitions"]), 1)
        self.assertGreaterEqual(len(detail["deltas"]), 1)
        self.assertEqual(instance_detail["traceback_count"], 1)
        self.assertEqual(len(instance_detail["tracebacks"]), 1)
        self.assertTrue(csv_content.startswith(b"\xef\xbb\xbf"))
        self.assertIn("X25 网络故障动态回溯".encode(), csv_content)

    def test_simulation_record_query_detail_and_export(self):
        instance = self.create_ready_instance("record")
        data_template = RehearsalDataTemplate.objects.get(code="NETWORK-METRIC-V1")
        process_template = RehearsalProcessTemplate.objects.get(code="PUBLIC-DISPOSAL-V1")
        scene_script = RehearsalSceneScript.objects.get(code="NETWORK-FAULT-SCENE-V1")
        alert_template = RehearsalAlertTemplate.objects.get(code="NETWORK-ALERT-V1")
        data_batch = self.service.generate_simulation_data(
            instance["id"],
            {
                "request_id": "record:data:001",
                "template_id": str(data_template.id),
                "name": "X26 查询指标批次",
                "count": 5,
                "start_at": timezone.now().isoformat(),
                "interval_minutes": 5,
            },
        )["batch"]
        process_run = self.service.create_process_run(
            instance["id"],
            {
                "request_id": "record:process:001",
                "template_id": str(process_template.id),
                "name": "X26 查询处置流程",
            },
        )["run"]
        scene_run = self.service.create_scene_run(
            instance["id"],
            {
                "request_id": "record:scene:001",
                "script_id": str(scene_script.id),
                "name": "X26 查询场景执行",
                "data_batch_ids": [data_batch["id"]],
                "process_run_ids": [process_run["id"]],
            },
        )["run"]
        self.service.create_alert(
            instance["id"],
            {
                "request_id": "record:alert:001",
                "template_id": str(alert_template.id),
                "rule_code": "LATENCY_HIGH",
                "scene_run_id": scene_run["id"],
                "data_batch_id": data_batch["id"],
                "title": "X26 查询模拟告警",
                "metric_value": 132,
            },
        )
        self.service.create_traceback(
            instance["id"],
            {
                "request_id": "record:traceback:001",
                "name": "X26 查询回溯快照",
                "scope": TracebackScope.FULL_INSTANCE,
            },
        )

        page = self.service.list_simulation_records({"instance_id": instance["id"], "page_size": 50})
        kinds = {item["kind"] for item in page["items"]}
        process_detail = self.service.simulation_record_detail("PROCESS_RUN", process_run["id"])
        filtered = self.service.list_simulation_records(
            {"instance_id": instance["id"], "kind": "DATA_BATCH", "state": "SUCCEEDED", "keyword": "X26"}
        )
        export_content = build_simulation_record_csv(self.service.simulation_records({"instance_id": instance["id"]}))

        self.assertEqual(
            kinds,
            {"INSTANCE", "COMPLAINT_BATCH", "DATA_BATCH", "PROCESS_RUN", "SCENE_RUN", "ALERT", "TRACEBACK"},
        )
        self.assertEqual(page["summary"]["total"], 7)
        self.assertEqual(page["summary"]["kind_counts"]["TRACEBACK"], 1)
        self.assertEqual(filtered["total"], 1)
        self.assertEqual(filtered["items"][0]["occurrence_no"], 1)
        self.assertEqual(process_detail["record"]["resource_id"], process_run["id"])
        self.assertGreaterEqual(process_detail["activity_count"], 1)
        self.assertTrue(export_content.startswith(b"\xef\xbb\xbf"))
        self.assertIn("X26 查询场景执行".encode(), export_content)

    def test_capability_mock_invoke_history_replay_and_warroom(self):
        instance = self.create_ready_instance("capability")
        mocks = self.service.list_capability_mocks()
        api_mock = next(item for item in mocks if item["kind"] == CapabilityMockKind.API)
        atomic_mock = next(item for item in mocks if item["kind"] == CapabilityMockKind.ATOMIC)

        succeeded = self.service.invoke_capability(
            instance["id"],
            {
                "request_id": "capability:invoke:001",
                "capability_id": api_mock["id"],
                "request_payload": {"resource_id": "NE-X27-001", "include_metrics": True},
                "warroom_group": "网络质量战情室",
            },
        )
        duplicate = self.service.invoke_capability(
            instance["id"],
            {
                "request_id": "capability:invoke:001",
                "capability_id": api_mock["id"],
                "request_payload": {"resource_id": "IGNORED"},
            },
        )
        failed = self.service.invoke_capability(
            instance["id"],
            {
                "request_id": "capability:invoke:002",
                "capability_id": atomic_mock["id"],
                "request_payload": {"node_id": "NODE-X27-01", "strategy": "DRAIN_AND_BYPASS"},
                "expected_status": CapabilityCallStatus.FAILED,
                "latency_ms": 360,
                "warroom_group": "网络质量战情室",
            },
        )
        replay = self.service.replay_capability_call(
            succeeded["call"]["id"],
            {
                "request_id": "capability:replay:001",
                "expected_status": CapabilityCallStatus.TIMEOUT,
                "latency_ms": 3000,
            },
        )
        page = self.service.list_capability_calls(
            {"instance_id": instance["id"], "warroom_group": "网络质量战情室", "page_size": 10}
        )
        detail = self.service.capability_call_detail(replay["call"]["id"])
        record_page = self.service.list_simulation_records(
            {"instance_id": instance["id"], "kind": "CAPABILITY_CALL", "page_size": 10}
        )
        instance_detail = self.service.instance_detail(instance["id"])

        self.assertTrue(succeeded["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(succeeded["call"]["http_status"], 200)
        self.assertEqual(failed["call"]["status"], CapabilityCallStatus.FAILED)
        self.assertEqual(failed["call"]["http_status"], 500)
        self.assertEqual(replay["call"]["source"], "REPLAY")
        self.assertEqual(replay["call"]["http_status"], 504)
        self.assertEqual(detail["replay_of_id"], succeeded["call"]["id"])
        self.assertEqual(page["summary"], {"total": 3, "succeeded": 1, "failed": 1, "timeout": 1, "average_latency_ms": 1149})
        self.assertEqual(page["warrooms"][0]["attention"], 2)
        self.assertEqual(record_page["total"], 3)
        self.assertEqual(record_page["summary"]["kind_counts"]["CAPABILITY_CALL"], 3)
        self.assertEqual(instance_detail["capability_call_count"], 3)
        self.assertEqual(len(instance_detail["capability_calls"]), 3)
        self.assertEqual(RehearsalCapabilityCall.objects.count(), 3)
        self.assertEqual(RehearsalCapabilityMock.objects.count(), 2)

    def test_fault_configuration_version_validation_apply_and_audit(self):
        instance = self.create_ready_instance("fault-config")
        default_configuration = self.service.list_fault_configurations({"status": "READY"})["items"][0]

        cloned = self.service.clone_fault_configuration(
            default_configuration["id"],
            {
                "request_id": "fault-config:clone:001",
                "name": "核心网络节点故障模拟配置 V2",
                "note": "服务层版本化验收",
            },
        )
        duplicate_clone = self.service.clone_fault_configuration(
            default_configuration["id"],
            {"request_id": "fault-config:clone:001"},
        )
        invalid = self.service.create_fault_configuration(
            {
                "request_id": "fault-config:invalid:001",
                "scenario_id": str(self.scenario.id),
                "code": "INVALID-FAULT-CONFIG",
                "name": "待完善故障配置",
                "alert_rules": [],
                "timeout_policy": {},
                "active_window": {},
                "topology_snapshot": {},
                "emergency_plan": {},
            }
        )
        applied = self.service.apply_fault_configuration(
            instance["id"],
            {
                "request_id": "fault-config:apply:001",
                "configuration_id": cloned["configuration"]["id"],
                "note": "绑定到 X28 验收实例",
            },
        )
        duplicate_apply = self.service.apply_fault_configuration(
            instance["id"],
            {
                "request_id": "fault-config:apply:001",
                "configuration_id": cloned["configuration"]["id"],
            },
        )
        detail = self.service.fault_configuration_detail(cloned["configuration"]["id"])
        bound_page = self.service.list_fault_configurations({"bound": "true", "page_size": 10})
        archived = self.service.archive_fault_configuration(
            cloned["configuration"]["id"],
            {"note": "保留已应用快照后归档版本"},
        )

        self.assertTrue(cloned["created"])
        self.assertFalse(duplicate_clone["created"])
        self.assertEqual(cloned["configuration"]["version"], default_configuration["version"] + 1)
        self.assertEqual(cloned["configuration"]["status"], FaultConfigurationStatus.READY)
        self.assertTrue(cloned["configuration"]["validation_report"]["passed"])
        self.assertEqual(invalid["configuration"]["status"], FaultConfigurationStatus.DRAFT)
        self.assertEqual(invalid["configuration"]["validation_report"]["error_count"], 5)
        self.assertTrue(applied["created"])
        self.assertFalse(duplicate_apply["created"])
        self.assertEqual(applied["instance"]["fault_configuration_id"], cloned["configuration"]["id"])
        self.assertEqual(applied["instance"]["fault_configuration_snapshot"]["version"], 2)
        self.assertEqual(bound_page["summary"]["bound_instances"], 1)
        self.assertEqual(detail["bound_instance_count"], 1)
        self.assertGreaterEqual(detail["event_count"], 3)
        self.assertEqual(archived["configuration"]["status"], FaultConfigurationStatus.ARCHIVED)
        self.assertTrue(
            RehearsalFaultConfigurationEvent.objects.filter(
                configuration_id=cloned["configuration"]["id"],
                action=FaultConfigurationAction.APPLIED,
            ).exists()
        )
        self.assertTrue(
            RehearsalInstanceEvent.objects.filter(
                instance_id=instance["id"],
                event_type="FAULT_CONFIG_APPLIED",
            ).exists()
        )

    def test_scene_replay_traffic_topology_nodes_mirror_and_idempotency(self):
        instance = self.create_ready_instance("scene-replay")
        configuration = self.service.list_fault_configurations({"status": "READY"})["items"][0]
        self.service.apply_fault_configuration(
            instance["id"],
            {
                "request_id": "scene-replay:config:001",
                "configuration_id": configuration["id"],
            },
        )
        capability = next(item for item in self.service.list_capability_mocks() if item["kind"] == "ATOMIC")
        self.service.invoke_capability(
            instance["id"],
            {
                "request_id": "scene-replay:call:001",
                "capability_id": capability["id"],
                "request_payload": {"source_node": "GW-01", "node_id": "CORE-01", "strategy": "DRAIN"},
                "expected_status": CapabilityCallStatus.FAILED,
                "latency_ms": 640,
            },
        )
        payload = {
            "request_id": "scene-replay:create:001",
            "name": "X29 核心链路场景动态回溯",
            "description": "验收流量帧、节点轨迹和拓扑前后镜像",
        }

        created = self.service.create_scene_replay(instance["id"], payload)
        duplicate = self.service.create_scene_replay(instance["id"], payload)
        replay_id = created["replay"]["id"]
        page = self.service.list_scene_replays(
            {"instance_id": instance["id"], "status": SceneReplayStatus.READY, "node_id": "CORE-01"}
        )
        detail = self.service.scene_replay_detail(replay_id)
        instance_detail = self.service.instance_detail(instance["id"])

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(RehearsalSceneReplay.objects.count(), 1)
        self.assertEqual(created["replay"]["status"], SceneReplayStatus.READY)
        self.assertEqual(created["replay"]["traffic_frame_count"], 1)
        self.assertEqual(page["total"], 1)
        self.assertEqual(page["summary"]["traffic_frames"], 1)
        self.assertEqual(len(detail["checksum"]), 64)
        self.assertEqual(detail["topology_before"]["version"], "TOPOLOGY-CORE-2026.09")
        core_trace = next(item for item in detail["node_traces"] if item["node_id"] == "CORE-01")
        self.assertEqual(core_trace["after_state"], "DEGRADED")
        self.assertTrue(core_trace["changed"])
        self.assertGreaterEqual(len(detail["mirror_comparison"]["node_deltas"]), 1)
        self.assertEqual(instance_detail["scene_replay_count"], 1)
        self.assertEqual(len(instance_detail["scene_replays"]), 1)

    def test_fault_report_template_lifecycle_countdown_attachments_and_history(self):
        instance = self.create_ready_instance("fault-report")
        configuration = self.service.list_fault_configurations({"status": "READY"})["items"][0]
        self.service.apply_fault_configuration(
            instance["id"],
            {
                "request_id": "fault-report:config:001",
                "configuration_id": configuration["id"],
            },
        )
        default_template = self.service.list_fault_report_templates({"status": "READY"})[0]
        template_version = self.service.create_fault_report_template(
            {
                "request_id": "fault-report:template:002",
                "scenario_id": str(self.scenario.id),
                "code": default_template["code"],
                "name": "核心网络故障分阶段上报模板 V2",
                "default_level": FaultReportLevel.LEVEL_3,
                "stage_definitions": [
                    {
                        "code": "INITIAL_REPORT",
                        "name": "首次上报",
                        "duration_minutes": 8,
                        "owner_role": "监控席",
                    }
                ],
                "required_fields": ["title", "description"],
            }
        )
        payload = {
            "request_id": "fault-report:create:001",
            "template_id": default_template["id"],
            "title": "核心节点故障上报",
            "description": "核心节点出现持续高时延并触发隔离预案。",
            "impact_scope": "东区宽带与政企专线模拟流量",
            "level": FaultReportLevel.LEVEL_3,
            "attachments": [
                {
                    "reference_id": "evidence:alarm:001",
                    "file_name": "alarm-snapshot.json",
                    "content_type": "application/json",
                    "file_size": 512,
                }
            ],
        }
        created = self.service.create_fault_report(instance["id"], payload)
        duplicate = self.service.create_fault_report(instance["id"], payload)
        report_id = created["report"]["id"]
        submitted = self.service.operate_fault_report(report_id, {"action": FaultReportAction.SUBMITTED})
        accepted = self.service.operate_fault_report(report_id, {"action": FaultReportAction.ACCEPTED})
        advanced = self.service.operate_fault_report(
            report_id,
            {"action": FaultReportAction.STAGE_ADVANCED, "note": "首次上报信息已确认"},
        )
        escalated = self.service.operate_fault_report(
            report_id,
            {
                "action": FaultReportAction.ESCALATED,
                "level": FaultReportLevel.LEVEL_2,
                "note": "影响范围扩大，升级故障等级",
            },
        )
        attached = self.service.operate_fault_report(
            report_id,
            {
                "action": FaultReportAction.ATTACHMENT_ADDED,
                "note": "补充处置记录",
                "attachments": [
                    {
                        "reference_id": "evidence:disposal:001",
                        "file_name": "disposal.log",
                        "content_type": "text/plain",
                        "file_size": 1024,
                    }
                ],
            },
        )
        recovered = self.service.operate_fault_report(
            report_id,
            {"action": FaultReportAction.RECOVERED, "note": "备用链路切换后业务恢复"},
        )
        closed = self.service.operate_fault_report(
            report_id,
            {"action": FaultReportAction.CLOSED, "note": "恢复验证完成，关闭故障上报"},
        )
        page = self.service.list_fault_reports(
            {"instance_id": instance["id"], "status": FaultReportStatus.CLOSED, "keyword": "核心节点"}
        )
        record_page = self.service.list_simulation_records(
            {"instance_id": instance["id"], "kind": "FAULT_REPORT"}
        )
        record_detail = self.service.simulation_record_detail("FAULT_REPORT", report_id)
        instance_detail = self.service.instance_detail(instance["id"])

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(template_version["template"]["version"], default_template["version"] + 1)
        self.assertEqual(submitted["report"]["status"], FaultReportStatus.REPORTED)
        self.assertEqual(accepted["report"]["status"], FaultReportStatus.PROCESSING)
        self.assertEqual(advanced["report"]["current_stage_code"], "IMPACT_CONFIRM")
        self.assertEqual(escalated["report"]["level"], FaultReportLevel.LEVEL_2)
        self.assertEqual(attached["report"]["attachment_count"], 2)
        self.assertEqual(recovered["report"]["status"], FaultReportStatus.RECOVERED)
        self.assertEqual(closed["report"]["status"], FaultReportStatus.CLOSED)
        self.assertEqual(closed["report"]["progress_percent"], 100)
        self.assertEqual(page["total"], 1)
        self.assertEqual(page["summary"]["attachments"], 2)
        self.assertEqual(record_page["summary"]["kind_counts"]["FAULT_REPORT"], 1)
        self.assertEqual(record_detail["activity_count"], 8)
        self.assertEqual(instance_detail["fault_report_count"], 1)
        self.assertEqual(len(instance_detail["fault_reports"]), 1)
        self.assertEqual(RehearsalFaultReport.objects.count(), 1)
        self.assertEqual(RehearsalFaultReportEvent.objects.filter(report_id=report_id).count(), 8)
        self.assertEqual(RehearsalFaultReportTemplate.objects.count(), 2)

    def test_scene_library_create_clone_operate_and_run(self):
        instance = self.create_ready_instance("library")
        instance_id = instance["id"]
        data_template = RehearsalDataTemplate.objects.get(code="NETWORK-METRIC-V1")
        process_template = RehearsalProcessTemplate.objects.get(code="PUBLIC-DISPOSAL-V1")
        data_batch = self.service.generate_simulation_data(
            instance_id,
            {
                "request_id": "library:data:001",
                "template_id": str(data_template.id),
                "name": "场景库网络指标",
                "count": 6,
                "start_at": timezone.now().isoformat(),
                "interval_minutes": 5,
            },
        )["batch"]
        process_run = self.service.create_process_run(
            instance_id,
            {
                "request_id": "library:process:001",
                "template_id": str(process_template.id),
                "name": "场景库处置流程",
            },
        )["run"]
        data_model = self.service.create_data_model(
            instance_id,
            {
                "request_id": "library:model:001",
                "name": "场景库统一数据模型",
                "build_strategy": DataModelBuildStrategy.FULL_INSTANCE,
                "source_scope": {"batch_id": data_batch["id"]},
            },
        )["data_model"]
        script = RehearsalSceneScript.objects.get(code="NETWORK-FAULT-SCENE-V1")

        created = self.service.create_scene_library(
            {
                "request_id": "library:create:001",
                "scenario_id": str(self.scenario.id),
                "script_id": str(script.id),
                "code": "LIBRARY-NETWORK-FAULT",
                "name": "网络故障可复用场景",
                "description": "服务层验收场景库",
                "default_instance_id": instance_id,
                "process_template_id": str(process_template.id),
                "data_model_id": data_model["id"],
                "permission_roles": ["workspace_manage", "WORKSPACE_MANAGE"],
                "configuration": {"acceptance": True},
            }
        )
        duplicate = self.service.create_scene_library(
            {
                "request_id": "library:create:001",
                "scenario_id": str(self.scenario.id),
                "script_id": str(script.id),
                "code": "LIBRARY-NETWORK-FAULT",
                "name": "网络故障可复用场景",
            }
        )
        library_id = created["library"]["id"]
        cloned = self.service.clone_scene_library(
            library_id,
            {
                "request_id": "library:clone:001",
                "name": "网络故障可复用场景 V2",
                "note": "验收克隆",
                "configuration": {"acceptance": True, "version": 2},
            },
        )
        archived = self.service.operate_scene_library(cloned["library"]["id"], {"action": "ARCHIVED", "note": "暂存"})
        published = self.service.operate_scene_library(
            cloned["library"]["id"],
            {"action": "PUBLISHED", "note": "恢复可复用"},
        )
        run = self.service.create_scene_run_from_library(
            library_id,
            {
                "request_id": "library:run:001",
                "include_latest_complaint_batch": True,
                "include_latest_data_batch": True,
                "include_latest_process_run": True,
            },
        )
        detail = self.service.scene_library_detail(library_id)
        instance_detail = self.service.instance_detail(instance_id)

        self.assertTrue(created["created"])
        self.assertFalse(duplicate["created"])
        self.assertEqual(created["library"]["permission_roles"], ["WORKSPACE_MANAGE"])
        self.assertTrue(cloned["created"])
        self.assertEqual(cloned["library"]["version"], created["library"]["version"] + 1)
        self.assertEqual(archived["library"]["status"], SceneLibraryStatus.ARCHIVED)
        self.assertEqual(published["library"]["status"], SceneLibraryStatus.READY)
        self.assertTrue(run["created"])
        self.assertEqual(run["run"]["library_id"], library_id)
        self.assertEqual(run["run"]["injected_snapshot"]["process_runs"][0]["id"], process_run["id"])
        self.assertEqual(detail["run_count"], 1)
        self.assertEqual(len(detail["events"]), 2)
        self.assertEqual(len(detail["scene_runs"]), 1)
        self.assertEqual(instance_detail["scene_library_count"], 2)
        self.assertEqual(len(instance_detail["scene_libraries"]), 2)

    def test_auxiliary_simulation_is_reproducible_and_exports_three_formats(self):
        instance = self.create_instance("auxiliary")
        factors = [
            {"code": "NETWORK_JITTER", "trigger_probability": 1},
            {"code": "PACKET_LOSS", "trigger_probability": 1},
            {"code": "TRAFFIC_PEAK", "trigger_probability": 1},
        ]
        first = self.service.create_auxiliary_simulation(
            instance["id"],
            {
                "request_id": "auxiliary:reproducible:001",
                "name": "可复现随机因子组合",
                "seed": 20260917,
                "duration_minutes": 30,
                "factors": factors,
            },
        )["simulation"]
        second = self.service.create_auxiliary_simulation(
            instance["id"],
            {
                "request_id": "auxiliary:reproducible:002",
                "name": "可复现随机因子组合复跑",
                "seed": 20260917,
                "duration_minutes": 30,
                "factors": factors,
            },
        )["simulation"]
        first_row = RehearsalAuxiliarySimulation.objects.get(id=first["id"])
        first_values = [{key: value for key, value in point.items() if key != "occurred_at"} for point in first["points"]]
        second_values = [{key: value for key, value in point.items() if key != "occurred_at"} for point in second["points"]]

        self.assertEqual(first["checksum"], second["checksum"])
        self.assertEqual(first_values, second_values)
        self.assertEqual(first["factor_count"], 3)
        self.assertEqual(first["triggered_factor_count"], 3)
        self.assertIn(first["status"], AuxiliarySimulationStatus.values)
        self.assertEqual(len(first["statistics"]), 6)
        self.assertEqual(len(first["points"]), 31)
        self.assertEqual(json.loads(build_auxiliary_json(first_row, "full"))["report_level"], "full")
        workbook = load_workbook(BytesIO(build_auxiliary_xlsx(first_row, "full")), read_only=True)
        self.assertIn("分钟级原始数据", workbook.sheetnames)
        pdf = PdfReader(BytesIO(build_auxiliary_pdf(first_row, "summary")))
        self.assertGreaterEqual(len(pdf.pages), 1)
        self.assertIn("辅助模拟结果报告", "".join(page.extract_text() or "" for page in pdf.pages))


class RehearsalApiTests(RehearsalFixture):
    def setUp(self):
        super().setUp()
        self.client = APIClient()
        self.client.force_authenticate(
            user=self.user,
            token=SimpleNamespace(
                role_list=[f"WORKSPACE_MANAGE:/WORKSPACE/{self.workspace_id}"],
                permission_list=[],
            ),
        )

    @staticmethod
    def response_data(response):
        assert response.status_code == 200, response.content
        return json.loads(response.content)["data"]

    def test_api_generate_operate_filter_and_download(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API 客户申告预演",
                },
                format="json",
            )
        )["instance"]
        batch = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/complaint/generate",
                {
                    "request_id": "api:batch:001",
                    "template_id": dashboard["templates"][0]["id"],
                    "count": 4,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )["batch"]
        prepared = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/operate",
                {"action": "PREPARE"},
                format="json",
            )
        )
        process = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/process",
                {
                    "request_id": "api:process:001",
                    "template_id": dashboard["process_templates"][0]["id"],
                    "name": "API 处置进度演练",
                },
                format="json",
            )
        )["run"]
        started_process = self.response_data(
            self.client.post(
                f"{prefix}/process/{process['id']}/operate",
                {"action": "START"},
                format="json",
            )
        )["run"]
        complaints = self.response_data(
            self.client.get(f"{prefix}/complaint", {"instance_id": instance["id"], "page_size": 20})
        )
        detail = self.response_data(self.client.get(f"{prefix}/instance/{instance['id']}"))
        export_response = self.client.get(f"{prefix}/complaint/export", {"instance_id": instance["id"]})
        template_response = self.client.get(f"{prefix}/complaint/import-template")
        validation_response = self.client.get(f"{prefix}/batch/{batch['id']}/validation/export")
        process_list = self.response_data(self.client.get(f"{prefix}/process", {"status": "RUNNING"}))
        process_detail = self.response_data(self.client.get(f"{prefix}/process/{process['id']}"))
        process_export = self.client.get(f"{prefix}/process/export", {"status": "RUNNING"})
        data_batch = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/data-simulation",
                {
                    "request_id": "api:data:001",
                    "template_id": dashboard["data_templates"][0]["id"],
                    "name": "API 数据模拟",
                    "count": 8,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 15,
                    "tags": ["API验收"],
                },
                format="json",
            )
        )["batch"]
        data_list = self.response_data(self.client.get(f"{prefix}/data-simulation", {"tag": "API验收"}))
        data_detail = self.response_data(self.client.get(f"{prefix}/data-simulation/{data_batch['id']}"))
        data_export = self.client.get(f"{prefix}/data-simulation/export", {"batch_id": data_batch["id"]})
        data_instance_detail = self.response_data(self.client.get(f"{prefix}/instance/{instance['id']}"))

        self.assertEqual(prepared["instance"]["status"], RehearsalInstanceStatus.READY)
        self.assertEqual(complaints["total"], 4)
        self.assertEqual(len(detail["batches"]), 1)
        self.assertEqual(len(detail["events"]), 3)
        self.assertEqual(export_response.status_code, 200)
        self.assertEqual(export_response["Content-Type"], "text/csv; charset=utf-8")
        self.assertEqual(template_response.status_code, 200)
        self.assertEqual(validation_response.status_code, 200)
        self.assertEqual(started_process["status"], ProcessRunStatus.RUNNING)
        self.assertEqual(process_list["total"], 1)
        self.assertEqual(len(process_detail["steps"]), process["total_steps"])
        self.assertEqual(process_export.status_code, 200)
        self.assertIn("API 处置进度演练".encode(), process_export.content)
        self.assertEqual(data_list["total"], 1)
        self.assertEqual(len(data_detail["points"]), 8)
        self.assertEqual(len(data_instance_detail["data_batches"]), 1)
        self.assertEqual(data_export.status_code, 200)
        self.assertIn("API 数据模拟".encode(), data_export.content)
        self.assertEqual(RehearsalInstance.objects.get(id=instance["id"]).status, RehearsalInstanceStatus.READY)

    def test_api_scene_run_and_work_order_closed_loop(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:scene-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API 场景预演",
                },
                format="json",
            )
        )["instance"]
        batch = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/complaint/generate",
                {
                    "request_id": "api:scene-batch:001",
                    "template_id": dashboard["templates"][0]["id"],
                    "count": 3,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )["batch"]
        scene = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/scene-run",
                {
                    "request_id": "api:scene-run:001",
                    "script_id": dashboard["scene_scripts"][0]["id"],
                    "complaint_batch_ids": [batch["id"]],
                },
                format="json",
            )
        )["run"]
        started = self.response_data(
            self.client.post(
                f"{prefix}/scene-run/{scene['id']}/operate",
                {"action": "START"},
                format="json",
            )
        )["run"]
        order = self.response_data(
            self.client.post(
                f"{prefix}/scene-run/{scene['id']}/work-order",
                {
                    "request_id": "api:scene-order:001",
                    "title": "API 场景遗留事项",
                    "priority": "MEDIUM",
                },
                format="json",
            )
        )["work_order"]
        accepted = self.response_data(
            self.client.post(
                f"{prefix}/work-order/{order['id']}/operate",
                {"action": "ACCEPT", "assignee": "API 验收员"},
                format="json",
            )
        )["work_order"]
        scene_list = self.response_data(self.client.get(f"{prefix}/scene-run", {"status": "RUNNING"}))
        scene_detail = self.response_data(self.client.get(f"{prefix}/scene-run/{scene['id']}"))
        order_list = self.response_data(self.client.get(f"{prefix}/work-order", {"status": "PROCESSING"}))

        self.assertEqual(started["status"], SceneRunStatus.RUNNING)
        self.assertEqual(accepted["status"], WorkOrderStatus.PROCESSING)
        self.assertEqual(scene_list["total"], 1)
        self.assertEqual(len(scene_detail["steps"]), 5)
        self.assertEqual(len(scene_detail["work_orders"]), 1)
        self.assertEqual(order_list["total"], 1)

    def test_api_scene_library_versioning_and_run(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:library-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API 场景库预演",
                },
                format="json",
            )
        )["instance"]
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/complaint/generate",
                {
                    "request_id": "api:library-batch:001",
                    "template_id": dashboard["templates"][0]["id"],
                    "count": 3,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/operate",
                {"action": "PREPARE"},
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/data-simulation",
                {
                    "request_id": "api:library-data:001",
                    "template_id": dashboard["data_templates"][0]["id"],
                    "name": "API 场景库时序指标",
                    "count": 6,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/process",
                {
                    "request_id": "api:library-process:001",
                    "template_id": dashboard["process_templates"][0]["id"],
                    "name": "API 场景库处置流程",
                },
                format="json",
            )
        )
        created = self.response_data(
            self.client.post(
                f"{prefix}/scene-library",
                {
                    "request_id": "api:library:create:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "script_id": dashboard["scene_scripts"][0]["id"],
                    "code": "API-LIBRARY-NETWORK",
                    "name": "API 网络故障场景库",
                    "default_instance_id": instance["id"],
                    "process_template_id": dashboard["process_templates"][0]["id"],
                    "configuration": {"acceptance": "api"},
                },
                format="json",
            )
        )
        library_id = created["library"]["id"]
        listed = self.response_data(self.client.get(f"{prefix}/scene-library", {"keyword": "API 网络"}))
        cloned = self.response_data(
            self.client.post(
                f"{prefix}/scene-library/{library_id}/clone",
                {"request_id": "api:library:clone:001", "name": "API 网络故障场景库 V2"},
                format="json",
            )
        )
        archived = self.response_data(
            self.client.post(
                f"{prefix}/scene-library/{cloned['library']['id']}/operate",
                {"action": "ARCHIVED", "note": "API 归档"},
                format="json",
            )
        )
        run = self.response_data(
            self.client.post(
                f"{prefix}/scene-library/{library_id}/run",
                {"request_id": "api:library:run:001"},
                format="json",
            )
        )
        detail = self.response_data(self.client.get(f"{prefix}/scene-library/{library_id}"))
        instance_detail = self.response_data(self.client.get(f"{prefix}/instance/{instance['id']}"))
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertTrue(created["created"])
        self.assertEqual(listed["total"], 1)
        self.assertEqual(cloned["library"]["version"], created["library"]["version"] + 1)
        self.assertEqual(archived["library"]["status"], SceneLibraryStatus.ARCHIVED)
        self.assertTrue(run["created"])
        self.assertEqual(run["run"]["library_id"], library_id)
        self.assertEqual(detail["run_count"], 1)
        self.assertEqual(len(detail["scene_runs"]), 1)
        self.assertEqual(len(instance_detail["scene_libraries"]), 2)
        self.assertGreaterEqual(updated_dashboard["summary"]["scene_library_count"], 3)

    def test_api_alert_lifecycle_and_unified_timeline(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:alert-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API 告警时序预演",
                },
                format="json",
            )
        )["instance"]
        alert = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/alert",
                {
                    "request_id": "api:alert:001",
                    "template_id": dashboard["alert_templates"][0]["id"],
                    "rule_code": "LATENCY_HIGH",
                    "resource_name": "API 核心路由器",
                    "metric_value": 128,
                },
                format="json",
            )
        )["alert"]
        active = self.response_data(
            self.client.post(
                f"{prefix}/alert/{alert['id']}/operate",
                {"action": "REPORT"},
                format="json",
            )
        )["alert"]
        acknowledged = self.response_data(
            self.client.post(
                f"{prefix}/alert/{alert['id']}/operate",
                {"action": "ACKNOWLEDGE"},
                format="json",
            )
        )["alert"]
        recovered = self.response_data(
            self.client.post(
                f"{prefix}/alert/{alert['id']}/operate",
                {"action": "RECOVER", "metric_value": 68},
                format="json",
            )
        )["alert"]
        alert_list = self.response_data(self.client.get(f"{prefix}/alert", {"status": "RECOVERED"}))
        alert_detail = self.response_data(self.client.get(f"{prefix}/alert/{alert['id']}"))
        timeline = self.response_data(self.client.get(f"{prefix}/instance/{instance['id']}/timeline"))

        self.assertEqual(active["status"], AlertStatus.ACTIVE)
        self.assertEqual(acknowledged["status"], AlertStatus.ACKNOWLEDGED)
        self.assertEqual(recovered["status"], AlertStatus.RECOVERED)
        self.assertEqual(alert_list["total"], 1)
        self.assertEqual(len(alert_detail["events"]), 4)
        self.assertEqual(timeline["kind_counts"]["ALERT"], 4)

    def test_api_data_model_build_detail_and_load(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:data-model-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API 数据模型预演",
                },
                format="json",
            )
        )["instance"]
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/complaint/generate",
                {
                    "request_id": "api:data-model-batch:001",
                    "template_id": dashboard["templates"][0]["id"],
                    "count": 4,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/operate",
                {"action": "PREPARE"},
                format="json",
            )
        )
        data_batch = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/data-simulation",
                {
                    "request_id": "api:data-model-metric:001",
                    "template_id": dashboard["data_templates"][0]["id"],
                    "count": 6,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )["batch"]
        created = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/data-model",
                {
                    "request_id": "api:data-model:001",
                    "name": "API 统一预演数据模型",
                    "build_strategy": "FULL_INSTANCE",
                    "source_scope": {"batch_id": data_batch["id"]},
                },
                format="json",
            )
        )
        model_id = created["data_model"]["id"]
        listed = self.response_data(self.client.get(f"{prefix}/data-model", {"quality_status": "PASS"}))
        detail = self.response_data(self.client.get(f"{prefix}/data-model/{model_id}"))
        loaded = self.response_data(
            self.client.post(
                f"{prefix}/data-model/{model_id}/load",
                {"request_id": "api:data-model:load:002", "note": "API 刷新"},
                format="json",
            )
        )
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertTrue(created["created"])
        self.assertEqual(created["data_model"]["quality_status"], DataModelQualityStatus.PASS)
        self.assertEqual(listed["total"], 1)
        self.assertEqual(len(detail["versions"]), 1)
        self.assertEqual(loaded["data_model"]["current_version"], 2)
        self.assertEqual(updated_dashboard["summary"]["data_model_count"], 1)
        self.assertEqual(updated_dashboard["summary"]["data_model_version_count"], 2)

    def test_api_traceback_create_list_detail_and_export(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:traceback-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API 数据回溯预演",
                },
                format="json",
            )
        )["instance"]
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/complaint/generate",
                {
                    "request_id": "api:traceback-complaint:001",
                    "template_id": dashboard["templates"][0]["id"],
                    "count": 3,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/operate",
                {"action": "PREPARE"},
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/data-simulation",
                {
                    "request_id": "api:traceback-metric:001",
                    "template_id": dashboard["data_templates"][0]["id"],
                    "count": 5,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )
        created = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/traceback",
                {
                    "request_id": "api:traceback:create:001",
                    "name": "API X25 动态回溯",
                    "scope": "FULL_INSTANCE",
                    "kind_scope": ["COMPLAINT", "METRIC"],
                },
                format="json",
            )
        )
        traceback_id = created["traceback"]["id"]
        listed = self.response_data(
            self.client.get(
                f"{prefix}/traceback",
                {"instance_id": instance["id"], "kind": "METRIC", "keyword": "X25"},
            )
        )
        detail = self.response_data(self.client.get(f"{prefix}/traceback/{traceback_id}"))
        exported = self.client.get(f"{prefix}/traceback/export", {"instance_id": instance["id"]})
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertTrue(created["created"])
        self.assertEqual(created["traceback"]["total_events"], 8)
        self.assertEqual(listed["total"], 1)
        self.assertEqual(len(detail["nodes"]), 2)
        self.assertGreaterEqual(len(detail["items"]), 8)
        self.assertEqual(exported.status_code, 200)
        self.assertTrue(exported.content.startswith(b"\xef\xbb\xbf"))
        self.assertIn("API X25 动态回溯".encode(), exported.content)
        self.assertEqual(updated_dashboard["summary"]["traceback_count"], 1)

    def test_api_simulation_record_list_detail_and_export(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:record-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API X26 模拟记录预演",
                },
                format="json",
            )
        )["instance"]
        batch = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/complaint/generate",
                {
                    "request_id": "api:record-batch:001",
                    "template_id": dashboard["templates"][0]["id"],
                    "count": 3,
                    "start_at": timezone.now().isoformat(),
                    "interval_minutes": 5,
                },
                format="json",
            )
        )["batch"]

        listed = self.response_data(
            self.client.get(
                f"{prefix}/simulation-record",
                {"instance_id": instance["id"], "keyword": "X26", "page_size": 20},
            )
        )
        detail = self.response_data(
            self.client.get(f"{prefix}/simulation-record/COMPLAINT_BATCH/{batch['id']}")
        )
        exported = self.client.get(f"{prefix}/simulation-record/export", {"instance_id": instance["id"]})
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertEqual(listed["total"], 2)
        self.assertEqual({item["kind"] for item in listed["items"]}, {"INSTANCE", "COMPLAINT_BATCH"})
        self.assertEqual(detail["record"]["kind"], "COMPLAINT_BATCH")
        self.assertEqual(detail["activity_count"], 3)
        self.assertEqual(exported.status_code, 200)
        self.assertTrue(exported.content.startswith(b"\xef\xbb\xbf"))
        self.assertIn("API X26 模拟记录预演".encode(), exported.content)
        self.assertGreaterEqual(updated_dashboard["summary"]["simulation_record_count"], 2)

    def test_api_capability_mock_invoke_history_detail_and_replay(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:capability-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API X27 能力调用预演",
                },
                format="json",
            )
        )["instance"]
        capability = next(item for item in dashboard["capability_mocks"] if item["kind"] == "API")
        invoked = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/capability-call",
                {
                    "request_id": "api:capability-call:001",
                    "capability_id": capability["id"],
                    "request_payload": {"resource_id": "API-X27-NE-01"},
                    "warroom_group": "API 验收战情室",
                },
                format="json",
            )
        )
        call_id = invoked["call"]["id"]
        listed = self.response_data(
            self.client.get(
                f"{prefix}/capability-call",
                {"instance_id": instance["id"], "warroom_group": "API 验收战情室"},
            )
        )
        detail = self.response_data(self.client.get(f"{prefix}/capability-call/{call_id}"))
        replayed = self.response_data(
            self.client.post(
                f"{prefix}/capability-call/{call_id}/replay",
                {
                    "request_id": "api:capability-replay:001",
                    "expected_status": "FAILED",
                    "latency_ms": 480,
                },
                format="json",
            )
        )
        records = self.response_data(
            self.client.get(
                f"{prefix}/simulation-record",
                {"instance_id": instance["id"], "kind": "CAPABILITY_CALL"},
            )
        )
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertTrue(invoked["created"])
        self.assertEqual(invoked["call"]["status"], CapabilityCallStatus.SUCCEEDED)
        self.assertEqual(listed["summary"]["total"], 1)
        self.assertEqual(listed["warrooms"][0]["warroom_group"], "API 验收战情室")
        self.assertEqual(detail["request_payload"]["resource_id"], "API-X27-NE-01")
        self.assertEqual(replayed["call"]["source"], "REPLAY")
        self.assertEqual(replayed["call"]["status"], CapabilityCallStatus.FAILED)
        self.assertEqual(records["total"], 2)
        self.assertEqual(updated_dashboard["summary"]["capability_mock_count"], 2)
        self.assertEqual(updated_dashboard["summary"]["capability_call_count"], 2)

    def test_api_fault_configuration_clone_validate_apply_and_detail(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        source = dashboard["fault_configurations"][0]
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:fault-config-instance:001",
                    "scenario_id": source["scenario_id"],
                    "name": "API X28 故障配置预演",
                },
                format="json",
            )
        )["instance"]
        cloned = self.response_data(
            self.client.post(
                f"{prefix}/fault-configuration/{source['id']}/clone",
                {
                    "request_id": "api:fault-config-clone:001",
                    "name": "API 核心网络故障配置 V2",
                    "note": "API 版本化验收",
                },
                format="json",
            )
        )["configuration"]
        validated = self.response_data(
            self.client.post(
                f"{prefix}/fault-configuration/{cloned['id']}/validate",
                {"note": "API 主动复核"},
                format="json",
            )
        )["configuration"]
        applied = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/fault-configuration",
                {
                    "request_id": "api:fault-config-apply:001",
                    "configuration_id": cloned["id"],
                    "note": "应用到 API 验收实例",
                },
                format="json",
            )
        )
        listed = self.response_data(
            self.client.get(f"{prefix}/fault-configuration", {"bound": "true", "page_size": 20})
        )
        detail = self.response_data(self.client.get(f"{prefix}/fault-configuration/{cloned['id']}"))
        instance_detail = self.response_data(self.client.get(f"{prefix}/instance/{instance['id']}"))

        self.assertEqual(cloned["version"], source["version"] + 1)
        self.assertEqual(validated["status"], FaultConfigurationStatus.READY)
        self.assertTrue(validated["validation_report"]["passed"])
        self.assertTrue(applied["created"])
        self.assertEqual(applied["instance"]["fault_configuration_id"], cloned["id"])
        self.assertEqual(listed["total"], 1)
        self.assertEqual(listed["summary"]["bound_instances"], 1)
        self.assertEqual(detail["bound_instances"][0]["id"], instance["id"])
        self.assertGreaterEqual(detail["event_count"], 4)
        self.assertEqual(instance_detail["fault_configuration_snapshot"]["code"], source["code"])

    def test_api_scene_replay_create_list_detail_and_dashboard(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        configuration = dashboard["fault_configurations"][0]
        capability = next(item for item in dashboard["capability_mocks"] if item["kind"] == "API")
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:scene-replay-instance:001",
                    "scenario_id": configuration["scenario_id"],
                    "name": "API X29 场景动态回溯",
                },
                format="json",
            )
        )["instance"]
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/fault-configuration",
                {
                    "request_id": "api:scene-replay-config:001",
                    "configuration_id": configuration["id"],
                },
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/capability-call",
                {
                    "request_id": "api:scene-replay-call:001",
                    "capability_id": capability["id"],
                    "request_payload": {"source_node": "GW-01", "resource_id": "CORE-01"},
                    "expected_status": "TIMEOUT",
                    "latency_ms": 1800,
                },
                format="json",
            )
        )
        created = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/scene-replay",
                {
                    "request_id": "api:scene-replay-create:001",
                    "name": "API X29 核心链路镜像回溯",
                },
                format="json",
            )
        )
        replay_id = created["replay"]["id"]
        listed = self.response_data(
            self.client.get(
                f"{prefix}/scene-replay",
                {"instance_id": instance["id"], "node_id": "CORE-01", "keyword": "X29"},
            )
        )
        detail = self.response_data(self.client.get(f"{prefix}/scene-replay/{replay_id}"))
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertTrue(created["created"])
        self.assertEqual(created["replay"]["traffic_frame_count"], 1)
        self.assertEqual(listed["total"], 1)
        self.assertEqual(detail["traffic_frames"][0]["status"], CapabilityCallStatus.TIMEOUT)
        self.assertGreaterEqual(detail["changed_node_count"], 1)
        self.assertEqual(updated_dashboard["summary"]["scene_replay_count"], 1)
        self.assertEqual(updated_dashboard["summary"]["scene_replay_frame_count"], 1)

    def test_api_fault_report_create_operate_list_detail_and_dashboard(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        configuration = dashboard["fault_configurations"][0]
        template = dashboard["fault_report_templates"][0]
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:fault-report-instance:001",
                    "scenario_id": configuration["scenario_id"],
                    "name": "API X30 故障上报模拟",
                },
                format="json",
            )
        )["instance"]
        self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/fault-configuration",
                {
                    "request_id": "api:fault-report-config:001",
                    "configuration_id": configuration["id"],
                },
                format="json",
            )
        )
        created = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/fault-report",
                {
                    "request_id": "api:fault-report-create:001",
                    "template_id": template["id"],
                    "title": "API 核心节点故障上报",
                    "description": "核心节点高时延，启动故障上报模拟。",
                    "impact_scope": "东区模拟客户流量",
                    "level": FaultReportLevel.LEVEL_3,
                    "attachments": [
                        {
                            "reference_id": "api:evidence:001",
                            "file_name": "api-alarm.json",
                            "content_type": "application/json",
                            "file_size": 256,
                        }
                    ],
                },
                format="json",
            )
        )
        report_id = created["report"]["id"]
        self.response_data(
            self.client.post(
                f"{prefix}/fault-report/{report_id}/operate",
                {"action": FaultReportAction.SUBMITTED},
                format="json",
            )
        )
        self.response_data(
            self.client.post(
                f"{prefix}/fault-report/{report_id}/operate",
                {"action": FaultReportAction.ACCEPTED},
                format="json",
            )
        )
        operated = self.response_data(
            self.client.post(
                f"{prefix}/fault-report/{report_id}/operate",
                {"action": FaultReportAction.STAGE_ADVANCED, "note": "API 首次上报已确认"},
                format="json",
            )
        )
        listed = self.response_data(
            self.client.get(
                f"{prefix}/fault-report",
                {"instance_id": instance["id"], "status": FaultReportStatus.PROCESSING},
            )
        )
        detail = self.response_data(self.client.get(f"{prefix}/fault-report/{report_id}"))
        templates = self.response_data(self.client.get(f"{prefix}/fault-report-template"))
        records = self.response_data(
            self.client.get(
                f"{prefix}/simulation-record",
                {"instance_id": instance["id"], "kind": "FAULT_REPORT"},
            )
        )
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertTrue(created["created"])
        self.assertEqual(operated["report"]["status"], FaultReportStatus.PROCESSING)
        self.assertEqual(operated["report"]["current_stage_code"], "IMPACT_CONFIRM")
        self.assertGreater(operated["report"]["remaining_seconds"], 0)
        self.assertEqual(listed["total"], 1)
        self.assertEqual(listed["summary"]["active"], 1)
        self.assertEqual(detail["attachment_count"], 1)
        self.assertEqual(detail["event_count"], 4)
        self.assertGreaterEqual(len(templates), 1)
        self.assertEqual(records["total"], 1)
        self.assertEqual(updated_dashboard["summary"]["fault_report_count"], 1)
        self.assertEqual(updated_dashboard["summary"]["fault_report_active_count"], 1)

    def test_api_auxiliary_simulation_create_query_detail_and_download(self):
        prefix = f"/admin/api/workspace/{self.workspace_id}/rehearsal"
        dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))
        instance = self.response_data(
            self.client.post(
                f"{prefix}/instance",
                {
                    "request_id": "api:auxiliary-instance:001",
                    "scenario_id": dashboard["scenarios"][0]["id"],
                    "name": "API X31 辅助模拟",
                },
                format="json",
            )
        )["instance"]
        created = self.response_data(
            self.client.post(
                f"{prefix}/instance/{instance['id']}/auxiliary-simulation",
                {
                    "request_id": "api:auxiliary-create:001",
                    "name": "API 随机因子联动模拟",
                    "seed": 310031,
                    "duration_minutes": 20,
                    "factors": [
                        {"code": "NETWORK_JITTER", "trigger_probability": 1},
                        {"code": "FAULT_SPREAD", "trigger_probability": 1},
                        {"code": "TRAFFIC_PEAK", "trigger_probability": 1},
                    ],
                },
                format="json",
            )
        )
        simulation_id = created["simulation"]["id"]
        listed = self.response_data(
            self.client.get(f"{prefix}/auxiliary-simulation", {"instance_id": instance["id"], "seed": 310031})
        )
        detail = self.response_data(self.client.get(f"{prefix}/auxiliary-simulation/{simulation_id}"))
        records = self.response_data(
            self.client.get(f"{prefix}/simulation-record", {"kind": "AUXILIARY_SIMULATION"})
        )
        pdf_response = self.client.get(
            f"{prefix}/auxiliary-simulation/{simulation_id}/export",
            {"export_format": "pdf", "report_level": "summary"},
        )
        xlsx_response = self.client.get(
            f"{prefix}/auxiliary-simulation/{simulation_id}/export",
            {"export_format": "xlsx", "report_level": "full"},
        )
        json_response = self.client.get(
            f"{prefix}/auxiliary-simulation/{simulation_id}/export",
            {"export_format": "json", "report_level": "full"},
        )
        updated_dashboard = self.response_data(self.client.get(f"{prefix}/dashboard"))

        self.assertTrue(created["created"])
        self.assertEqual(listed["total"], 1)
        self.assertEqual(detail["seed"], 310031)
        self.assertEqual(len(detail["points"]), 21)
        self.assertEqual(records["total"], 1)
        self.assertEqual(pdf_response.status_code, 200)
        self.assertTrue(pdf_response.content.startswith(b"%PDF"))
        self.assertEqual(xlsx_response.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(xlsx_response.content), read_only=True)["分钟级原始数据"].max_row, 22)
        self.assertEqual(json.loads(json_response.content)["report_level"], "full")
        self.assertEqual(updated_dashboard["summary"]["auxiliary_simulation_count"], 1)
