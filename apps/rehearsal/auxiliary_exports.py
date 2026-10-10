# coding=utf-8
import json
from html import escape
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.shapes import Drawing, Line, PolyLine, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from marketing.report_exports import REPORT_FONT_NAME, REPORT_FONT_PATH


AUXILIARY_EXPORT_CONTENT_TYPES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json; charset=utf-8",
}


def auxiliary_export_filename(simulation, export_format, report_level):
    return f"{simulation.code}-{report_level}.{export_format}"


def _payload(simulation, report_level):
    result = simulation.result_snapshot or {}
    payload = {
        "report_level": report_level,
        "basic_information": {
            "code": simulation.code,
            "name": simulation.name,
            "scenario_name": simulation.instance.scenario.name,
            "instance_code": simulation.instance.code,
            "instance_name": simulation.instance.name,
            "seed": simulation.seed,
            "duration_minutes": simulation.duration_minutes,
            "status": simulation.get_status_display(),
            "started_at": simulation.started_at.isoformat(),
            "completed_at": simulation.completed_at.isoformat(),
            "checksum": simulation.checksum,
        },
        "random_factors": simulation.factor_snapshot,
        "summary": result.get("summary", {}),
        "impact_analysis": result.get("impact_analysis", []),
        "statistics": result.get("statistics", []),
        "anomalies": result.get("anomalies", []),
        "key_moments": result.get("key_moments", []),
    }
    if report_level == "full":
        payload["baseline"] = result.get("baseline", {})
        payload["raw_points"] = result.get("points", [])
    return payload


def build_auxiliary_json(simulation, report_level):
    return json.dumps(_payload(simulation, report_level), ensure_ascii=False, indent=2).encode("utf-8")


def _style_sheet(sheet, widths):
    header_fill = PatternFill("solid", fgColor="123638")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for key, value in widths.items():
        sheet.column_dimensions[key].width = value


def build_auxiliary_xlsx(simulation, report_level):
    payload = _payload(simulation, report_level)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "报告概览"
    summary_sheet.append(["项目", "结果"])
    for key, value in payload["basic_information"].items():
        summary_sheet.append([key, value])
    summary_sheet.append(["结论", payload["summary"].get("conclusion", "")])
    _style_sheet(summary_sheet, {"A": 24, "B": 88})

    factor_sheet = workbook.create_sheet("随机因子")
    factor_sheet.append(["分类", "因子", "分布", "范围", "触发概率", "是否触发", "采样值", "触发分钟"])
    for item in payload["random_factors"]:
        factor_sheet.append(
            [
                item["category"],
                item["name"],
                item["distribution"],
                f"{item['min_value']}~{item['max_value']} {item['unit']}",
                item["trigger_probability"],
                "是" if item.get("triggered") else "否",
                item.get("sampled_value", 0),
                item.get("trigger_minute", ""),
            ]
        )
    _style_sheet(factor_sheet, {"A": 14, "B": 22, "C": 16, "D": 24, "E": 14, "F": 12, "G": 14, "H": 14})

    statistics_sheet = workbook.create_sheet("统计对比")
    statistics_sheet.append(["指标", "注入前基线", "注入后均值", "峰值", "方差", "均值变化(%)"])
    for item in payload["statistics"]:
        statistics_sheet.append(
            [item["name"], item["baseline"], item["mean"], item["peak"], item["variance"], item["delta_percent"]]
        )
    _style_sheet(statistics_sheet, {"A": 24, "B": 16, "C": 16, "D": 16, "E": 16, "F": 18})

    anomaly_sheet = workbook.create_sheet("异常点")
    anomaly_sheet.append(["分钟", "发生时间", "异常原因"])
    for item in payload["anomalies"]:
        anomaly_sheet.append([item["minute"], item["occurred_at"], "；".join(item["reasons"])])
    _style_sheet(anomaly_sheet, {"A": 12, "B": 30, "C": 88})

    if report_level == "full":
        point_sheet = workbook.create_sheet("分钟级原始数据")
        point_sheet.append(
            ["分钟", "发生时间", "时延(ms)", "丢包率(%)", "流量倍率", "申告量", "工单积压", "超时率(%)"]
        )
        for item in payload["raw_points"]:
            point_sheet.append(
                [
                    item["minute"], item["occurred_at"], item["latency_ms"], item["packet_loss_pct"],
                    item["traffic_multiplier"], item["complaint_count"], item["ticket_backlog"],
                    item["timeout_rate_pct"],
                ]
            )
        _style_sheet(point_sheet, {"A": 10, "B": 30, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14})

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _table_style():
    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123638")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CAD5D4")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7F6")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
    )


def _trend_chart(points):
    drawing = Drawing(165 * mm, 64 * mm)
    left, bottom, width, height = 34, 24, 420, 132
    drawing.add(Rect(left, bottom, width, height, fillColor=colors.HexColor("#F7FAF9"), strokeColor=colors.HexColor("#CBD7D4")))
    drawing.add(String(left, height + bottom + 10, "分钟级影响趋势（归一化）", fontName=REPORT_FONT_NAME, fontSize=9, fillColor=colors.HexColor("#123638")))
    series = [
        ("latency_ms", colors.HexColor("#E56B4A"), "时延"),
        ("packet_loss_pct", colors.HexColor("#D9A441"), "丢包"),
        ("traffic_multiplier", colors.HexColor("#0D7C78"), "流量"),
        ("timeout_rate_pct", colors.HexColor("#5271A3"), "超时"),
    ]
    sampled = points[:: max(1, len(points) // 80)]
    for series_index, (key, color, label) in enumerate(series):
        values = [float(item[key]) for item in sampled]
        maximum = max(values) if values else 1
        minimum = min(values) if values else 0
        span = maximum - minimum or 1
        coords = []
        for index, value in enumerate(values):
            x = left + (index / max(len(values) - 1, 1)) * width
            y = bottom + ((value - minimum) / span) * height
            coords.extend([x, y])
        drawing.add(PolyLine(coords, strokeColor=color, strokeWidth=1.5))
        legend_x = left + 245 + series_index * 48
        drawing.add(Line(legend_x, height + bottom + 13, legend_x + 12, height + bottom + 13, strokeColor=color, strokeWidth=2))
        drawing.add(String(legend_x + 15, height + bottom + 9, label, fontName=REPORT_FONT_NAME, fontSize=7, fillColor=colors.HexColor("#43545A")))
    drawing.add(String(left, 8, "0", fontName=REPORT_FONT_NAME, fontSize=7))
    drawing.add(String(left + width - 18, 8, f"{points[-1]['minute']} 分", fontName=REPORT_FONT_NAME, fontSize=7))
    return drawing


def build_auxiliary_pdf(simulation, report_level):
    payload = _payload(simulation, report_level)
    if REPORT_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(REPORT_FONT_NAME, REPORT_FONT_PATH))
    output = BytesIO()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("AuxBody", parent=styles["BodyText"], fontName=REPORT_FONT_NAME, fontSize=8.5, leading=13, textColor=colors.HexColor("#35464A"))
    title = ParagraphStyle("AuxTitle", parent=body, fontSize=20, leading=28, alignment=TA_CENTER, textColor=colors.HexColor("#123638"), spaceAfter=6)
    heading = ParagraphStyle("AuxHeading", parent=body, fontSize=12.5, leading=18, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#0D7C78"))
    compact_body = ParagraphStyle("AuxCompact", parent=body, fontSize=7.2, leading=9)
    table_header = ParagraphStyle(
        "AuxTableHeader",
        parent=body,
        textColor=colors.white,
        alignment=TA_CENTER,
    )
    compact_header = ParagraphStyle(
        "AuxCompactHeader",
        parent=compact_body,
        textColor=colors.white,
        alignment=TA_CENTER,
    )

    def paragraph(value):
        return Paragraph(escape(str(value)), body)

    def header_paragraph(value):
        return Paragraph(escape(str(value)), table_header)

    def page(canvas, document):
        canvas.saveState()
        canvas.setFont(REPORT_FONT_NAME, 8)
        canvas.setFillColor(colors.HexColor("#6B777B"))
        canvas.drawString(18 * mm, 11 * mm, f"辅助模拟结果报告 · {simulation.code}")
        canvas.drawRightString(A4[0] - 18 * mm, 11 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title=simulation.name, author="电信智能体故障与场景预演平台")
    basic = payload["basic_information"]
    story = [
        Paragraph("辅助模拟结果报告", title),
        Paragraph(f"{paragraph(simulation.name).text} · {'完整版' if report_level == 'full' else '精简版'}", body),
        Spacer(1, 4 * mm),
        Paragraph("基础信息", heading),
    ]
    basic_table = Table(
        [
            [header_paragraph("场景 / 实例"), header_paragraph(f"{basic['scenario_name']} / {basic['instance_name']}")],
            [paragraph("随机种子"), paragraph(basic["seed"])],
            [paragraph("模拟时长"), paragraph(f"{basic['duration_minutes']} 分钟")],
            [paragraph("状态 / 校验和"), paragraph(f"{basic['status']} / {basic['checksum'][:16]}…")],
            [paragraph("核心结论"), paragraph(payload["summary"].get("conclusion", ""))],
        ],
        colWidths=[38 * mm, 132 * mm],
    )
    basic_table.setStyle(_table_style())
    story.extend([basic_table, Paragraph("随机因子组合", heading)])
    factor_rows = [[header_paragraph(value) for value in ("类型", "因子", "规则", "触发", "采样")]]
    for item in payload["random_factors"]:
        factor_rows.append(
            [
                paragraph(item["category"]), paragraph(item["name"]), paragraph(item["distribution"]),
                paragraph("是" if item.get("triggered") else "否"),
                paragraph(f"{item.get('sampled_value', 0)} {item['unit']}"),
            ]
        )
    factor_table = Table(factor_rows, colWidths=[24 * mm, 42 * mm, 30 * mm, 20 * mm, 42 * mm], repeatRows=1)
    factor_table.setStyle(_table_style())
    story.extend([factor_table, Paragraph("关键指标趋势", heading), _trend_chart((simulation.result_snapshot or {}).get("points", []))])

    story.append(Paragraph("注入前后统计对比", heading))
    stat_rows = [[header_paragraph(value) for value in ("指标", "基线", "均值", "峰值", "方差", "变化")]]
    for item in payload["statistics"]:
        stat_rows.append([paragraph(item["name"]), paragraph(item["baseline"]), paragraph(item["mean"]), paragraph(item["peak"]), paragraph(item["variance"]), paragraph(f"{item['delta_percent']}%")])
    stat_table = Table(stat_rows, colWidths=[40 * mm, 24 * mm, 24 * mm, 24 * mm, 24 * mm, 24 * mm], repeatRows=1)
    stat_table.setStyle(_table_style())
    story.extend([stat_table, Paragraph("随机因子影响分析", heading)])
    story.extend(Paragraph(f"• {item}", body) for item in payload["impact_analysis"])
    story.append(Paragraph("异常点标注", heading))
    if payload["anomalies"]:
        anomaly_rows = [[header_paragraph(value) for value in ("分钟", "时间", "异常原因")]]
        anomaly_rows.extend([paragraph(item["minute"]), paragraph(item["occurred_at"]), paragraph("；".join(item["reasons"]))] for item in payload["anomalies"])
        anomaly_table = Table(anomaly_rows, colWidths=[18 * mm, 48 * mm, 104 * mm], repeatRows=1)
        anomaly_table.setStyle(_table_style())
        story.append(anomaly_table)
    else:
        story.append(Paragraph("本批次未识别到超阈值异常点。", body))

    if report_level == "full":
        story.extend([PageBreak(), Paragraph("分钟级原始数据", heading)])
        def compact_paragraph(value):
            return Paragraph(escape(str(value)), compact_body)

        def compact_header_paragraph(value):
            return Paragraph(escape(str(value)), compact_header)

        point_rows = [[compact_header_paragraph(value) for value in ("分钟", "时延", "丢包", "流量", "申告", "积压", "超时")]]
        for item in payload.get("raw_points", []):
            point_rows.append([compact_paragraph(item["minute"]), compact_paragraph(item["latency_ms"]), compact_paragraph(item["packet_loss_pct"]), compact_paragraph(item["traffic_multiplier"]), compact_paragraph(item["complaint_count"]), compact_paragraph(item["ticket_backlog"]), compact_paragraph(item["timeout_rate_pct"])])
        point_table = Table(point_rows, colWidths=[20 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm], repeatRows=1)
        point_table.setStyle(_table_style())
        point_table.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        story.append(point_table)
    document.build(story, onFirstPage=page, onLaterPages=page)
    return output.getvalue()


def build_auxiliary_export(simulation, export_format, report_level):
    if export_format == "json":
        return build_auxiliary_json(simulation, report_level)
    if export_format == "xlsx":
        return build_auxiliary_xlsx(simulation, report_level)
    if export_format == "pdf":
        return build_auxiliary_pdf(simulation, report_level)
    raise ValueError("不支持的辅助模拟报告格式")
